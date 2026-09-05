"""Locate, open, and append rows on Sassy_Closet_SoT.xlsx (header-aware).

Never invents mã. Never embeds pictures. Writes photo_link as text/hyperlink.
"""

from __future__ import annotations

import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Mapping
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent.parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (
    BODY_FONT,
    ONEDRIVE_SOT,
    apply_photo_link_style,
    header_row_for_sheet,
    normalize_header,
    read_headers,
    resolve_header_key,
    resolve_sheet_name,
    table_name_for_sheet,
)

PT = ZoneInfo("America/Los_Angeles")

ENV_SOT = "SASSY_SOT"
ENV_SOT_DIR = "SASSY_SOT_DIR"


class WorkbookError(Exception):
    """User-facing workbook problem (missing file, missing sheet, bad header)."""


def _home() -> Path:
    return Path.home()


def candidate_sot_paths(explicit: Path | None = None) -> list[Path]:
    """Ordered places Mini Boss / Kit might keep the live SoT book."""
    found: list[Path] = []
    if explicit is not None:
        found.append(Path(explicit).expanduser())
    env = os.environ.get(ENV_SOT)
    if env:
        found.append(Path(env).expanduser())
    env_dir = os.environ.get(ENV_SOT_DIR)
    if env_dir:
        found.append(Path(env_dir).expanduser() / "Sassy_Closet_SoT.xlsx")
    home = _home()
    extras = [
        home / "OneDrive" / ONEDRIVE_SOT,
        home / "OneDrive - Personal" / ONEDRIVE_SOT,
        home / "Library/CloudStorage/OneDrive-Personal" / ONEDRIVE_SOT,
        home / "Library/CloudStorage/OneDrive-Personal" / "Documents/Sassy Closet/Sassy_Closet_SoT.xlsx",
        Path.cwd() / "Sassy_Closet_SoT.xlsx",
        Path.cwd() / "out" / "Sassy_Closet_SoT.xlsx",
        Path.cwd() / "Documents/Sassy Closet/Sassy_Closet_SoT.xlsx",
    ]
    found.extend(extras)
    # Dedup while keeping order.
    seen: set[str] = set()
    unique: list[Path] = []
    for path in found:
        key = str(path)
        if key in seen:
            continue
        seen.add(key)
        unique.append(path)
    return unique


def locate_sot(explicit: Path | None = None, *, must_exist: bool = True) -> Path:
    """Return the SoT path. Prefer --workbook / $SASSY_SOT / OneDrive."""
    for path in candidate_sot_paths(explicit):
        if path.is_file():
            return path
    if explicit is not None:
        path = Path(explicit).expanduser()
        if must_exist and not path.is_file():
            raise WorkbookError(
                f"workbook not found: {path}. "
                f"Pass --workbook or set {ENV_SOT}. Kit syncs to OneDrive {ONEDRIVE_SOT}."
            )
        return path
    if not must_exist:
        return Path(ONEDRIVE_SOT).name
    searched = "\n  ".join(str(p) for p in candidate_sot_paths(explicit)[:8])
    raise WorkbookError(
        "Sassy_Closet_SoT.xlsx not found.\n"
        f"  Kit sync target: OneDrive {ONEDRIVE_SOT}\n"
        f"  Or set {ENV_SOT} / pass --workbook.\n"
        f"  Looked at:\n  {searched}"
    )


def open_sot(path: Path, *, data_only: bool = False, read_only: bool = False) -> Workbook:
    if not path.is_file():
        raise WorkbookError(f"workbook not found: {path}")
    return load_workbook(path, data_only=data_only, read_only=read_only)


def require_sheet(wb: Workbook, logical: str) -> Worksheet:
    name = resolve_sheet_name(wb.sheetnames, logical)
    if name is None:
        raise WorkbookError(
            f"sheet {logical!r} not found. Have: {wb.sheetnames}. "
            "The ONE desktop book is Sassy_Closet_SoT.xlsx (Official / Wishlist / "
            "Orders / Dashboard). Lean desktop books use Ma_List / Candidates."
        )
    return wb[name]


def optional_sheet(wb: Workbook, logical: str) -> Worksheet | None:
    name = resolve_sheet_name(wb.sheetnames, logical)
    return wb[name] if name else None


def sheet_headers(ws: Worksheet) -> tuple[int, list[str]]:
    header_row = header_row_for_sheet(ws.title)
    headers = read_headers(ws, header_row)
    if not headers:
        raise WorkbookError(f"{ws.title}: no headers on row {header_row}")
    return header_row, headers


def first_empty_data_row(ws: Worksheet, header_row: int, headers: list[str]) -> int:
    max_col = max(len(headers), 1)
    last = max(ws.max_row or header_row, header_row)
    first_empty = last + 1
    for row in range(header_row + 1, last + 1):
        values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
        if not any(v is not None and str(v).strip() != "" for v in values):
            first_empty = row
            break
    else:
        first_empty = last + 1
    return first_empty


def existing_column_values(ws: Worksheet, header_row: int, col_index0: int) -> list[str]:
    values: list[str] = []
    col = col_index0 + 1
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        value = ws.cell(row, col).value
        if value is None or str(value).strip() == "":
            continue
        values.append(str(value).strip())
    return values


def find_existing_ma_row(ws: Worksheet, ma: str) -> int | None:
    header_row, headers = sheet_headers(ws)
    idx = resolve_header_key(headers, "ma")
    if idx is None:
        return None
    wanted = ma.strip().upper()
    col = idx + 1
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        value = ws.cell(row, col).value
        if value is not None and str(value).strip().upper() == wanted:
            return row
    return None


def next_numeric_id(existing: list[str], *, prefix: str, width: int = 3) -> str:
    """Next WISH-001 / ORD-001 style id. Not a mã — never used as stock SKU."""
    highest = 0
    for raw in existing:
        text = raw.strip().upper()
        if text.startswith(prefix.upper()):
            tail = text[len(prefix) :].lstrip("-_")
        else:
            tail = text.lstrip("#")
        if tail.isdigit():
            highest = max(highest, int(tail))
    return f"{prefix}{highest + 1:0{width}d}"


def today_iso(now: datetime | None = None) -> str:
    stamp = now or datetime.now(PT)
    return stamp.date().isoformat()


def now_pt(now: datetime | None = None) -> datetime:
    stamp = now or datetime.now(PT)
    if stamp.tzinfo is None:
        return stamp.replace(tzinfo=PT)
    return stamp.astimezone(PT)


def write_cell(ws: Worksheet, row: int, col: int, header: str, value: object) -> None:
    cell = ws.cell(row, col)
    if normalize_header(header) == "photo_link":
        apply_photo_link_style(cell, value)
        return
    if isinstance(value, (datetime, date)) and not isinstance(value, datetime):
        cell.value = value
    else:
        cell.value = value
    if cell.font is None or cell.font.name is None:
        cell.font = BODY_FONT


def resize_table_and_filter(ws: Worksheet, header_row: int, last_row: int, ncols: int) -> None:
    last_row = max(last_row, header_row + 1)
    ref = f"A{header_row}:{get_column_letter(ncols)}{last_row}"
    table_name = table_name_for_sheet(ws.title)
    if table_name and table_name in ws.tables:
        table: Table = ws.tables[table_name]
        table.ref = ref
    if ws.auto_filter is not None:
        ws.auto_filter.ref = ref


def append_mapped_row(
    ws: Worksheet,
    fields: Mapping[str, object],
    *,
    skip_empty: bool = True,
) -> tuple[int, list[str]]:
    """Write logical fields onto the next empty row. Returns (row, written headers)."""
    header_row, headers = sheet_headers(ws)
    dest = first_empty_data_row(ws, header_row, headers)
    written: list[str] = []
    unknown: list[str] = []
    used_cols: set[int] = set()
    for logical, value in fields.items():
        if skip_empty and (value is None or (isinstance(value, str) and value.strip() == "")):
            continue
        idx = resolve_header_key(headers, logical)
        if idx is None:
            unknown.append(logical)
            continue
        if idx in used_cols:
            continue
        header = headers[idx]
        write_cell(ws, dest, idx + 1, header, value)
        used_cols.add(idx)
        written.append(header)
    resize_table_and_filter(ws, header_row, dest, max(len(headers), 1))
    if unknown:
        # Not fatal — lean desktop books have fewer columns than SoT Official.
        written.append(f"(skipped unknown on {ws.title}: {', '.join(unknown)})")
    return dest, written


def add_workbook_path_arg(parser) -> None:
    parser.add_argument(
        "--workbook",
        "-w",
        type=Path,
        default=None,
        help=f"Sassy_Closet_SoT.xlsx (default: ${ENV_SOT} or OneDrive {ONEDRIVE_SOT})",
    )


def add_dry_run_arg(parser) -> None:
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="validate and print the row, do not save",
    )
