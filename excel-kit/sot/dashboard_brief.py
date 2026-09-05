#!/usr/bin/env python3
"""Print the SoT morning brief from Dashboard formula values when cached.

openpyxl cannot calculate Excel formulas. If B43 (or the KPI cells) still
show formulas / blanks, Boss copies Dashboard!B43 from desktop Excel after
the book has been opened once.

Usage:
  python3 excel-kit/sot/dashboard_brief.py
  python3 excel-kit/sot/dashboard_brief.py -w ./out/Sassy_Closet_SoT.xlsx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl.workbook import Workbook

_KIT_DIR = Path(__file__).resolve().parent.parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (  # noqa: E402
    SOT_DASHBOARD_BRIEF_CELL,
    SOT_DASHBOARD_BRIEF_FORMULA,
    SOT_DASHBOARD_BRIEF_NOTE,
    SOT_DASHBOARD_ROWS,
    SQUARE_SOT_SHORT,
)

from sot.workbook import (  # noqa: E402
    WorkbookError,
    add_workbook_path_arg,
    locate_sot,
    open_sot,
    optional_sheet,
    require_sheet,
)


def _display(value: object) -> str:
    if value is None:
        return "(empty — Excel has not cached a value)"
    text = str(value).strip()
    if text.startswith("="):
        return f"(formula, not cached) {text}"
    return text


def _is_usable(value: object) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    return bool(text) and not text.startswith("=")


def read_brief(wb: Workbook) -> dict[str, object]:
    ws = require_sheet(wb, "dashboard")
    cells: dict[str, object] = {}
    for row, (label, _formula) in SOT_DASHBOARD_ROWS.items():
        cells[f"B{row}"] = {"label": label, "value": ws[f"B{row}"].value}
    cells[SOT_DASHBOARD_BRIEF_CELL] = {
        "label": ws["A43"].value or "Morning brief (copy this cell)",
        "value": ws[SOT_DASHBOARD_BRIEF_CELL].value,
    }
    return {"sheet": ws.title, "cells": cells}


def print_brief(path: Path, data_only_cells: dict[str, object], formula_cells: dict[str, object]) -> int:
    print(SQUARE_SOT_SHORT)
    print(f"workbook: {path}")
    print(f"dashboard: {data_only_cells.get('sheet') or formula_cells.get('sheet')}")
    print()
    brief = data_only_cells.get("cells", {}).get(SOT_DASHBOARD_BRIEF_CELL, {})
    formula_brief = formula_cells.get("cells", {}).get(SOT_DASHBOARD_BRIEF_CELL, {})
    brief_value = brief.get("value") if isinstance(brief, dict) else None
    formula_value = formula_brief.get("value") if isinstance(formula_brief, dict) else None

    print(f"=== {SOT_DASHBOARD_BRIEF_CELL} morning brief ===")
    if _is_usable(brief_value):
        print(_display(brief_value))
        cached = True
    else:
        cached = False
        print(_display(brief_value if brief_value is not None else formula_value))
        print()
        print(SOT_DASHBOARD_BRIEF_NOTE)
        print(f"Expected formula in {SOT_DASHBOARD_BRIEF_CELL}:")
        print(f"  {SOT_DASHBOARD_BRIEF_FORMULA}")

    print()
    print("=== KPI cells (B5–B18, next-mã B21–B27) ===")
    rows = sorted(SOT_DASHBOARD_ROWS)
    usable_kpis = 0
    for row in rows:
        coord = f"B{row}"
        label = SOT_DASHBOARD_ROWS[row][0]
        data_val = None
        formula_val = None
        data_map = data_only_cells.get("cells", {})
        form_map = formula_cells.get("cells", {})
        if isinstance(data_map, dict) and isinstance(data_map.get(coord), dict):
            data_val = data_map[coord].get("value")
        if isinstance(form_map, dict) and isinstance(form_map.get(coord), dict):
            formula_val = form_map[coord].get("value")
        shown = data_val if _is_usable(data_val) else (data_val if data_val is not None else formula_val)
        if _is_usable(data_val):
            usable_kpis += 1
        print(f"  {coord}  {label}: {_display(shown)}")

    print()
    if cached or usable_kpis:
        print(f"cached values read: brief={'yes' if cached else 'no'}, KPI cells={usable_kpis}")
        return 0
    print("no cached formula values — Boss copies Dashboard!B43 from Excel.")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Print morning brief numbers from Dashboard cached values, "
            f"or document that Boss copies {SOT_DASHBOARD_BRIEF_CELL}."
        ),
    )
    add_workbook_path_arg(parser)
    args = parser.parse_args(argv)
    try:
        path = locate_sot(args.workbook)
        data_wb = open_sot(path, data_only=True)
        if optional_sheet(data_wb, "dashboard") is None:
            print(
                "error: no Dashboard sheet. Morning brief lives on "
                f"Sassy_Closet_SoT.xlsx → Dashboard!{SOT_DASHBOARD_BRIEF_CELL}. "
                "Boss copies B43 after opening the book in Excel.",
                file=sys.stderr,
            )
            data_wb.close()
            return 2
        data_cells = read_brief(data_wb)
        data_wb.close()
        formula_wb = open_sot(path, data_only=False)
        formula_cells = read_brief(formula_wb)
        formula_wb.close()
        return print_brief(path, data_cells, formula_cells)
    except WorkbookError as exc:
        print(f"error: {exc}", file=sys.stderr)
        print(
            f"Without a local workbook, Boss still copies Dashboard!{SOT_DASHBOARD_BRIEF_CELL} "
            "from the OneDrive SoT book after Excel has calculated it.",
            file=sys.stderr,
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
