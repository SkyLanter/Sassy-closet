#!/usr/bin/env python3
"""Build an empty Sassy_Closet_SoT.xlsx (the ONE desktop working copy).

Headers + Dashboard formulas + Lists dropdowns. Zero live inventory rows.
Square Free stays on-hand SoT. photo_link last on Wishlist. No embeds.

Usage:
  python3 excel-kit/sot/build_sot_desktop.py --out-dir ./out
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent.parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (  # noqa: E402
    BLUSH,
    BLUSH_ROW,
    BODY_FONT,
    BOT_ACTIVITY,
    BOT_ACTIVITY_WIDTHS,
    BOTS_DRAFT_ONLY,
    INK,
    MA_PREFIX_MEANS,
    MUTED,
    ONEDRIVE_PHOTOS,
    ONEDRIVE_SOT,
    ORDER_STATUS,
    PHOTO_RULE,
    SECTION,
    SOT_CHANNELS,
    SOT_CONDITION,
    SOT_DASHBOARD_BRIEF_CELL,
    SOT_DASHBOARD_BRIEF_FORMULA,
    SOT_DASHBOARD_ROWS,
    SOT_FULFILL,
    SOT_OFFICIAL,
    SOT_OFFICIAL_STATUS,
    SOT_ORDERS,
    SOT_PAY,
    SOT_WISHLIST,
    SOT_WISHLIST_STATUS,
    SQUARE_SOT_LINE,
    SQUARE_SOT_SHORT,
    YELLOW_INPUT,
    add_list_dropdown,
    apply_widths,
    header_index,
    new_boutique_workbook,
    paint_tab,
    strip_workbook_images,
    style_header_row,
    write_how_to_use,
)

DIVIDER = "·  ·  ·  ·  ·  ·  ·  ·  ·  ·  ·  ·"
YELLOW = PatternFill("solid", fgColor=YELLOW_INPUT)
HEADER_ROW = 4
RUNWAY = 8


def _title_block(ws: Worksheet, title: str, subtitle: str) -> None:
    ws.merge_cells("A1:D1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=INK)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUSH)
    ws.merge_cells("A2:D2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    ws.merge_cells("A3:D3")
    ws["A3"] = SQUARE_SOT_SHORT
    ws["A3"].font = Font(name="Calibri", size=10, color=INK)
    ws["A3"].fill = PatternFill("solid", fgColor=SECTION)
    paint_tab(ws)


def _input_runway(ws: Worksheet, headers: tuple[str, ...], header_row: int, rows: int) -> None:
    fill = PatternFill("solid", fgColor=BLUSH_ROW)
    for row in range(header_row + 1, header_row + 1 + rows):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col, None)
            cell.font = BODY_FONT
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")
    last = header_row + rows
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last}"


def _add_table(ws: Worksheet, name: str, headers: tuple[str, ...], header_row: int, rows: int) -> None:
    last = header_row + max(rows, 1)
    ref = f"A{header_row}:{get_column_letter(len(headers))}{last}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _widths_for(headers: tuple[str, ...], default: float = 14) -> dict[str, float]:
    widths: dict[str, float] = {}
    special = {
        "ma": 10,
        "photo_link": 36,
        "source": 36,
        "notes": 24,
        "name_vi": 22,
        "name_en": 22,
        "what_vi": 22,
        "what_en": 22,
        "snapshot": 20,
        "thread_note": 22,
        "square_name": 16,
        "caption_ready": 16,
    }
    for i, name in enumerate(headers, start=1):
        widths[get_column_letter(i)] = special.get(name, default)
    return widths


def _build_lists(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    columns = [
        ("official_status", SOT_OFFICIAL_STATUS),
        ("wishlist_status", SOT_WISHLIST_STATUS),
        ("order_status", ORDER_STATUS),
        ("condition", SOT_CONDITION),
        ("channel", SOT_CHANNELS),
        ("pay", SOT_PAY),
        ("fulfill", SOT_FULFILL),
        ("prefix", tuple(MA_PREFIX_MEANS)),
    ]
    for col, (title, values) in enumerate(columns, start=1):
        head = ws.cell(1, col, title)
        head.font = Font(name="Calibri", size=10, bold=True, color=INK)
        head.fill = PatternFill("solid", fgColor=BLUSH)
        for row, value in enumerate(values, start=2):
            ws.cell(row, col, value).font = BODY_FONT
        ws.column_dimensions[get_column_letter(col)].width = 22
    paint_tab(ws)


def _build_official(wb: Workbook) -> None:
    ws = wb.create_sheet("Official")
    _title_block(
        ws,
        "Official — owned pieces (working copy)",
        "Square Free = on-hand SoT. This table is mã / captions / photos — not a second warehouse.",
    )
    style_header_row(ws, SOT_OFFICIAL, row=HEADER_ROW, underline_photo_link=True)
    apply_widths(ws, _widths_for(SOT_OFFICIAL))
    _input_runway(ws, SOT_OFFICIAL, HEADER_ROW, RUNWAY)
    _add_table(ws, "Official", SOT_OFFICIAL, HEADER_ROW, RUNWAY)
    add_list_dropdown(
        ws,
        header_index(SOT_OFFICIAL, "status") + 1,
        SOT_OFFICIAL_STATUS,
        start_row=HEADER_ROW + 1,
        prompt="Official status",
    )
    # Freeze header row only (phone-safe).
    ws.freeze_panes = f"A{HEADER_ROW + 1}"


def _build_wishlist(wb: Workbook) -> None:
    ws = wb.create_sheet("Wishlist")
    _title_block(
        ws,
        "Wishlist — finds, not stock",
        "Link/source required if you have a URL. photo_link last / optional. Off Square until bought.",
    )
    style_header_row(ws, SOT_WISHLIST, row=HEADER_ROW, underline_photo_link=True)
    apply_widths(ws, _widths_for(SOT_WISHLIST))
    _input_runway(ws, SOT_WISHLIST, HEADER_ROW, RUNWAY)
    _add_table(ws, "Wishlist", SOT_WISHLIST, HEADER_ROW, RUNWAY)
    add_list_dropdown(
        ws,
        header_index(SOT_WISHLIST, "status") + 1,
        SOT_WISHLIST_STATUS,
        start_row=HEADER_ROW + 1,
        prompt="wishlist status",
    )
    ws.freeze_panes = f"A{HEADER_ROW + 1}"


def _build_orders(wb: Workbook) -> None:
    ws = wb.create_sheet("Orders")
    _title_block(
        ws,
        "Orders — Facebook hello → pay / ship",
        "Inquiry | Reserved | Paid | Shipped | Picked up | Cancelled. Bots draft only.",
    )
    style_header_row(ws, SOT_ORDERS, row=HEADER_ROW, underline_photo_link=False)
    apply_widths(ws, _widths_for(SOT_ORDERS))
    _input_runway(ws, SOT_ORDERS, HEADER_ROW, RUNWAY)
    _add_table(ws, "Orders", SOT_ORDERS, HEADER_ROW, RUNWAY)
    add_list_dropdown(
        ws,
        header_index(SOT_ORDERS, "order_status") + 1,
        ORDER_STATUS,
        start_row=HEADER_ROW + 1,
        prompt="order status",
    )
    ws.freeze_panes = f"A{HEADER_ROW + 1}"


def _build_bot(wb: Workbook) -> None:
    ws = wb.create_sheet("Bot_Activity")
    style_header_row(ws, BOT_ACTIVITY, underline_photo_link=False)
    apply_widths(ws, BOT_ACTIVITY_WIDTHS)
    paint_tab(ws)
    fill = PatternFill("solid", fgColor=BLUSH_ROW)
    for row in range(2, 12):
        for col in range(1, len(BOT_ACTIVITY) + 1):
            cell = ws.cell(row, col, None)
            cell.fill = fill
            cell.font = BODY_FONT
    ws.auto_filter.ref = f"A1:{get_column_letter(len(BOT_ACTIVITY))}11"


def _build_dashboard(wb: Workbook) -> None:
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 28
    ws["A1"] = "🎀 Sassy Closet · Morning Dashboard"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=INK)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUSH)
    ws.merge_cells("A1:B1")
    ws["A2"] = SQUARE_SOT_LINE
    ws["A2"].font = Font(name="Calibri", size=10, color=INK)
    ws.merge_cells("A2:B2")
    ws["A4"] = "TODAY (formulas — open in Excel so values cache)"
    ws["A4"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    ws["A4"].fill = PatternFill("solid", fgColor=SECTION)

    for row, (label, formula) in SOT_DASHBOARD_ROWS.items():
        ws.cell(row, 1, label).font = BODY_FONT
        cell = ws.cell(row, 2, formula)
        cell.font = BODY_FONT
        cell.fill = YELLOW
        if row >= 21:
            cell.number_format = "@"

    ws["A20"] = "NEXT MÃ — Stock reads these. Bots never mint."
    ws["A20"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    ws["A20"].fill = PatternFill("solid", fgColor=SECTION)
    ws.merge_cells("A20:B20")

    ws["A29"] = "COPY FOR BOSS"
    ws["A29"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    ws["A29"].fill = PatternFill("solid", fgColor=SECTION)
    ws.merge_cells("A29:B29")
    ws["A30"] = (
        "Open this book in Microsoft Excel (not chat preview). "
        f"Copy {SOT_DASHBOARD_BRIEF_CELL} — that is the morning brief. "
        "If a script prints a formula, Excel has not cached values yet."
    )
    ws["A30"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A30"].font = BODY_FONT
    ws.merge_cells("A30:B32")
    ws.row_dimensions[30].height = 36
    ws["A34"] = BOTS_DRAFT_ONLY
    ws["A34"].font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    ws.merge_cells("A34:B34")
    ws["A36"] = PHOTO_RULE
    ws["A36"].font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    ws.merge_cells("A36:B36")

    ws["A43"] = "Morning brief (Boss copies this cell → Mini Boss / Slack)"
    ws["A43"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    brief = ws[SOT_DASHBOARD_BRIEF_CELL]
    brief.value = SOT_DASHBOARD_BRIEF_FORMULA
    brief.font = Font(name="Calibri", size=11, bold=True, color=INK)
    brief.fill = PatternFill("solid", fgColor=BLUSH)
    brief.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[43].height = 48
    ws.freeze_panes = "A5"
    paint_tab(ws)


def _covers(wb: Workbook) -> None:
    start = wb.active
    start.title = "START HERE"
    write_how_to_use(
        start,
        [
            ("title", "🎀 Sassy Closet SoT — the ONE desktop book"),
            ("sub", "Working copy / mã index / captions. Square Free is on-hand inventory SoT."),
            ("divider", DIVIDER),
            ("section", "How to use"),
            ("rule", "1. Official = pieces you already own. One row, one Boss-assigned mã. Never invent AO/QU/VA/AK/GI/PK/SET+3."),
            ("rule", "2. Wishlist = finds. Paste Link/source if you have it. photo_link optional / last. Off Square until bought."),
            ("rule", "3. Orders = Facebook hello → Inquiry → Reserved → Paid → Shipped or Picked up (or Cancelled)."),
            ("rule", "4. Dashboard formulas need Excel. Boss copies Dashboard!B43 for the morning brief."),
            ("rule", "5. Bot_Activity is a draft log. Bots never Save, post, or take Zelle."),
            ("rule", f"6. Photos live in {ONEDRIVE_PHOTOS}/ as AO001.jpg / #001.jpg. Excel stores photo_link only."),
            ("divider", DIVIDER),
            ("section", "Shop law"),
            ("body", SQUARE_SOT_LINE),
            ("body", BOTS_DRAFT_ONLY),
            ("body", "Cursor Cloud Agents change kit code on GitHub. Kit syncs this .xlsx to OneDrive."),
            ("divider", DIVIDER),
            ("foot", "♡ Empty data tables are intentional. No sample stock."),
        ],
    )

    tomorrow = wb.create_sheet("TOMORROW")
    write_how_to_use(
        tomorrow,
        [
            ("title", "Tomorrow — soft daily list"),
            ("sub", "Not inventory. Check Square first if a count disagrees."),
            ("section", "Morning"),
            ("body", "1. Open Dashboard in Excel. Copy B43. Send to Mini Boss if she asks."),
            ("body", "2. Filter Official Available. Confirm Square on-hand for anything reserved."),
            ("body", "3. Wishlist: any bought row still missing a mã? Ask Stock — do not invent."),
            ("section", "Inbox"),
            ("body", "4. New Facebook hello → Orders Inquiry (append_order_row.py)."),
            ("body", "5. Paid? Owner Zelle + Square Save. Script only logs the working copy."),
            ("section", "Night"),
            ("body", "6. Kit / Mini Boss: if Agents changed excel-kit, rebuild or append, then sync OneDrive."),
            ("foot", SQUARE_SOT_SHORT),
        ],
    )

    prompt = wb.create_sheet("BOT PROMPT")
    write_how_to_use(
        prompt,
        [
            ("title", "BOT PROMPT — paste into a Cloud Agent"),
            ("sub", "Agents edit GitHub excel-kit. They do not Save in Square."),
            ("section", "Standing"),
            ("body", "You are working on Sassy Closet excel-kit. Square Free = on-hand SoT. Official Excel = working copy."),
            ("body", "Never invent mã. Never embed images. photo_link last on Wishlist / Ma_List / Candidates."),
            ("body", "Do not commit live inventory, customer names, or Square tokens."),
            ("section", "Useful commands"),
            ("body", "python3 excel-kit/sot/append_official_row.py --ma AO001 ..."),
            ("body", "python3 excel-kit/sot/append_wishlist_row.py --source URL ..."),
            ("body", "python3 excel-kit/sot/append_order_row.py --status Inquiry ..."),
            ("body", "python3 excel-kit/sot/dashboard_brief.py"),
            ("body", "python3 excel-kit/tests/run_checks.py"),
            ("foot", "Full prompt pack: excel-kit/PROMPTS.md"),
        ],
    )


def build(out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / Path(ONEDRIVE_SOT).name
    wb = new_boutique_workbook()
    _covers(wb)
    _build_lists(wb)
    _build_official(wb)
    _build_wishlist(wb)
    _build_orders(wb)
    _build_dashboard(wb)
    _build_bot(wb)
    # Keep START HERE first; Excel already has it as active.
    order = [
        "START HERE",
        "TOMORROW",
        "BOT PROMPT",
        "Lists",
        "Official",
        "Wishlist",
        "Orders",
        "Dashboard",
        "Bot_Activity",
    ]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))
    strip_workbook_images(wb)
    wb.save(path)
    wb.close()
    verify_sot_book(path)
    print(f"wrote {path}")
    print("empty Official / Wishlist / Orders on purpose — no live inventory")
    print(f"sync target: OneDrive {ONEDRIVE_SOT}")
    print(f"morning brief: Dashboard!{SOT_DASHBOARD_BRIEF_CELL}")
    return path


def verify_sot_book(path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        for required in ("START HERE", "Official", "Wishlist", "Orders", "Dashboard"):
            if required not in wb.sheetnames:
                raise AssertionError(f"missing sheet {required}")
        official = wb["Official"]
        headers = [official.cell(HEADER_ROW, c).value for c in range(1, len(SOT_OFFICIAL) + 1)]
        if headers != list(SOT_OFFICIAL):
            raise AssertionError(f"Official headers drifted: {headers}")
        wish = wb["Wishlist"]
        wheaders = [wish.cell(HEADER_ROW, c).value for c in range(1, len(SOT_WISHLIST) + 1)]
        if wheaders != list(SOT_WISHLIST):
            raise AssertionError(f"Wishlist headers drifted: {wheaders}")
        if wheaders[-1] != "photo_link":
            raise AssertionError("Wishlist photo_link must be last")
        if wb["Dashboard"][SOT_DASHBOARD_BRIEF_CELL].value != SOT_DASHBOARD_BRIEF_FORMULA:
            raise AssertionError("Dashboard!B43 formula missing")
        # No invented mã in Official data rows.
        for row in range(HEADER_ROW + 1, official.max_row + 1):
            ma = official.cell(row, 1).value
            if ma not in (None, ""):
                raise AssertionError(f"Official must ship empty, found mã {ma!r}")
        for ws in wb.worksheets:
            if getattr(ws, "_images", None):
                raise AssertionError(f"{ws.title}: embedded images forbidden")
    finally:
        wb.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out-dir", type=Path, default=Path("out"), help="folder for Sassy_Closet_SoT.xlsx")
    args = parser.parse_args(argv)
    build(args.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
