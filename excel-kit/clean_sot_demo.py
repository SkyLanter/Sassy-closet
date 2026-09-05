#!/usr/bin/env python3
"""Clear kit-seeded demo rows from a Sassy Closet SoT workbook.

Also patches outdated inventory wording so Square Free stays the on-hand
source of truth, and strips any embedded pictures (photo_link only).

Does not invent mã or stock. Does not write customer PII.

Usage:
  python3 clean_sot_demo.py path/to/Sassy_Closet_SoT.xlsx
  python3 clean_sot_demo.py SoT.xlsx --output cleaned.xlsx
  python3 clean_sot_demo.py SoT.xlsx --in-place
  python3 clean_sot_demo.py SoT.xlsx --wipe-data   # drop every data row
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (
    SOT_DATA_SHEETS,
    SQUARE_SOT_LINE,
    SQUARE_SOT_SHORT,
    SQUARE_WORDING_PATCHES,
    is_valid_ma,
    looks_like_demo_row,
    resolve_header_key,
    strip_workbook_images,
)


def _cell_text(value: object) -> str:
    return "" if value is None else str(value)


def _already_square_sot(text: str) -> bool:
    lower = text.lower()
    return "square free" in lower and (
        "source of truth" in lower or "on-hand" in lower
    )


def patch_square_wording(wb: Workbook) -> list[str]:
    """Replace outdated SoT wording. Skip cells that already state Square Free."""
    hits: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.strip():
                    continue
                if value.startswith("="):
                    continue
                if _already_square_sot(value):
                    continue
                updated = value
                for pattern, replacement in SQUARE_WORDING_PATCHES:
                    updated = pattern.sub(replacement, updated)
                if updated != value:
                    cell.value = updated
                    hits.append(f"{ws.title}!{cell.coordinate}")
    return hits


def _ensure_start_here_square_line(wb: Workbook) -> bool:
    if "START HERE" not in wb.sheetnames:
        return False
    ws = wb["START HERE"]
    for row in ws.iter_rows(values_only=True):
        blob = " ".join(_cell_text(v) for v in row)
        if _already_square_sot(blob):
            return False
    # Park the standing line in A25 if that cell is empty; otherwise append.
    target = ws["A25"]
    if target.value in (None, ""):
        target.value = SQUARE_SOT_SHORT
        return True
    last = ws.max_row + 2
    ws.cell(last, 1, SQUARE_SOT_SHORT)
    return True


def _header_row_for(ws: Worksheet) -> int:
    meta = SOT_DATA_SHEETS.get(ws.title, {})
    configured = meta.get("header_row")
    if isinstance(configured, int):
        return configured
    return 1


def _table_name_for(ws: Worksheet) -> str | None:
    meta = SOT_DATA_SHEETS.get(ws.title, {})
    name = meta.get("table")
    return name if isinstance(name, str) and name else None


def _row_values(ws: Worksheet, row: int, max_col: int) -> list[object]:
    return [ws.cell(row, col).value for col in range(1, max_col + 1)]


def _row_payload(ws: Worksheet, row: int, max_col: int) -> list[tuple[object, str | None]]:
    payload: list[tuple[object, str | None]] = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row, col)
        target = cell.hyperlink.target if cell.hyperlink is not None else None
        payload.append((cell.value, target))
    return payload


def _row_is_empty(values: list[object]) -> bool:
    return not any(v is not None and str(v).strip() != "" for v in values)


def _clear_row(ws: Worksheet, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row, col)
        cell.value = None
        if cell.hyperlink is not None:
            cell.hyperlink = None


def _resize_table(ws: Worksheet, header_row: int, last_data_row: int, max_col: int) -> None:
    table_name = _table_name_for(ws)
    if not table_name or table_name not in ws.tables:
        return
    table: Table = ws.tables[table_name]
    end = max(last_data_row, header_row)
    table.ref = f"A{header_row}:{get_column_letter(max_col)}{end}"


def clear_demo_rows(ws: Worksheet, *, wipe_data: bool) -> tuple[int, int]:
    """Clear demo (or all) data rows. Returns (cleared, kept)."""
    if ws.title not in SOT_DATA_SHEETS:
        return 0, 0
    header_row = _header_row_for(ws)
    max_col = max(ws.max_column or 1, 1)
    cleared = 0
    kept_rows: list[list[tuple[object, str | None]]] = []
    last_row = ws.max_row or header_row
    for row in range(header_row + 1, last_row + 1):
        values = _row_values(ws, row, max_col)
        if _row_is_empty(values):
            continue
        if wipe_data or looks_like_demo_row(values):
            cleared += 1
        else:
            kept_rows.append(_row_payload(ws, row, max_col))
    for row in range(header_row + 1, last_row + 1):
        _clear_row(ws, row, max_col)
    for offset, payload in enumerate(kept_rows, start=1):
        dest_row = header_row + offset
        for col, (value, link) in enumerate(payload, start=1):
            cell = ws.cell(dest_row, col, value)
            if link:
                cell.hyperlink = link
    kept = len(kept_rows)
    # One blank input row when the sheet is empty so the Excel table still has a body.
    empty_tail = header_row + (kept if kept else 1)
    _resize_table(ws, header_row, empty_tail, max_col)
    if ws.auto_filter and ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{empty_tail}"
    return cleared, kept


def report_invalid_ma(ws: Worksheet) -> list[str]:
    """Flag mã-shaped cells that fail MA_RE. Never invent replacements."""
    if ws.title not in SOT_DATA_SHEETS:
        return []
    header_row = _header_row_for(ws)
    max_col = max(ws.max_column or 1, 1)
    headers = [ws.cell(header_row, col).value for col in range(1, max_col + 1)]
    header_names = ["" if h is None else str(h) for h in headers]
    ma_idx = resolve_header_key(header_names, "ma")
    if ma_idx is None:
        return []
    bad: list[str] = []
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        value = ws.cell(row, ma_idx + 1).value
        if value is None or str(value).strip() == "":
            continue
        if not is_valid_ma(value):
            bad.append(f"{ws.title}!{ws.cell(row, ma_idx + 1).coordinate}={value!r}")
    return bad


def clean_workbook(wb: Workbook, *, wipe_data: bool) -> dict[str, object]:
    report: dict[str, object] = {
        "sheets": {},
        "images_removed": strip_workbook_images(wb),
        "wording_cells": [],
        "start_here_square_added": False,
        "invalid_ma": [],
        "square_law": SQUARE_SOT_LINE,
    }
    for title in list(wb.sheetnames):
        if title not in SOT_DATA_SHEETS:
            continue
        cleared, kept = clear_demo_rows(wb[title], wipe_data=wipe_data)
        report["sheets"][title] = {"cleared": cleared, "kept": kept}
        report["invalid_ma"].extend(report_invalid_ma(wb[title]))
    report["wording_cells"] = patch_square_wording(wb)
    report["start_here_square_added"] = _ensure_start_here_square_line(wb)
    # Second pass — drawings can reappear if a table copy kept them.
    report["images_removed"] = int(report["images_removed"]) + strip_workbook_images(wb)
    return report


def _print_report(path: Path, out: Path, report: dict[str, object]) -> None:
    print(f"cleaned {path} → {out}")
    sheets = report["sheets"]
    if isinstance(sheets, dict):
        if not sheets:
            print("  data sheets: none found (headers unchanged)")
        for name, stats in sheets.items():
            print(f"  {name}: cleared {stats['cleared']} demo/data row(s), kept {stats['kept']}")
    print(f"  embedded images removed: {report['images_removed']}")
    wording = report["wording_cells"]
    if isinstance(wording, list) and wording:
        print(f"  Square wording patched: {', '.join(wording)}")
    else:
        print("  Square wording: already current or no outdated phrases")
    if report["start_here_square_added"]:
        print("  START HERE: added Square Free on-hand SoT line")
    invalid = report.get("invalid_ma") or []
    if isinstance(invalid, list) and invalid:
        print(f"  invalid mã (ask Stock, do not invent): {', '.join(invalid)}")
    print(f"  shop law: {SQUARE_SOT_SHORT}")
    print("  Official Excel = working copy. Square Free = on-hand SoT. No embeds.")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path, help="SoT or desktop .xlsx to clean")
    parser.add_argument("--output", "-o", type=Path, help="write here (default: *_cleaned.xlsx)")
    parser.add_argument("--in-place", action="store_true", help="overwrite the input file")
    parser.add_argument(
        "--wipe-data",
        action="store_true",
        help="clear every Official/Wishlist/Orders/Ma_List/Candidates data row",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="report only — do not write a cleaned workbook",
    )
    args = parser.parse_args(argv)

    src: Path = args.workbook
    if not src.is_file():
        print(f"error: not a file: {src}", file=sys.stderr)
        return 2
    if args.in_place and args.output:
        print("error: use either --in-place or --output, not both", file=sys.stderr)
        return 2
    if args.check and (args.in_place or args.output):
        print("error: --check cannot be combined with --in-place or --output", file=sys.stderr)
        return 2

    wb = load_workbook(src)
    report = clean_workbook(wb, wipe_data=args.wipe_data)
    if args.check:
        _print_report(src, src, report)
        print("check only — workbook not written")
        wb.close()
        return 0
    dest = src if args.in_place else (args.output or src.with_name(f"{src.stem}_cleaned{src.suffix}"))
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    _print_report(src, dest, report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
