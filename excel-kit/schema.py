"""Sassy Closet Excel kit — shared headers, mã rules, and workbook helpers.

Desktop Official / Wishlist use MA_LIST, ORDERS, BOT_ACTIVITY, CANDIDATES.
The OneDrive SoT workbook (Sassy_Closet_SoT.xlsx) is the ONE desktop working
copy: Official / Wishlist / Orders / Dashboard (plus Bot_Activity when present).
Square Free remains on-hand inventory SoT; Official Excel is a working copy /
mã index / captions — not a second stock. Never embed images. photo_link last
on Ma_List, Candidates, and SoT Wishlist.
"""

from __future__ import annotations

import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Literal, Never, Sequence

# openpyxl is required to *write* Excel. Constant / mã helpers used by
# append_official_row --help must import without it (local bugcheck).
try:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.numbers import FORMAT_TEXT
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as exc:  # pragma: no cover — --help / smoke imports
    _OPENPYXL_IMPORT_ERROR: ImportError | None = exc
    Alignment = Border = Font = PatternFill = Side = None  # type: ignore[misc, assignment]
    FORMAT_TEXT = "@"
    get_column_letter = None  # type: ignore[assignment]
    Workbook = None  # type: ignore[misc, assignment]
    DataValidation = None  # type: ignore[misc, assignment]
    Worksheet = None  # type: ignore[misc, assignment]
else:
    _OPENPYXL_IMPORT_ERROR = None

# ---------------------------------------------------------------------------
# Shop paths (OneDrive). Kit syncs generated files here; this repo is code.
# ---------------------------------------------------------------------------

ONEDRIVE_SHOP_DIR = "Documents/Sassy Closet"
ONEDRIVE_SOT = f"{ONEDRIVE_SHOP_DIR}/Sassy_Closet_SoT.xlsx"
ONEDRIVE_PHOTOS = f"{ONEDRIVE_SHOP_DIR}/Photos"
ONEDRIVE_FROM_GF = f"{ONEDRIVE_SHOP_DIR}/From GF"
ONEDRIVE_OFFICIAL_DESKTOP = f"{ONEDRIVE_SHOP_DIR}/Sassy_Closet_Official_desktop.xlsx"
ONEDRIVE_WISHLIST_DESKTOP = f"{ONEDRIVE_SHOP_DIR}/Sassy_Closet_Wishlist_desktop.xlsx"
ONEDRIVE_SQUARE = f"{ONEDRIVE_SHOP_DIR}/Square.xlsx"
ONEDRIVE_FINANCE = f"{ONEDRIVE_SHOP_DIR}/Finance.xlsx"
SQUARE_XLSX_NAME = "Square.xlsx"
FINANCE_XLSX_NAME = "Finance.xlsx"
SQUARE_PHOTO_FOLDER_PREFIX = f"{ONEDRIVE_PHOTOS}/"

# ---------------------------------------------------------------------------
# Mã
# ---------------------------------------------------------------------------

MA_PREFIXES: tuple[str, ...] = ("AO", "QU", "VA", "AK", "GI", "PK", "SET")
MA_PREFIX_MEANS: dict[str, str] = {
    "AO": "Áo / tops",
    "QU": "Quần / bottoms",
    "VA": "Váy / dresses",
    "AK": "Áo khoác / jackets",
    "GI": "Giày / shoes",
    "PK": "Phụ kiện / accessories",
    "SET": "Set",
}
MA_RE = re.compile(r"^(AO|QU|VA|AK|GI|PK|SET)(\d{3})$")

# ---------------------------------------------------------------------------
# Desktop kit headers — photo_link is last on Ma_List and Candidates
# ---------------------------------------------------------------------------

MA_LIST: tuple[str, ...] = (
    "ma",
    "item",
    "sizes_in_stock",
    "color",
    "cost",
    "price",
    "on_hand",
    "status",
    "notes",
    "photo_link",
)

ORDERS: tuple[str, ...] = (
    "date",
    "buyer",
    "ma",
    "size",
    "color",
    "channel",
    "status",
    "pay",
    "ship_or_local",
    "agreed_price",
    "notes",
)

BOT_ACTIVITY: tuple[str, ...] = (
    "Date",
    "Time PT",
    "Bot",
    "Ma",
    "Action",
    "Summary",
    "Status",
)

CANDIDATES: tuple[str, ...] = (
    "#",
    "Date",
    "Status",
    "Type",
    "Brand + name",
    "Size Asia + cm",
    "Color",
    "Cost ¥",
    "Sell $",
    "Link (reopen)",
    "Photos / notes",
    "photo_link",
)

MA_LIST_STATUS: tuple[str, ...] = ("in_stock", "held", "sold", "archived")
ORDER_STATUS: tuple[str, ...] = (
    "Inquiry",
    "Reserved",
    "Paid",
    "Shipped",
    "Picked up",
    "Cancelled",
)
CANDIDATE_STATUS: tuple[str, ...] = ("candidate", "watching", "skip", "bought")
CANDIDATE_TYPES: tuple[str, ...] = (
    "Dress",
    "Top",
    "Skirt",
    "Bag",
    "Shoes",
    "Other",
    "SET",
)

# ---------------------------------------------------------------------------
# Square.xlsx + Finance.xlsx (Boss 2026-09-07 ~11:30 PT; SJ tax 2026-09-18)
# Bought / on-hand team tracker + tax-ready ledger. Staged site mãs are NOT
# bought. Square Free remains on-hand SoT. Empty data is correct. No cute.
# Hub / site sync must not drop these FINANCE_* / SQUARE_* book exports.
# ---------------------------------------------------------------------------

SQUARE_TEMPLATE_ROWS = 20
SQUARE_FORMULA_LAST_ROW = 1001
FINANCE_TAX_YEAR = 2026
FINANCE_ZELLE_DISPLAY_NAME = "Thang Tien Huynh"

# Combined district rate published for San Jose, CA (CDTFA), effective Apr 1, 2026.
# Do not treat the older 9.375% figure as current.
SAN_JOSE_JURISDICTION = "San Jose, California"
SAN_JOSE_SALES_TAX_RATE = Decimal("0.10")
SAN_JOSE_SALES_TAX_RATE_DISPLAY = "10.000%"
SAN_JOSE_SALES_TAX_EFFECTIVE = "2026-04-01"
SAN_JOSE_SALES_TAX_SOURCE = "CDTFA combined district rate"
LEGACY_SAN_JOSE_SALES_TAX_RATE_NOT_CURRENT = Decimal("0.09375")

SQUARE_ON_HAND: tuple[str, ...] = (
    "ma",
    "kind",
    "colors",
    "size",
    "qty_on_hand",
    "cost_cny",
    "cost_usd",
    "cost_currency",
    "buy_date",
    "source_link",
    "photo_folder",
    "square_item_name",
    "track_on",
    "status",
    "sold_date",
    "notes",
)
SQUARE_ON_HAND_STATUS: tuple[str, ...] = ("on_hand", "reserved", "sold", "dead")
SQUARE_TRACK_ON: tuple[str, ...] = ("Y", "N")
SQUARE_COST_CURRENCY: tuple[str, ...] = ("CNY", "USD")

# Contract Sold_Log — finance_ref points at the Finance.xlsx Sales row / pay_ref.
SQUARE_SOLD_LOG: tuple[str, ...] = (
    "ma",
    "sold_date",
    "qty",
    "finance_ref",
    "notes",
)
SQUARE_SHEETS: tuple[str, ...] = ("On_Hand", "Sold_Log", "Readme")
SQUARE_XLSX_SHEETS = SQUARE_SHEETS

SQUARE_README_LINES: tuple[str, ...] = (
    "Square.xlsx tracks bought / on-hand pieces for the team. On_Hand starts empty — staged site mãs are not bought.",
    "Square Free Dashboard is on-hand source of truth. This Excel is the team tracker, not a second warehouse.",
    "When a mã sells: qty_on_hand 0, status sold, sold_date; Sold_Log (ma, sold_date, qty, finance_ref); Finance.xlsx Sales row (required).",
    "photo_folder = Documents/Sassy Closet/Photos/{ma}/. Never invent mã, stock, or $.",
    "Never Square Save. No cute / embeds. Rebuild leaves On_Hand empty unless Cap seeds a confirmed buy.",
)
SQUARE_XLSX_README_LINES = SQUARE_README_LINES

# Contract Sales columns stay in A–N. Tax / COGS helpers append before notes so
# E/F/G/H remain gross / ship / discount / net (formula-stable).
FINANCE_SALES: tuple[str, ...] = (
    "date",
    "ma",
    "description",
    "qty",
    "gross_usd",
    "ship_usd",
    "discount_usd",
    "net_usd",
    "pay_method",
    "pay_ref",
    "customer_note",
    "channel",
    "square_xlsx_ma",
    "tax_category",
    "cogs_usd",
    "taxable_base_usd",
    "sales_tax_rate",
    "sales_tax_usd",
    "notes",
)
FINANCE_PAY_METHOD: tuple[str, ...] = ("zelle", "square", "square_online", "cash", "other")
FINANCE_CHANNEL: tuple[str, ...] = ("facebook", "meetup", "website", "other")
FINANCE_SALES_CHANNEL = FINANCE_CHANNEL
FINANCE_TAX_CATEGORY: tuple[str, ...] = ("product_sale", "shipping", "other")
FINANCE_SALES_TAX_CATEGORY = FINANCE_TAX_CATEGORY

FINANCE_FEES: tuple[str, ...] = (
    "date",
    "source",
    "amount_usd",
    "fee_type",
    "related_sale_ref",
    "notes",
)
FINANCE_FEE_SOURCE: tuple[str, ...] = ("square", "square_online", "bank", "other")
FINANCE_FEE_TYPE: tuple[str, ...] = ("square_processing", "shipping_label", "ads", "other")

FINANCE_PAYOUTS: tuple[str, ...] = (
    "date",
    "from_method",
    "to_account_note",
    "amount_usd",
    "confirmation",
    "notes",
)
FINANCE_FROM_METHOD: tuple[str, ...] = (
    "zelle",
    "square",
    "square_online",
    "cash",
    "bank",
    "other",
)
FINANCE_PAYOUT_KIND: tuple[str, ...] = ("square_payout", "zelle", "bank", "other")

FINANCE_EXPENSES: tuple[str, ...] = (
    "date",
    "vendor",
    "category",
    "amount_usd",
    "payment_method",
    "receipt_note",
    "notes",
)
FINANCE_EXPENSE_CATEGORY: tuple[str, ...] = (
    "inventory_cogs",
    "shipping_supplies",
    "packaging",
    "software",
    "ads",
    "other",
)
FINANCE_EXPENSE_PAY: tuple[str, ...] = ("zelle", "square", "square_online", "cash", "other")

FINANCE_SHEETS: tuple[str, ...] = (
    "Sales",
    "Fees",
    "Payouts_Transfers",
    "Expenses",
    "Tax_Summary",
    "Readme",
)

FinanceMetric = Literal[
    "gross",
    "shipping",
    "discount",
    "net",
    "taxable_base",
    "sales_tax",
    "fees",
    "cogs",
    "expenses",
    "profit",
]

FINANCE_TAX_SUMMARY_METRICS: tuple[tuple[str, FinanceMetric], ...] = (
    ("Gross sales (USD)", "gross"),
    ("Shipping income (USD)", "shipping"),
    ("Discounts (USD)", "discount"),
    ("Net receipts (gross + ship − discount, no sales tax)", "net"),
    ("Taxable base (USD)", "taxable_base"),
    ("Sales tax collected (USD, liability — not income)", "sales_tax"),
    ("Fees (USD)", "fees"),
    ("COGS (Sales cogs_usd — copy Square.xlsx cost_usd when sold)", "cogs"),
    ("Expenses (USD)", "expenses"),
    ("Net profit (income-tax base — not legal advice)", "profit"),
)

FINANCE_TAX_SUMMARY_HEADERS: tuple[str, ...] = (
    "metric",
    *(f"{FINANCE_TAX_YEAR}-{month:02d}" for month in range(1, 13)),
    "YTD",
    "accountant_note",
)

FINANCE_TAX_SUMMARY_NOTE_ROWS: tuple[tuple[str, str], ...] = (
    ("Shop jurisdiction", SAN_JOSE_JURISDICTION),
    (
        "Combined sales tax rate",
        f"{SAN_JOSE_SALES_TAX_RATE} ({SAN_JOSE_SALES_TAX_RATE_DISPLAY} {SAN_JOSE_SALES_TAX_SOURCE}, effective {SAN_JOSE_SALES_TAX_EFFECTIVE}). Do not use 9.375% as current.",
    ),
    (
        "Money model",
        "Sales tax is on the taxable sell price to the customer and is collected separately. Income tax applies to profit, not COGS. Columns are Schedule C-style summary only — not legal advice.",
    ),
    (
        "Customer shipping",
        "Customer flat ship $ is TBD. Never invent a ship rate into official numbers. Enter ship_usd only when the customer actually paid shipping.",
    ),
    (
        "Zelle display name (customers see)",
        f"{FINANCE_ZELLE_DISPLAY_NAME} — note only, never store bank passwords",
    ),
    (
        "Sold path (finance_ref)",
        "Square.xlsx On_Hand status=sold, qty_on_hand 0, sold_date → Sold_Log.finance_ref → Finance.xlsx Sales.square_xlsx_ma (required). Dead write-off = no Sales row.",
    ),
    (
        "Empty book",
        "Sales / Fees / Payouts_Transfers / Expenses start empty. Rebuild does not invent sales or $. Cap enters real rows only.",
    ),
    (
        "Payouts / transfers / 1099-K",
        "Payouts_Transfers are not income. Paste 1099-K / Square export notes on that sheet. Give Tax_Summary to an accountant.",
    ),
    (
        "Open questions (Cap / accountant)",
        "CDTFA seller permit, nexus, and tax-included vs +tax pricing are not decided in this book. Default formulas treat sales tax as extra (not stuffed into profit).",
    ),
)

# YTD formulas. Letters must match FINANCE_* header order (tested).
FINANCE_TAX_SUMMARY_ROWS: tuple[tuple[str, str], ...] = (
    ("metric", "YTD"),
    ("Gross sales (USD)", f"=SUM(Sales!E2:E{SQUARE_FORMULA_LAST_ROW})"),
    ("Shipping income (USD)", f"=SUM(Sales!F2:F{SQUARE_FORMULA_LAST_ROW})"),
    ("Discounts (USD)", f"=SUM(Sales!G2:G{SQUARE_FORMULA_LAST_ROW})"),
    ("Net receipts (gross + ship − discount, no sales tax)", f"=SUM(Sales!H2:H{SQUARE_FORMULA_LAST_ROW})"),
    ("Taxable base (USD)", f"=SUM(Sales!P2:P{SQUARE_FORMULA_LAST_ROW})"),
    ("Sales tax collected (USD, liability — not income)", f"=SUM(Sales!R2:R{SQUARE_FORMULA_LAST_ROW})"),
    ("Fees (USD)", f"=SUM(Fees!C2:C{SQUARE_FORMULA_LAST_ROW})"),
    ("COGS (Sales cogs_usd — copy Square.xlsx cost_usd when sold)", f"=SUM(Sales!O2:O{SQUARE_FORMULA_LAST_ROW})"),
    ("Expenses (USD)", f"=SUM(Expenses!D2:D{SQUARE_FORMULA_LAST_ROW})"),
    (
        "Net profit (income-tax base — not legal advice)",
        (
            f"=SUM(Sales!H2:H{SQUARE_FORMULA_LAST_ROW})"
            f"-SUM(Fees!C2:C{SQUARE_FORMULA_LAST_ROW})"
            f"-SUM(Sales!O2:O{SQUARE_FORMULA_LAST_ROW})"
            f"-SUM(Expenses!D2:D{SQUARE_FORMULA_LAST_ROW})"
        ),
    ),
)

FINANCE_README_LINES: tuple[str, ...] = (
    "Finance.xlsx is tax-ready for San Jose, California. Sales / Fees / Payouts_Transfers / Expenses start empty — no invented sales or $.",
    "Sold in Square.xlsx → one Sales row (square_xlsx_ma + Sold_Log.finance_ref). Copy Square cost_usd into cogs_usd. Never invent sell prices; formulas compute from entered rows.",
    "San Jose combined sales tax is 10.000% (CDTFA, effective Apr 1, 2026). Tax is on taxable sell price to the customer (sales_tax_usd), collected separately — do not stuff tax into profit. Income tax is on profit, not COGS. Do not use 9.375% as current.",
    "Example shape only (Cap types real $): cost 100 + profit 20 → sell 120 before tax; SJ sales tax ≈ 12 on 120; income-tax base ≈ 20. Customer flat ship $ is TBD — never invent a ship rate. 35% GM is a Cap playbook floor, not a formula in this book.",
    "pay_method zelle|square|square_online|cash|other. channel facebook|meetup|website|other. tax_category product_sale|shipping|other. fee_type square_processing|shipping_label|ads|other. expense category inventory_cogs|shipping_supplies|packaging|software|ads|other.",
    "Zelle display customers see: Thang Tien Huynh (note only — never store bank passwords). Payouts / transfers are not income. Tax_Summary is monthly + YTD Schedule C-style — not legal advice. Never Square Save. No cute / embeds.",
    "Land path: Documents/Sassy Closet/Finance.xlsx beside Square.xlsx + sassycloset.xlsx. Open Tax_Summary in Excel so formulas calculate. Empty totals of 0 are correct until real sales.",
)

# Symbols build_square_finance.py imports. Deleting any of these is a regression.
FINANCE_BOOK_EXPORTS: tuple[str, ...] = (
    "FINANCE_CHANNEL",
    "FINANCE_EXPENSE_CATEGORY",
    "FINANCE_EXPENSE_PAY",
    "FINANCE_EXPENSES",
    "FINANCE_FEE_SOURCE",
    "FINANCE_FEE_TYPE",
    "FINANCE_FEES",
    "FINANCE_FROM_METHOD",
    "FINANCE_PAY_METHOD",
    "FINANCE_PAYOUT_KIND",
    "FINANCE_PAYOUTS",
    "FINANCE_README_LINES",
    "FINANCE_SALES",
    "FINANCE_SALES_CHANNEL",
    "FINANCE_SALES_TAX_CATEGORY",
    "FINANCE_SHEETS",
    "FINANCE_TAX_CATEGORY",
    "FINANCE_TAX_SUMMARY_HEADERS",
    "FINANCE_TAX_SUMMARY_METRICS",
    "FINANCE_TAX_SUMMARY_NOTE_ROWS",
    "FINANCE_TAX_SUMMARY_ROWS",
    "FINANCE_TAX_YEAR",
    "FINANCE_XLSX_NAME",
    "FINANCE_ZELLE_DISPLAY_NAME",
    "ONEDRIVE_FINANCE",
    "ONEDRIVE_PHOTOS",
    "ONEDRIVE_SQUARE",
    "SAN_JOSE_JURISDICTION",
    "SAN_JOSE_SALES_TAX_RATE",
    "SAN_JOSE_SALES_TAX_RATE_DISPLAY",
    "SQUARE_COST_CURRENCY",
    "SQUARE_FORMULA_LAST_ROW",
    "SQUARE_ON_HAND",
    "SQUARE_ON_HAND_STATUS",
    "SQUARE_README_LINES",
    "SQUARE_SHEETS",
    "SQUARE_SOLD_LOG",
    "SQUARE_TEMPLATE_ROWS",
    "SQUARE_TRACK_ON",
    "SQUARE_XLSX_NAME",
    "add_list_dropdown",
    "finance_net_usd_formula",
    "header_index",
    "san_jose_sales_tax_usd",
    "square_photo_folder_formula",
    "strip_workbook_images",
)

# ---------------------------------------------------------------------------
# SoT workbook (pink stock book) — used by clean_sot_demo.py
# ---------------------------------------------------------------------------

SOT_SHEETS: tuple[str, ...] = (
    "START HERE",
    "TOMORROW",
    "BOT PROMPT",
    "Lists",
    "Official",
    "Wishlist",
    "Orders",
    "Dashboard",
    "Bot_Activity",
)

# Optional on the SoT book; required on the lean Official desktop book.
SOT_OPTIONAL_SHEETS: tuple[str, ...] = ("Bot_Activity",)

SHEET_ALIASES: dict[str, tuple[str, ...]] = {
    "official": ("Official", "Ma_List", "Ma List"),
    "wishlist": ("Wishlist", "Candidates"),
    "orders": ("Orders",),
    "bot_activity": ("Bot_Activity", "Bot Activity", "BOT_ACTIVITY"),
    "dashboard": ("Dashboard",),
    "lists": ("Lists",),
    "start_here": ("START HERE", "How_to_use"),
}

SOT_OFFICIAL: tuple[str, ...] = (
    "ma",
    "date_in",
    "brand",
    "name_vi",
    "name_en",
    "category",
    "size",
    "color",
    "condition",
    "qty_on_hand",
    "cost",
    "list_price",
    "margin_$",
    "margin_pct",
    "status",
    "photo_file",
    "photo_link",
    "fb_url",
    "storage",
    "source",
    "bust_chest",
    "waist",
    "length",
    "flaws",
    "reserved_for",
    "days_in",
    "square_name",
    "caption_ready",
    "notes",
    "last_updated",
    "ma_num",
    "prefix",
)

# photo_link is last. source = Link/source (Taobao / shop URL) — required if available.
SOT_WISHLIST: tuple[str, ...] = (
    "wish_id",
    "date_added",
    "prefix",
    "what_vi",
    "what_en",
    "size",
    "color",
    "max_cost",
    "target_price",
    "reason",
    "requested_by",
    "status",
    "source",
    "notes",
    "photo_link",
)

SOT_ORDERS: tuple[str, ...] = (
    "order_id",
    "date",
    "channel",
    "buyer",
    "fb_handle",
    "ma",
    "snapshot",
    "qty",
    "agreed_price",
    "ship_fee",
    "net",
    "order_status",
    "pay_method",
    "fulfill",
    "meetup_or_ship",
    "thread_note",
    "last_touch",
)

SOT_OFFICIAL_STATUS: tuple[str, ...] = (
    "Available",
    "Reserved",
    "Sold",
    "Hold",
    "Damaged",
    "Donated",
)

SOT_WISHLIST_STATUS: tuple[str, ...] = ("candidate", "watching", "skip", "bought")
SOT_ORDER_STATUS: tuple[str, ...] = ORDER_STATUS
SOT_CONDITION: tuple[str, ...] = ("New", "Like new", "Good", "Fair", "As-is")
SOT_CHANNELS: tuple[str, ...] = ("Facebook", "IG", "Walk-in", "Friend", "Other")
SOT_PAY: tuple[str, ...] = ("Zelle", "Cash", "Venmo", "Other", "Unpaid")
SOT_FULFILL: tuple[str, ...] = ("Ship", "Local pickup", "Hold", "TBD")

# Official working-copy status → lean Ma_List status (Square still wins on-hand).
OFFICIAL_TO_MA_LIST_STATUS: dict[str, str] = {
    "Available": "in_stock",
    "Reserved": "held",
    "Sold": "sold",
    "Hold": "held",
    "Damaged": "archived",
    "Donated": "archived",
}

# Header aliases so append scripts work on both SoT and lean desktop books.
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "ma": ("ma", "mã", "Ma"),
    "item": ("item", "name_vi", "Brand + name", "what_vi"),
    "name_vi": ("name_vi", "what_vi", "item", "Brand + name"),
    "name_en": ("name_en", "what_en"),
    "what_vi": ("what_vi", "name_vi", "Brand + name", "item"),
    "what_en": ("what_en", "name_en"),
    "size": ("size", "sizes_in_stock", "Size Asia + cm"),
    "sizes_in_stock": ("sizes_in_stock", "size", "Size Asia + cm"),
    "color": ("color", "Color", "mau", "Mau"),
    "cost": ("cost", "Cost ¥", "max_cost", "Von Y"),
    "list_price": ("list_price", "price", "Sell $", "Gia ban $", "target_price", "agreed_price"),
    "price": ("price", "list_price", "Sell $", "Gia ban $"),
    "qty_on_hand": ("qty_on_hand", "on_hand", "qty"),
    "on_hand": ("on_hand", "qty_on_hand", "qty"),
    "status": ("status", "Status", "order_status"),
    "order_status": ("order_status", "status", "Status"),
    "notes": ("notes", "Photos / notes", "thread_note"),
    "photo_link": ("photo_link",),
    "source": (
        "source",
        "link",
        "Link (reopen)",
        "Link/source",
        "link_source",
        "Link",
    ),
    "date": ("date", "Date", "date_in", "date_added"),
    "date_in": ("date_in", "date", "Date"),
    "date_added": ("date_added", "date", "Date"),
    "buyer": ("buyer",),
    "channel": ("channel",),
    "pay": ("pay", "pay_method"),
    "pay_method": ("pay_method", "pay"),
    "ship_or_local": ("ship_or_local", "fulfill", "meetup_or_ship"),
    "agreed_price": ("agreed_price", "price", "list_price"),
    "wish_id": ("wish_id", "#"),
    "order_id": ("order_id",),
    "brand": ("brand", "Brand + name"),
    "prefix": ("prefix",),
    "requested_by": ("requested_by",),
    "reason": ("reason",),
    "max_cost": ("max_cost", "cost", "Cost ¥"),
    "target_price": ("target_price", "Sell $", "list_price", "price"),
    "type": ("Type", "type", "category"),
    "category": ("category", "Type", "type"),
    "photo_file": ("photo_file",),
    "square_name": ("square_name",),
    "caption_ready": ("caption_ready",),
    "fb_handle": ("fb_handle",),
    "snapshot": ("snapshot",),
    "qty": ("qty", "qty_on_hand", "on_hand"),
    "ship_fee": ("ship_fee",),
    "net": ("net",),
    "fulfill": ("fulfill", "ship_or_local"),
    "meetup_or_ship": ("meetup_or_ship", "ship_or_local"),
    "thread_note": ("thread_note", "notes"),
    "last_touch": ("last_touch", "last_updated"),
    "last_updated": ("last_updated", "last_touch"),
    "reserved_for": ("reserved_for",),
    "condition": ("condition",),
    "storage": ("storage",),
    "fb_url": ("fb_url",),
}

# Sheets whose data rows may hold kit-seeded demo stock / fake buyers.
SOT_DATA_SHEETS: dict[str, dict[str, object]] = {
    "Official": {"header_row": 4, "table": "Official"},
    "Wishlist": {"header_row": 4, "table": "Wishlist"},
    "Orders": {"header_row": 4, "table": "Orders"},
    "Bot_Activity": {"header_row": 1, "table": None},
    "Ma_List": {"header_row": 1, "table": None},
    "Candidates": {"header_row": 1, "table": None},
    "Ma List": {"header_row": 1, "table": None},
    "Bot Activity": {"header_row": 1, "table": None},
}

# ---------------------------------------------------------------------------
# Dashboard — morning brief. Boss copies Dashboard!B43 when formulas are live.
# ---------------------------------------------------------------------------

SOT_DASHBOARD_BRIEF_CELL = "B43"
SOT_DASHBOARD_BRIEF_LABEL_CELL = "A43"

# row -> (A-label, B-formula-or-kind). Formulas use Excel table names.
SOT_DASHBOARD_ROWS: dict[int, tuple[str, str]] = {
    5: ("Official pieces (ma rows)", '=COUNTA(Official[ma])'),
    6: ("Available", '=COUNTIF(Official[status],"Available")'),
    7: ("Reserved (Official)", '=COUNTIF(Official[status],"Reserved")'),
    8: ("Sold", '=COUNTIF(Official[status],"Sold")'),
    9: (
        "Other Official (Hold / Damaged / Donated)",
        '=COUNTIF(Official[status],"Hold")+COUNTIF(Official[status],"Damaged")+COUNTIF(Official[status],"Donated")',
    ),
    10: (
        "Qty on hand (Available rows — working copy)",
        '=SUMIF(Official[status],"Available",Official[qty_on_hand])',
    ),
    11: (
        "Wishlist open (candidate + watching)",
        '=COUNTIF(Wishlist[status],"candidate")+COUNTIF(Wishlist[status],"watching")',
    ),
    12: ("Wishlist bought (not Square until Boss Save)", '=COUNTIF(Wishlist[status],"bought")'),
    13: ("Orders Inquiry", '=COUNTIF(Orders[order_status],"Inquiry")'),
    14: ("Orders Reserved", '=COUNTIF(Orders[order_status],"Reserved")'),
    15: ("Orders Paid", '=COUNTIF(Orders[order_status],"Paid")'),
    16: ("Orders Shipped", '=COUNTIF(Orders[order_status],"Shipped")'),
    17: ("Orders Picked up", '=COUNTIF(Orders[order_status],"Picked up")'),
    18: ("Orders Cancelled", '=COUNTIF(Orders[order_status],"Cancelled")'),
    21: ("Next AO (Stock reads — bots do not mint)", '="AO"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"AO")+1,"000")'),
    22: ("Next QU (Stock reads — bots do not mint)", '="QU"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"QU")+1,"000")'),
    23: ("Next VA (Stock reads — bots do not mint)", '="VA"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"VA")+1,"000")'),
    24: ("Next AK (Stock reads — bots do not mint)", '="AK"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"AK")+1,"000")'),
    25: ("Next GI (Stock reads — bots do not mint)", '="GI"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"GI")+1,"000")'),
    26: ("Next PK (Stock reads — bots do not mint)", '="PK"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"PK")+1,"000")'),
    27: ("Next SET (Stock reads — bots do not mint)", '="SET"&TEXT(MAXIFS(Official[ma_num],Official[prefix],"SET")+1,"000")'),
}

SOT_DASHBOARD_BRIEF_FORMULA = (
    '="Morning brief · Official "&TEXT(B5,"0")'
    '&" · Available "&TEXT(B6,"0")'
    '&" · Reserved "&TEXT(B7,"0")'
    '&" · on-hand "&TEXT(B10,"0")'
    '&" · wishlist open "&TEXT(B11,"0")'
    '&" · inquiry "&TEXT(B13,"0")'
    '&" · reserved orders "&TEXT(B14,"0")'
    '&" · paid "&TEXT(B15,"0")'
    '&" · next AO "&B21'
)

SOT_DASHBOARD_BRIEF_NOTE = (
    "If this script prints a formula instead of numbers, Excel has not cached "
    "values yet. Boss: open Sassy_Closet_SoT.xlsx → Dashboard → copy B43 "
    "(morning brief). Paste that text to Mini Boss / Slack. "
    "openpyxl cannot calculate Excel formulas."
)

# ---------------------------------------------------------------------------
# Square Free catalog import (headers only in git — never live stock rows)
# ---------------------------------------------------------------------------

# Stable Square item-library columns Stock maps in Dashboard → Import library.
# Location-specific headers use [Default]; Stock replaces Default with the
# live Square location name from a fresh export before any Boss-approved Save.
SQUARE_IMPORT_HEADERS: tuple[str, ...] = (
    "Token",
    "Item Name",
    "Variation Name",
    "SKU",
    "Description",
    "Category",
    "Reporting Category",
    "SEO Title",
    "SEO Description",
    "Permalink",
    "GTIN",
    "Square Online Item Visibility",
    "Item Type",
    "Weight",
    "Price",
    "Online Sale Price",
    "Archived",
    "Sellable",
    "Stockable",
    "Skip Detail Screen in POS",
    "Option Name 1",
    "Option Value 1",
    "Default Unit Cost",
    "Enabled [Default]",
    "Current Quantity [Default]",
    "New Quantity [Default]",
    "Stock Alert Enabled [Default]",
    "Stock Alert Count [Default]",
)

SQUARE_SKU_IS_MA = (
    "SKU = mã (AO001). Item Name = mã + short name. "
    "Variation Name = Asia size + color (M / đen). "
    "Stockable = Y (Track stock ON). Never put No in New Quantity to kill tracking."
)

# ---------------------------------------------------------------------------
# Standing copy
# ---------------------------------------------------------------------------

SQUARE_SOT_LINE = (
    "Square Free = on-hand inventory source of truth. "
    "Official Excel (Sassy_Closet_SoT.xlsx) is a working copy / mã index / captions — "
    "NOT a second inventory."
)

SQUARE_SOT_SHORT = (
    "Square Free = on-hand source of truth. "
    "Official Excel = working copy / mã index / captions / photos — "
    "not a second inventory brain. Wishlist ≠ stock. Bots draft only."
)

ASK_STOCK_MA = (
    "Never invent a mã. Ask Stock to read Dashboard next-mã "
    "(MAXIFS on Official[ma_num] + prefix, cells B21:B27) and Boss to assign it. "
    "Do not reuse a Sold mã."
)

STAY_OFF_SQUARE = (
    "Stay off Square until Boss confirms bought/received AND says Save. "
    "Approvals live in Slack #shop-decisions. Scripts never Square Save. "
    "Wishlist / candidates are not on-hand."
)

BOTS_DRAFT_ONLY = (
    "Bots draft only. Owner posts on Facebook, sends the message, "
    "takes Zelle, and taps Save in Square."
)

PHOTO_RULE = (
    f"Photos live in {ONEDRIVE_PHOTOS}/ named #001.jpg / AO001.jpg. "
    "Excel stores photo_link only — never embed images."
)

# Outdated phrases → current shop law. Applied cell-by-cell; skip cells
# that already state Square Free is on-hand SoT.
SQUARE_WORDING_PATCHES: tuple[tuple[re.Pattern[str], str], ...] = (
    (
        re.compile(
            r"(?i)\bgoogle (?:drive|sheets) is (?:the )?source of truth\b"
        ),
        SQUARE_SOT_LINE,
    ),
    (
        re.compile(
            r"(?i)\b(?:this )?(?:excel|workbook|spreadsheet|sheet) is "
            r"(?:the )?(?:on-hand )?source of truth"
            r"(?:\s+for(?:\s+the)?(?:\s+on-hand)?\s+inventory)?[.]?"
        ),
        SQUARE_SOT_LINE,
    ),
    (
        re.compile(r"(?i)\bexcel is (?:the )?inventory\b"),
        (
            "Square Free is on-hand inventory; "
            "Excel is a working copy / mã index / captions"
        ),
    ),
    (
        re.compile(r"(?i)\blive inventory (?:lives|is) in (?:this )?excel\b"),
        SQUARE_SOT_LINE,
    ),
)

# Tokens that mark kit-seeded demo / example rows (never real stock).
DEMO_ROW_TOKENS: tuple[str, ...] = (
    "demo",
    "sample",
    "example",
    "placeholder",
    "lorem",
    "fake stock",
    "test buyer",
    "demo buyer",
    "ao999",
    "qu999",
    "va999",
    "ak999",
    "gi999",
    "pk999",
    "set999",
    "ao000",
)

# ---------------------------------------------------------------------------
# Pink boutique theme (matches desktop books already on OneDrive)
# ---------------------------------------------------------------------------

ROSE = "C43B6E"
BLUSH = "F7C6D5"
BLUSH_ROW = "FFF7FA"
SECTION = "F3E6EE"
INK = "4A3038"
MUTED = "8B6B75"
YELLOW_INPUT = "FFF3B0"
if _OPENPYXL_IMPORT_ERROR is None:
    THIN = Border(
        left=Side(style="thin", color="E8C4D0"),
        right=Side(style="thin", color="E8C4D0"),
        top=Side(style="thin", color="E8C4D0"),
        bottom=Side(style="thin", color="E8C4D0"),
    )
    HEADER_FILL = PatternFill("solid", fgColor=BLUSH)
    HEADER_FONT = Font(name="Calibri", size=10, bold=True, color=INK, underline="single")
    HEADER_FONT_PLAIN = Font(name="Calibri", size=10, bold=True, color=INK)
    BODY_FONT = Font(name="Calibri", size=10, color=INK)
    TITLE_FONT = Font(name="Calibri", size=17, bold=True, color=INK)
    SUB_FONT = Font(name="Calibri", size=11, color=MUTED)
    SECTION_FONT = Font(name="Calibri", size=11, bold=True, color=INK)
    LINK_FONT = Font(name="Calibri", size=10, color="2B6CB0", underline="single")
else:  # pragma: no cover — constants still import; Excel helpers need openpyxl
    THIN = HEADER_FILL = HEADER_FONT = HEADER_FONT_PLAIN = None
    BODY_FONT = TITLE_FONT = SUB_FONT = SECTION_FONT = LINK_FONT = None

MA_LIST_WIDTHS = {
    "A": 10,
    "B": 22,
    "C": 16,
    "D": 12,
    "E": 10,
    "F": 10,
    "G": 10,
    "H": 12,
    "I": 22,
    "J": 36,
}
ORDERS_WIDTHS = {
    "A": 12,
    "B": 16,
    "C": 10,
    "D": 12,
    "E": 12,
    "F": 12,
    "G": 14,
    "H": 12,
    "I": 14,
    "J": 14,
    "K": 22,
}
BOT_ACTIVITY_WIDTHS = {
    "A": 12,
    "B": 12,
    "C": 12,
    "D": 10,
    "E": 16,
    "F": 28,
    "G": 12,
}
CANDIDATES_WIDTHS = {
    "A": 6,
    "B": 12,
    "C": 12,
    "D": 12,
    "E": 22,
    "F": 16,
    "G": 12,
    "H": 10,
    "I": 10,
    "J": 36,
    "K": 22,
    "L": 36,
}

# ---------------------------------------------------------------------------
# Mã / photo helpers
# Live site still uses legacy AO001-style (MA_RE). Do not migrate parse_ma
# to A01 here — that is a separate handoff (bc-43bc86e1).
# ---------------------------------------------------------------------------


def _require_openpyxl() -> None:
    """Excel builders need openpyxl; SoT helper constants do not."""
    if _OPENPYXL_IMPORT_ERROR is not None:
        raise ImportError(
            "openpyxl is required for Excel helpers. pip install -r requirements.txt"
        ) from _OPENPYXL_IMPORT_ERROR


def parse_ma(value: object) -> tuple[str, int] | None:
    """Return (prefix, number) for a valid mã; else None. Never invent one.

    Legacy AO001-style until the A01 handoff. Never invent a mã.
    """
    if value is None:
        return None
    text = str(value).strip().upper()
    match = MA_RE.fullmatch(text)
    if not match:
        return None
    return match.group(1), int(match.group(2))


def is_valid_ma(value: object) -> bool:
    return parse_ma(value) is not None


def format_ma(prefix: str, number: int) -> str:
    """Format a Boss/Stock-assigned prefix+number. Does not pick the next code."""
    key = prefix.strip().upper()
    if key not in MA_PREFIX_MEANS:
        raise ValueError(f"unknown mã prefix: {prefix!r}")
    if number < 1 or number > 999:
        raise ValueError(f"mã number out of range: {number}")
    return f"{key}{number:03d}"


class MissingMaError(ValueError):
    """Raised when a script would have to invent a mã — ask Stock instead."""


def require_ma(value: object) -> str:
    """Return a normalized mã or raise MissingMaError (never invent)."""
    parsed = parse_ma(value)
    if parsed is None:
        given = "" if value is None else str(value).strip()
        hint = f" got {given!r}." if given else ""
        raise MissingMaError(f"mã required.{hint} {ASK_STOCK_MA}")
    prefix, number = parsed
    return format_ma(prefix, number)


def normalize_header(name: object) -> str:
    return "" if name is None else str(name).strip().lower()


def resolve_header_key(headers: Sequence[str], logical: str) -> int | None:
    """Column index (0-based) for a logical field, using HEADER_ALIASES."""
    aliases = HEADER_ALIASES.get(logical, (logical,))
    wanted = {normalize_header(a) for a in aliases}
    for i, header in enumerate(headers):
        if normalize_header(header) in wanted:
            return i
    return None


def read_headers(ws: Worksheet, header_row: int) -> list[str]:
    max_col = max(ws.max_column or 1, 1)
    values = [ws.cell(header_row, col).value for col in range(1, max_col + 1)]
    # Trim trailing empty headers.
    while values and (values[-1] is None or str(values[-1]).strip() == ""):
        values.pop()
    return ["" if v is None else str(v).strip() for v in values]


def header_row_for_sheet(title: str) -> int:
    meta = SOT_DATA_SHEETS.get(title, {})
    configured = meta.get("header_row")
    return configured if isinstance(configured, int) else 1


def table_name_for_sheet(title: str) -> str | None:
    meta = SOT_DATA_SHEETS.get(title, {})
    name = meta.get("table")
    return name if isinstance(name, str) and name else None


def resolve_sheet_name(sheetnames: Sequence[str], logical: str) -> str | None:
    aliases = SHEET_ALIASES.get(logical, (logical,))
    wanted = {a.strip().lower() for a in aliases}
    for name in sheetnames:
        if name.strip().lower() in wanted:
            return name
    return None


def photo_filename_for_ma(ma: str, extra: int | None = None) -> str:
    parsed = parse_ma(ma)
    if parsed is None:
        raise ValueError(f"not a valid mã: {ma!r} — never invent stock")
    prefix, num = parsed
    stem = f"{prefix}{num:03d}"
    return f"{stem}_{extra}.jpg" if extra else f"{stem}.jpg"


def photo_filename_for_wish(number: int, extra: int | None = None) -> str:
    if number < 1 or number > 999:
        raise ValueError("wishlist photo numbers are #001–#999")
    stem = f"#{number:03d}"
    return f"{stem}_{extra}.jpg" if extra else f"{stem}.jpg"


def square_photo_folder(ma: object) -> str:
    """On_Hand photo_folder path. Requires a typed mã — never invent one."""
    text = "" if ma is None else str(ma).strip()
    if not text:
        raise ValueError("ma required for photo_folder — never invent")
    return f"{ONEDRIVE_PHOTOS}/{text}/"


def square_photo_folder_formula(row: int) -> str:
    """Fill photo_folder from On_Hand!A{row}. Blank while mã is blank."""
    if row < 2:
        raise ValueError(f"photo_folder formula row must be a data row, got {row}")
    return f'=IF(A{row}="","","{ONEDRIVE_PHOTOS}/"&A{row}&"/")'


def money_usd(value: object) -> Decimal:
    """Parse a USD amount. Blank / None → 0.00."""
    if value is None or str(value).strip() == "":
        return Decimal("0.00")
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def san_jose_sales_tax_usd(taxable_base: object) -> Decimal:
    """San Jose 10.000% on taxable sell price. Never uses 9.375%."""
    base = money_usd(taxable_base)
    return (base * SAN_JOSE_SALES_TAX_RATE).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def finance_net_usd_amount(
    gross_usd: object,
    ship_usd: object = 0,
    discount_usd: object = 0,
) -> Decimal:
    """net = gross + ship − discount. Sales tax is collected separately."""
    return (
        money_usd(gross_usd) + money_usd(ship_usd) - money_usd(discount_usd)
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def finance_taxable_base_amount(gross_usd: object, discount_usd: object = 0) -> Decimal:
    """Default taxable base = max(gross − discount, 0). Shipping taxability is TBD."""
    base = money_usd(gross_usd) - money_usd(discount_usd)
    if base < 0:
        base = Decimal("0.00")
    return base.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def finance_profit_usd(
    *,
    gross_usd: object = 0,
    ship_usd: object = 0,
    discount_usd: object = 0,
    fees_usd: object = 0,
    cogs_usd: object = 0,
    expenses_usd: object = 0,
) -> Decimal:
    """Income-tax style profit: net receipts − fees − COGS − expenses. Not legal advice."""
    receipts = finance_net_usd_amount(gross_usd, ship_usd, discount_usd)
    return (
        receipts - money_usd(fees_usd) - money_usd(cogs_usd) - money_usd(expenses_usd)
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def looks_like_demo_row(values: Iterable[object]) -> bool:
    """True when a data row is kit-seeded demo, not a real piece or buyer."""
    blob = " ".join("" if v is None else str(v) for v in values).strip().lower()
    if not blob:
        return False
    return any(token in blob for token in DEMO_ROW_TOKENS)


def is_photo_link_header(name: str) -> bool:
    return name.strip().lower() == "photo_link"


# ---------------------------------------------------------------------------
# openpyxl helpers — freeze, filter, dropdowns, no embeds
# ---------------------------------------------------------------------------


def style_header_row(
    ws: Worksheet,
    headers: Sequence[str],
    row: int = 1,
    *,
    underline_photo_link: bool = True,
) -> None:
    _require_openpyxl()
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row, col, name)
        font = HEADER_FONT if (underline_photo_link and is_photo_link_header(name)) else HEADER_FONT_PLAIN
        cell.font = font
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
        cell.number_format = FORMAT_TEXT
    ws.row_dimensions[row].height = 22
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"
    ws.freeze_panes = f"A{row + 1}"


def apply_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def paint_tab(ws: Worksheet, color: str = BLUSH) -> None:
    ws.sheet_properties.tabColor = color


def add_list_dropdown(
    ws: Worksheet,
    col: int,
    choices: Sequence[str],
    *,
    start_row: int = 2,
    end_row: int = 200,
    prompt: str | None = None,
) -> DataValidation:
    letter = get_column_letter(col)
    # Quoted comma list — no hidden Lists sheet needed for the lean desktop books.
    formula = '"' + ",".join(choices) + '"'
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        error="Pick a value from the list.",
        errorTitle="status",
        prompt=prompt,
        showInputMessage=bool(prompt),
    )
    dv.add(f"{letter}{start_row}:{letter}{end_row}")
    ws.add_data_validation(dv)
    return dv


def strip_embedded_images(ws: Worksheet) -> int:
    """Remove worksheet drawings. Excel stores photo_link only."""
    images = getattr(ws, "_images", None)
    if not images:
        return 0
    count = len(images)
    images.clear()
    return count


def strip_workbook_images(wb: Workbook) -> int:
    return sum(strip_embedded_images(ws) for ws in wb.worksheets)


def apply_photo_link_style(cell, value: object) -> None:
    """Store a URL as text + hyperlink. Never add a picture."""
    if value is None or str(value).strip() == "":
        cell.value = None
        cell.hyperlink = None
        return
    text = str(value).strip()
    cell.value = text
    if text.startswith(("http://", "https://")):
        cell.hyperlink = text
        cell.font = LINK_FONT
    else:
        # OneDrive path or file name — keep as text, owner pastes the share URL.
        cell.font = BODY_FONT


def header_index(headers: Sequence[str], name: str) -> int:
    needle = name.strip().lower()
    for i, header in enumerate(headers):
        if header.strip().lower() == needle:
            return i
    raise KeyError(name)


def excel_column_letter(index_1based: int) -> str:
    """1-based index → Excel column letter. Works without openpyxl."""
    if index_1based < 1:
        raise ValueError(f"column index must be >= 1, got {index_1based}")
    n = index_1based
    letters: list[str] = []
    while n:
        n, rem = divmod(n - 1, 26)
        letters.append(chr(65 + rem))
    return "".join(reversed(letters))


def finance_header_letter(headers: Sequence[str], name: str) -> str:
    return excel_column_letter(header_index(headers, name) + 1)


def finance_net_usd_formula(row: int) -> str:
    """net = gross + ship − discount. Blank when those three cells are blank.

    Sales tax is not part of net — it is collected separately in sales_tax_usd.
    """
    if row < 2:
        raise ValueError(f"net_usd formula row must be a data row, got {row}")
    gross = finance_header_letter(FINANCE_SALES, "gross_usd")
    ship = finance_header_letter(FINANCE_SALES, "ship_usd")
    discount = finance_header_letter(FINANCE_SALES, "discount_usd")
    return (
        f'=IF(COUNTA({gross}{row},{ship}{row},{discount}{row})=0,"",'
        f"N({gross}{row})+N({ship}{row})-N({discount}{row}))"
    )


def finance_taxable_base_formula(row: int) -> str:
    """Default taxable base = max(gross − discount, 0). Blank until a sale is typed."""
    if row < 2:
        raise ValueError(f"taxable_base formula row must be a data row, got {row}")
    gross = finance_header_letter(FINANCE_SALES, "gross_usd")
    discount = finance_header_letter(FINANCE_SALES, "discount_usd")
    return (
        f'=IF(COUNTA({gross}{row},{discount}{row})=0,"",'
        f"MAX(N({gross}{row})-N({discount}{row}),0))"
    )


def finance_sales_tax_rate_formula(row: int) -> str:
    """San Jose 10.000% once a taxable base exists. Cap may override the cell."""
    if row < 2:
        raise ValueError(f"sales_tax_rate formula row must be a data row, got {row}")
    taxable = finance_header_letter(FINANCE_SALES, "taxable_base_usd")
    rate = f"{SAN_JOSE_SALES_TAX_RATE:.2f}"
    return f'=IF({taxable}{row}="","",{rate})'


def finance_sales_tax_usd_formula(row: int) -> str:
    """ROUND(taxable_base * rate, 2). Blank when no sale."""
    if row < 2:
        raise ValueError(f"sales_tax_usd formula row must be a data row, got {row}")
    taxable = finance_header_letter(FINANCE_SALES, "taxable_base_usd")
    rate = finance_header_letter(FINANCE_SALES, "sales_tax_rate")
    return (
        f'=IF(OR({taxable}{row}="",{rate}{row}=""),"",'
        f"ROUND(N({taxable}{row})*N({rate}{row}),2))"
    )


def finance_tax_month_col_letter(month: int) -> str:
    """January=1 → column B. December=12 → column M."""
    if month < 1 or month > 12:
        raise ValueError(f"month out of range: {month}")
    return excel_column_letter(month + 1)


def finance_month_bounds(year: int, month: int) -> tuple[int, int, int, int]:
    """Return (start_year, start_month, end_year, end_month) exclusive end."""
    if month < 1 or month > 12:
        raise ValueError(f"month out of range: {month}")
    if month == 12:
        return year, 12, year + 1, 1
    return year, month, year, month + 1


def finance_month_sumifs(
    sheet: str,
    amount_col: str,
    date_col: str,
    year: int,
    month: int,
    last_row: int = SQUARE_FORMULA_LAST_ROW,
) -> str:
    start_year, start_month, end_year, end_month = finance_month_bounds(year, month)
    return (
        f"=SUMIFS({sheet}!${amount_col}$2:${amount_col}${last_row},"
        f"{sheet}!${date_col}$2:${date_col}${last_row},"
        f'">="&DATE({start_year},{start_month},1),'
        f"{sheet}!${date_col}$2:${date_col}${last_row},"
        f'"<"&DATE({end_year},{end_month},1))'
    )


def finance_sheet_amount_col(sheet: str, header: str) -> str:
    if sheet == "Sales":
        headers = FINANCE_SALES
    elif sheet == "Fees":
        headers = FINANCE_FEES
    elif sheet == "Expenses":
        headers = FINANCE_EXPENSES
    elif sheet == "Payouts_Transfers":
        headers = FINANCE_PAYOUTS
    else:
        raise KeyError(f"unknown finance sheet {sheet!r}")
    return finance_header_letter(headers, header)


def finance_metric_month_formula(
    kind: FinanceMetric,
    year: int,
    month: int,
    last_row: int = SQUARE_FORMULA_LAST_ROW,
) -> str:
    date_col = "A"
    if kind == "gross":
        return finance_month_sumifs("Sales", finance_sheet_amount_col("Sales", "gross_usd"), date_col, year, month, last_row)
    if kind == "shipping":
        return finance_month_sumifs("Sales", finance_sheet_amount_col("Sales", "ship_usd"), date_col, year, month, last_row)
    if kind == "discount":
        return finance_month_sumifs("Sales", finance_sheet_amount_col("Sales", "discount_usd"), date_col, year, month, last_row)
    if kind == "net":
        return finance_month_sumifs("Sales", finance_sheet_amount_col("Sales", "net_usd"), date_col, year, month, last_row)
    if kind == "taxable_base":
        return finance_month_sumifs(
            "Sales", finance_sheet_amount_col("Sales", "taxable_base_usd"), date_col, year, month, last_row
        )
    if kind == "sales_tax":
        return finance_month_sumifs(
            "Sales", finance_sheet_amount_col("Sales", "sales_tax_usd"), date_col, year, month, last_row
        )
    if kind == "fees":
        return finance_month_sumifs("Fees", finance_sheet_amount_col("Fees", "amount_usd"), date_col, year, month, last_row)
    if kind == "cogs":
        return finance_month_sumifs("Sales", finance_sheet_amount_col("Sales", "cogs_usd"), date_col, year, month, last_row)
    if kind == "expenses":
        return finance_month_sumifs(
            "Expenses", finance_sheet_amount_col("Expenses", "amount_usd"), date_col, year, month, last_row
        )
    if kind == "profit":
        letter = finance_tax_month_col_letter(month)
        return finance_net_month_formula(letter)
    unused: Never = kind
    raise AssertionError(f"unhandled finance metric: {unused!r}")


def finance_net_month_formula(col_letter: str) -> str:
    """Profit = net receipts − fees − COGS − expenses. Sales tax is not income."""
    # Metric rows start at Excel row 2 in the same order as FINANCE_TAX_SUMMARY_METRICS.
    row = {kind: index + 2 for index, (_label, kind) in enumerate(FINANCE_TAX_SUMMARY_METRICS)}
    return (
        f"=N({col_letter}{row['net']})"
        f"-N({col_letter}{row['fees']})"
        f"-N({col_letter}{row['cogs']})"
        f"-N({col_letter}{row['expenses']})"
    )


def finance_ytd_formula(row: int) -> str:
    return f"=SUM(B{row}:M{row})"


def finance_ytd_label_formula(label: str) -> str:
    for row_label, formula in FINANCE_TAX_SUMMARY_ROWS:
        if row_label == label:
            return formula
    raise KeyError(label)


def assert_photo_link_last(headers: Sequence[str], sheet_label: str) -> None:
    if not headers or headers[-1] != "photo_link":
        raise AssertionError(f"{sheet_label}: photo_link must be the last column, got {headers!r}")


assert_photo_link_last(MA_LIST, "MA_LIST")
assert_photo_link_last(CANDIDATES, "CANDIDATES")
assert_photo_link_last(SOT_WISHLIST, "SOT_WISHLIST")


def photo_link_col_index(headers: Sequence[str]) -> int | None:
    """0-based index of photo_link, if present."""
    return resolve_header_key(headers, "photo_link")


def official_status_or_raise(value: object) -> str:
    text = "" if value is None else str(value).strip()
    for allowed in SOT_OFFICIAL_STATUS:
        if text.lower() == allowed.lower():
            return allowed
    raise ValueError(
        f"Official status must be one of {list(SOT_OFFICIAL_STATUS)}, got {text!r}"
    )


# Keep these names importable from schema through the A01 cutover (bc-43bc86e1)
# and any site↔Wishlist slimming. Re-export is fine; dropping them breaks
# excel-kit/sot/append_official_row.py.
APPEND_OFFICIAL_ROW_EXPORTS: tuple[str, ...] = (
    "ASK_STOCK_MA",
    "MissingMaError",
    "OFFICIAL_TO_MA_LIST_STATUS",
    "SOT_OFFICIAL_STATUS",
    "SQUARE_SOT_SHORT",
    "official_status_or_raise",
    "parse_ma",
    "photo_filename_for_ma",
    "require_ma",
    "resolve_sheet_name",
)


def order_status_or_raise(value: object) -> str:
    text = "" if value is None else str(value).strip()
    for allowed in ORDER_STATUS:
        if text.lower() == allowed.lower():
            return allowed
    raise ValueError(
        f"order status must be one of {list(ORDER_STATUS)}, got {text!r}"
    )


def wishlist_status_or_raise(value: object) -> str:
    text = "" if value is None else str(value).strip()
    for allowed in SOT_WISHLIST_STATUS:
        if text.lower() == allowed.lower():
            return allowed
    raise ValueError(
        f"wishlist status must be one of {list(SOT_WISHLIST_STATUS)}, got {text!r}"
    )


def new_boutique_workbook() -> Workbook:
    _require_openpyxl()
    wb = Workbook()
    wb.properties.creator = "Sassy Closet excel-kit"
    return wb


def write_how_to_use(ws: Worksheet, lines: Sequence[tuple[str, str]]) -> None:
    """lines: (role, text) where role is title|sub|rule|section|body|divider|foot."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 28
    fills = {
        "title": PatternFill("solid", fgColor=BLUSH),
        "sub": PatternFill("solid", fgColor=BLUSH_ROW),
        "rule": PatternFill("solid", fgColor=BLUSH_ROW),
        "section": PatternFill("solid", fgColor=SECTION),
        "body": PatternFill("solid", fgColor=BLUSH_ROW),
        "divider": PatternFill("solid", fgColor=BLUSH_ROW),
        "foot": PatternFill("solid", fgColor=BLUSH),
    }
    fonts = {
        "title": TITLE_FONT,
        "sub": SUB_FONT,
        "rule": BODY_FONT,
        "section": SECTION_FONT,
        "body": BODY_FONT,
        "divider": Font(name="Calibri", size=10, color=MUTED),
        "foot": Font(name="Calibri", size=10, italic=True, color=INK),
    }
    for row, (role, text) in enumerate(lines, start=1):
        cell = ws.cell(row, 1, text)
        cell.font = fonts[role]
        cell.fill = fills[role]
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 22 if role != "title" else 26
    paint_tab(ws)


def default_out_dir() -> Path:
    return Path.cwd() / "out"
