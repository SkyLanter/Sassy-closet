#!/usr/bin/env python3
"""Build the ONE sassycloset hub from the live site export.

All + category sheets + Orders + Readme in sassycloset.xlsx.
Sync Photos to out/Photos/{ma}/ from the live photo API.
Cute / pink / emoji / embeds are retired.
Website stays GF intake. Square Free remains on-hand inventory truth.
Never invent mã. No Square Save. No Facebook Post.

Usage:
  python3 excel-kit/build_sassycloset_hub.py --out-dir ./out
  python3 excel-kit/build_sassycloset_hub.py --from-csv path.csv --skip-photos --out-dir ./out
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import shutil
import subprocess
import sys
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (  # noqa: E402
    EXPORT_HEADERS,
    HUB_ALL,
    HUB_EXPORT_URL,
    HUB_KIND_SHEETS,
    HUB_NAME,
    HUB_ORDERS,
    HUB_ORDER_PAY,
    HUB_ORDER_STATUS,
    HUB_PHOTO_API,
    HUB_README_LINES,
    HUB_README_TXT,
    HUB_RETIRED_MA,
    HUB_SHEETS,
    HUB_TEMPLATE_ROWS,
    HUB_XLSX_NAME,
    ONEDRIVE_HUB,
    ONEDRIVE_HUB_DIR,
    ONEDRIVE_HUB_PHOTOS,
    add_list_dropdown,
    apply_photo_link_style,
    assert_hub_backup_columns,
    header_index,
    hub_photo_folder,
    is_hub_ma,
    strip_workbook_images,
)

PLAIN_FONT = Font(name="Calibri", size=11)
PLAIN_HEADER_FONT = Font(name="Calibri", size=11, bold=True)
NUMBER_HEADERS = frozenset({"sell_usd", "cost", "amount_usd"})
RETIRED_FILL_RGB = frozenset({"F7C6D5", "C43B6E", "FFF7FA", "F3E6EE", "FFF3B0"})
MAX_PHOTOS_PER_MA = 40
ALL_WIDTHS = {
    "ma": 10,
    "kind": 8,
    "colors": 18,
    "sell_usd": 11,
    "cost": 11,
    "currency": 10,
    "square": 12,
    "status": 12,
    "flag": 12,
    "next_desk": 14,
    "photo_folder": 42,
    "source_link": 36,
}
ORDERS_WIDTHS = {
    "date": 12,
    "ma": 10,
    "customer": 18,
    "pay": 10,
    "amount_usd": 12,
    "ship_or_meetup": 16,
    "status": 12,
    "notes": 24,
}
FORBIDDEN_SHEETS = frozenset(
    {
        "Dashboard",
        "Bot_Activity",
        "Wishlist",
        "START HERE",
        "How_to_use",
        "Candidates",
        "Inventory",
        "Track",
    }
)
USER_AGENT = "sassy-closet-excel-kit/sassycloset-hub"


class ExportError(RuntimeError):
    """Live export could not be read or did not match the locked headers."""


class PhotoSyncError(RuntimeError):
    """Photo API failed in a way that is not a simple missing file."""


def fetch_export_csv(url: str) -> str:
    request = urllib.request.Request(
        url,
        headers={
            "Accept": "text/csv, text/plain;q=0.9, */*;q=0.1",
            "User-Agent": USER_AGENT,
        },
        method="GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=45) as response:
            raw = response.read()
    except urllib.error.URLError as exc:
        raise ExportError(f"could not fetch {url}: {exc}") from exc
    return raw.decode("utf-8-sig")


def coerce_number(value: object) -> object:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return int(number)
    return number


def text_or_none(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def pick_cost(raw: dict[str, str]) -> tuple[object, str | None]:
    """Use the export's own cost + currency. Do not convert or invent."""
    currency = text_or_none(raw.get("cost_currency"))
    if currency is None:
        return None, None
    key = currency.upper()
    if key == "CNY":
        return coerce_number(raw.get("cost_cny")), currency
    if key == "USD":
        return coerce_number(raw.get("cost_usd")), currency
    return None, currency


def flag_from_export(raw: dict[str, str]) -> str | None:
    """Copy a flag only if the export already has one. Never invent stock."""
    return text_or_none(raw.get("flag"))


def map_export_row(raw: dict[str, str]) -> dict[str, object] | None:
    ma = text_or_none(raw.get("ma"))
    if ma is None:
        return None
    if not is_hub_ma(ma):
        raise ExportError(f"export mã {ma!r} is not a live hub mã — never invent")
    cost, currency = pick_cost(raw)
    return {
        "ma": ma,
        "kind": text_or_none(raw.get("kind")),
        "colors": text_or_none(raw.get("color")),
        "sell_usd": coerce_number(raw.get("sell_usd")),
        "cost": cost,
        "currency": currency,
        "square": text_or_none(raw.get("square")),
        "status": text_or_none(raw.get("status")),
        "flag": flag_from_export(raw),
        "next_desk": text_or_none(raw.get("next_desk")),
        "photo_folder": hub_photo_folder(ma),
        "source_link": text_or_none(raw.get("source_link")),
    }


def read_export_text(text: str) -> list[dict[str, object]]:
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ExportError("export has no headers")
    headers = [str(name).strip() for name in reader.fieldnames]
    missing = [name for name in EXPORT_HEADERS if name not in headers]
    if missing:
        raise ExportError(f"export missing columns: {missing} got={headers}")
    rows: list[dict[str, object]] = []
    skipped_empty_ma = 0
    for raw in reader:
        cleaned = {str(k).strip(): ("" if v is None else str(v)) for k, v in raw.items()}
        item = map_export_row(cleaned)
        if item is None:
            if any(str(cleaned.get(h) or "").strip() for h in EXPORT_HEADERS):
                skipped_empty_ma += 1
            continue
        rows.append(item)
    if skipped_empty_ma:
        print(
            f"note: skipped {skipped_empty_ma} export row(s) with empty mã "
            "(never invent)",
            file=sys.stderr,
        )
    return rows


def retired_mas(rows: Iterable[dict[str, object]]) -> list[str]:
    banned = {code.upper() for code in HUB_RETIRED_MA}
    found: list[str] = []
    for row in rows:
        ma = str(row.get("ma") or "").strip()
        if ma.upper() in banned:
            found.append(ma)
    return found


def style_plain_header(ws: Worksheet, headers: tuple[str, ...]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.font = PLAIN_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.number_format = FORMAT_TEXT
    ws.row_dimensions[1].height = 18
    last = get_column_letter(len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last}{max(ws.max_row or 1, 1)}"


def apply_named_widths(ws: Worksheet, headers: tuple[str, ...], widths: dict[str, float]) -> None:
    for index, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(name, 12)


def write_item_sheet(ws: Worksheet, title: str, rows: list[dict[str, object]]) -> None:
    ws.title = title
    for col, name in enumerate(HUB_ALL, start=1):
        ws.cell(1, col, name)
    wrap = {"colors", "photo_folder", "source_link"}
    textish = {"ma", "photo_folder", "source_link", "flag", "next_desk"}
    for row_index, item in enumerate(rows, start=2):
        for col, name in enumerate(HUB_ALL, start=1):
            cell = ws.cell(row_index, col)
            if name == "source_link":
                apply_photo_link_style(cell, item.get(name))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.number_format = FORMAT_TEXT
                continue
            cell.value = item.get(name)
            cell.font = PLAIN_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=name in wrap)
            if name in NUMBER_HEADERS and isinstance(item.get(name), (int, float)):
                cell.number_format = "0.##"
            elif name in textish:
                cell.number_format = FORMAT_TEXT
    style_plain_header(ws, HUB_ALL)
    last_row = max(len(rows) + 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HUB_ALL))}{last_row}"
    apply_named_widths(ws, HUB_ALL, ALL_WIDTHS)


def rows_for_kind(rows: list[dict[str, object]], letter: str) -> list[dict[str, object]]:
    want = letter.strip().upper()
    matched: list[dict[str, object]] = []
    for row in rows:
        kind = str(row.get("kind") or "").strip().upper()
        if kind == want:
            matched.append(row)
    return matched


def write_orders(ws: Worksheet) -> None:
    ws.title = "Orders"
    for col, name in enumerate(HUB_ORDERS, start=1):
        ws.cell(1, col, name)
    for row in range(2, HUB_TEMPLATE_ROWS + 2):
        for col in range(1, len(HUB_ORDERS) + 1):
            cell = ws.cell(row, col, None)
            cell.font = PLAIN_FONT
            cell.alignment = Alignment(vertical="center")
    style_plain_header(ws, HUB_ORDERS)
    last = HUB_TEMPLATE_ROWS + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HUB_ORDERS))}{last}"
    apply_named_widths(ws, HUB_ORDERS, ORDERS_WIDTHS)
    add_list_dropdown(
        ws,
        header_index(HUB_ORDERS, "pay") + 1,
        HUB_ORDER_PAY,
        start_row=2,
        end_row=last,
        prompt="pay",
    )
    add_list_dropdown(
        ws,
        header_index(HUB_ORDERS, "status") + 1,
        HUB_ORDER_STATUS,
        start_row=2,
        end_row=last,
        prompt="status",
    )


def write_readme(ws: Worksheet) -> None:
    ws.title = "Readme"
    ws.column_dimensions["A"].width = 92
    for row, line in enumerate(HUB_README_LINES, start=1):
        cell = ws.cell(row, 1, line)
        cell.font = PLAIN_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 18


def new_hub_workbook() -> WorkbookType:
    workbook = Workbook()
    workbook.properties.creator = "Sassy Closet excel-kit"
    workbook.properties.title = "sassycloset"
    return workbook


def build_workbook(rows: list[dict[str, object]]) -> WorkbookType:
    workbook = new_hub_workbook()
    write_item_sheet(workbook.active, "All", rows)
    for sheet_name, letter in HUB_KIND_SHEETS:
        write_item_sheet(workbook.create_sheet(sheet_name), sheet_name, rows_for_kind(rows, letter))
    write_orders(workbook.create_sheet("Orders"))
    write_readme(workbook.create_sheet("Readme"))
    for index, name in enumerate(HUB_SHEETS):
        workbook.move_sheet(name, offset=index - workbook.sheetnames.index(name))
    strip_workbook_images(workbook)
    return workbook


def fetch_photo_bytes(ma: str, name: str, photo_api: str) -> bytes | None:
    url = f"{photo_api.rstrip('/')}/{ma}/{name}"
    request = urllib.request.Request(
        url,
        headers={"Accept": "image/jpeg, image/*;q=0.9", "User-Agent": USER_AGENT},
        method="GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=45) as response:
            status = getattr(response, "status", 200)
            if status == 404:
                return None
            payload = response.read()
            content_type = str(response.headers.get("Content-Type") or "").lower()
    except urllib.error.HTTPError as exc:
        if exc.code == 404:
            return None
        raise PhotoSyncError(f"photo GET {url} failed: HTTP {exc.code}") from exc
    except urllib.error.URLError as exc:
        raise PhotoSyncError(f"photo GET {url} failed: {exc}") from exc
    if not payload:
        return None
    if "json" in content_type or payload.startswith(b"{") or payload.startswith(b"<!DOCTYPE"):
        return None
    if not (payload.startswith(b"\xff\xd8") or "image/" in content_type):
        raise PhotoSyncError(f"photo GET {url} returned non-image ({content_type})")
    return payload


def sync_photos(
    mas: list[str],
    dest_root: Path,
    *,
    skip: bool = False,
    photo_api: str = HUB_PHOTO_API,
) -> dict[str, list[str]]:
    dest_root.mkdir(parents=True, exist_ok=True)
    live = {ma.strip() for ma in mas if ma.strip()}
    banned = {code.upper() for code in HUB_RETIRED_MA}
    if dest_root.is_dir():
        for child in dest_root.iterdir():
            if not child.is_dir():
                continue
            if child.name not in live or child.name.upper() in banned:
                shutil.rmtree(child)
    written: dict[str, list[str]] = {}
    for ma in mas:
        if ma.upper() in banned:
            raise PhotoSyncError(
                f"refusing to sync photos for retired mã {ma} "
                "(A02 was renamed to P02, not P05)"
            )
        if not is_hub_ma(ma):
            raise PhotoSyncError(f"refusing to sync photos for invented mã {ma!r}")
        folder = dest_root / ma
        folder.mkdir(parents=True, exist_ok=True)
        files: list[str] = []
        if not skip:
            for index in range(1, MAX_PHOTOS_PER_MA + 1):
                name = f"{index:03d}.jpg"
                payload = fetch_photo_bytes(ma, name, photo_api)
                if payload is None:
                    break
                (folder / name).write_bytes(payload)
                files.append(name)
            keep = set(files)
            for leftover in folder.iterdir():
                if leftover.is_file() and leftover.name not in keep:
                    leftover.unlink()
        written[ma] = files
    return written


def write_hub_readme_txt(path: Path) -> None:
    path.write_text(HUB_README_TXT, encoding="utf-8")


def unzip_test(path: Path) -> None:
    proc = subprocess.run(
        ["unzip", "-t", str(path)],
        check=False,
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        raise AssertionError(f"unzip -t failed:\n{proc.stdout}\n{proc.stderr}")
    if "No errors detected" not in proc.stdout:
        raise AssertionError(f"unzip -t did not confirm a clean zip:\n{proc.stdout}")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    media = [
        name
        for name in names
        if "/media/" in name.lower() or name.lower().endswith((".png", ".jpg", ".jpeg", ".gif"))
    ]
    if media:
        raise AssertionError(f"embedded media in xlsx zip: {media}")


def _fill_rgb(fill: object) -> str | None:
    if fill is None:
        return None
    if getattr(fill, "patternType", None) in (None, "none"):
        return None
    fg = getattr(fill, "fgColor", None)
    if fg is None:
        return None
    rgb = getattr(fg, "rgb", None)
    if not rgb:
        return None
    text = str(rgb).upper()
    if len(text) == 8:
        text = text[2:]
    return text


def column_values(ws: Worksheet, headers: tuple[str, ...], header: str) -> list[str]:
    idx = header_index(list(headers), header) + 1
    values: list[str] = []
    for row in range(2, (ws.max_row or 1) + 1):
        value = ws.cell(row, idx).value
        if value is None or str(value).strip() == "":
            continue
        values.append(str(value).strip())
    return values


def nonempty_row_count(ws: Worksheet, headers: tuple[str, ...], header_row: int = 1) -> int:
    count = 0
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        values = [ws.cell(row, col).value for col in range(1, len(headers) + 1)]
        if any(value is not None and str(value).strip() != "" for value in values):
            count += 1
    return count


def _assert_item_sheet(ws: Worksheet, title: str) -> list[str]:
    if ws.title != title:
        raise AssertionError(f"sheet title {ws.title!r} != {title!r}")
    headers = [ws.cell(1, col).value for col in range(1, len(HUB_ALL) + 1)]
    if headers != list(HUB_ALL):
        raise AssertionError(f"{title} headers drifted: {headers}")
    if ws.freeze_panes != "A2":
        raise AssertionError(f"{title} freeze_panes={ws.freeze_panes}")
    if not ws.auto_filter.ref or not str(ws.auto_filter.ref).startswith("A1:"):
        raise AssertionError(f"{title} autofilter={ws.auto_filter.ref}")
    if ws.merged_cells.ranges:
        raise AssertionError(f"{title} merged cells forbidden: {ws.merged_cells.ranges}")
    return column_values(ws, HUB_ALL, "ma")


def verify_hub_book(path: Path, *, expected_mas: Iterable[str] | None = None) -> list[str]:
    unzip_test(path)
    workbook = load_workbook(path)
    notes: list[str] = []
    try:
        if list(workbook.sheetnames) != list(HUB_SHEETS):
            raise AssertionError(f"sheets must be {list(HUB_SHEETS)}, got {workbook.sheetnames}")
        extra = FORBIDDEN_SHEETS.intersection(workbook.sheetnames)
        if extra:
            raise AssertionError(f"retired/forbidden sheets present: {sorted(extra)}")

        all_sheet = workbook["All"]
        assert_hub_backup_columns(
            [all_sheet.cell(1, col).value for col in range(1, len(HUB_ALL) + 1)],
            "All",
        )
        mas = _assert_item_sheet(all_sheet, "All")
        links = column_values(all_sheet, HUB_ALL, "source_link")
        if not links and mas:
            print(
                "note: All has mã rows but no source_link values "
                "(column kept; do not invent Taobao URLs)",
                file=sys.stderr,
            )
        if expected_mas is not None:
            want = [str(ma).strip() for ma in expected_mas]
            if list(mas) != want:
                raise AssertionError(f"All mã mismatch: {list(mas)} != {want}")
        leftover = [ma for ma in mas if ma.upper() in {code.upper() for code in HUB_RETIRED_MA}]
        if leftover:
            raise AssertionError(f"leftover retired mã {leftover} (A02 was renamed to P02)")

        folders = column_values(all_sheet, HUB_ALL, "photo_folder")
        for ma, folder in zip(mas, folders, strict=True):
            want = hub_photo_folder(ma)
            if folder != want:
                raise AssertionError(f"photo_folder for {ma} must be {want!r}, got {folder!r}")
        if len(folders) != len(mas):
            raise AssertionError("every All row needs a photo_folder")

        kind_col = header_index(list(HUB_ALL), "kind") + 1
        kind_by_ma: dict[str, str] = {}
        for row_index, ma in enumerate(mas, start=2):
            kind_by_ma[ma] = str(all_sheet.cell(row_index, kind_col).value or "").strip().upper()
        for sheet_name, letter in HUB_KIND_SHEETS:
            sheet = workbook[sheet_name]
            sheet_mas = _assert_item_sheet(sheet, sheet_name)
            expected = [ma for ma, kind in kind_by_ma.items() if kind == letter]
            if sheet_mas != expected:
                raise AssertionError(
                    f"{sheet_name} must be All filtered by kind={letter}: "
                    f"{sheet_mas} != {expected}"
                )

        orders = workbook["Orders"]
        order_headers = [orders.cell(1, col).value for col in range(1, len(HUB_ORDERS) + 1)]
        if order_headers != list(HUB_ORDERS):
            raise AssertionError(f"Orders headers drifted: {order_headers}")
        if orders.freeze_panes != "A2":
            raise AssertionError(f"Orders freeze_panes={orders.freeze_panes}")
        if nonempty_row_count(orders, HUB_ORDERS) != 0:
            raise AssertionError("Orders must be a blank log (no invented customers)")
        if len(orders.data_validations.dataValidation) < 2:
            raise AssertionError("Orders needs pay + status dropdowns")

        readme = workbook["Readme"]
        lines = [
            str(readme.cell(row, 1).value).strip()
            for row in range(1, 8)
            if readme.cell(row, 1).value and str(readme.cell(row, 1).value).strip()
        ]
        if not (4 <= len(lines) <= 6):
            raise AssertionError(f"Readme must be hub-rule lines, got {len(lines)}: {lines}")
        blob = " ".join(lines).lower()
        for needle in ("teammate", "onedrive", "square", "invent", "save", "post"):
            if needle not in blob:
                raise AssertionError(f"Readme should mention {needle!r}")
        if any(ch in "".join(lines) for ch in "🎀♡💗✨"):
            raise AssertionError("Readme must stay plain (no cute emoji)")
        if readme.merged_cells.ranges:
            raise AssertionError("Readme merged title art is retired")

        for sheet in workbook.worksheets:
            if getattr(sheet, "_images", None):
                raise AssertionError(f"{sheet.title}: embedded images forbidden")
            if sheet.merged_cells.ranges:
                raise AssertionError(f"{sheet.title}: merged title art is retired")
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 1, 40), max_col=sheet.max_column or 1):
                for cell in row:
                    rgb = _fill_rgb(cell.fill)
                    if rgb and rgb in RETIRED_FILL_RGB:
                        raise AssertionError(f"{sheet.title}!{cell.coordinate} uses retired cute fill {rgb}")
                    font = cell.font
                    if font and font.size and font.size > 14:
                        raise AssertionError(f"{sheet.title}!{cell.coordinate} title-art font size {font.size}")

        notes.append(f"all_rows {len(mas)}")
        notes.append(f"mas {','.join(mas)}")
        notes.append(f"photo_root {ONEDRIVE_HUB_PHOTOS}/{{ma}}/")
        notes.append(f"sheets {','.join(workbook.sheetnames)}")
    finally:
        workbook.close()
    return notes


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_rows(from_csv: Path | None, export_url: str) -> list[dict[str, object]]:
    if from_csv is not None:
        return read_export_text(Path(from_csv).read_text(encoding="utf-8-sig"))
    return read_export_text(fetch_export_csv(export_url))


def print_land_path() -> None:
    print(f"OneDrive land path (Build lands later): {ONEDRIVE_HUB_DIR}/")
    print(f"  {HUB_XLSX_NAME}")
    print("  Photos/{MA}/001.jpg")
    print("  README.txt")
    print(f"sync target: OneDrive {ONEDRIVE_HUB}")


def build(
    out_dir: Path,
    *,
    from_csv: Path | None = None,
    export_url: str = HUB_EXPORT_URL,
    skip_photos: bool = False,
    photo_api: str = HUB_PHOTO_API,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / HUB_XLSX_NAME
    rows = load_rows(from_csv, export_url)
    leftover = retired_mas(rows)
    if leftover:
        raise ExportError(
            f"leftover retired mã {leftover} in export. "
            "A02 was renamed to P02 (not P05). Never invent mã."
        )
    workbook = build_workbook(rows)
    workbook.save(path)
    workbook.close()
    mas = [str(row["ma"]) for row in rows]
    photos = sync_photos(mas, out_dir / "Photos", skip=skip_photos, photo_api=photo_api)
    write_hub_readme_txt(out_dir / "README.txt")
    notes = verify_hub_book(path, expected_mas=mas)
    digest = sha256_file(path)
    print(f"wrote {path}")
    print(f"sha256 {digest}")
    for note in notes:
        print(note)
    photo_count = sum(len(names) for names in photos.values())
    print(f"photos {photo_count} files under {out_dir / 'Photos'}/{{ma}}/")
    if skip_photos:
        print("photos skipped (folders only; never invented mã)")
    if from_csv is None:
        print(f"source {export_url}")
    else:
        print(f"source csv {from_csv}")
    print_land_path()
    print(f"{HUB_NAME} hub for teammates; cute Excel is retired")
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=Path("out"),
        help="folder for sassycloset.xlsx + Photos/{ma}/ + README.txt (default: ./out)",
    )
    parser.add_argument(
        "--from-csv",
        type=Path,
        default=None,
        help="use a local export CSV instead of fetching the live site (tests)",
    )
    parser.add_argument(
        "--export-url",
        default=HUB_EXPORT_URL,
        help=f"live GET CSV (default: {HUB_EXPORT_URL})",
    )
    parser.add_argument(
        "--skip-photos",
        action="store_true",
        help="create Photos/{ma}/ folders only; do not GET the live photo API",
    )
    parser.add_argument(
        "--photo-api",
        default=HUB_PHOTO_API,
        help=f"live photo root (default: {HUB_PHOTO_API})",
    )
    args = parser.parse_args(argv)
    try:
        build(
            args.out_dir,
            from_csv=args.from_csv,
            export_url=args.export_url,
            skip_photos=args.skip_photos,
            photo_api=args.photo_api,
        )
    except (ExportError, PhotoSyncError, AssertionError, OSError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
