#!/usr/bin/env python3
"""Build one plain Sassy_Closet_Data.xlsx from the live site export.

Cute / pink / emoji / phone Excel themes are retired.
Website stays GF input. This book is Boss's simple mirror / edit buffer.
Square Free remains on-hand inventory truth. Never invent mã. No embeds.

Usage:
  python3 excel-kit/build_plain_data.py --out-dir ./out
  python3 excel-kit/build_plain_data.py --from-csv path.csv --out-dir ./out
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import subprocess
import sys
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (  # noqa: E402
    ONEDRIVE_PHOTOS,
    ONEDRIVE_PLAIN_DATA,
    PLAIN_CANDIDATE_STATUS,
    PLAIN_CANDIDATES,
    PLAIN_EXPORT_URL,
    PLAIN_INVENTORY,
    PLAIN_ORDER_PAY,
    PLAIN_ORDER_STATUS,
    PLAIN_ORDERS,
    PLAIN_RETIRED_MA,
    PLAIN_SHEETS,
    PLAIN_TEMPLATE_ROWS,
    PLAIN_XLSX_NAME,
    add_list_dropdown,
    header_index,
    photo_folder_link,
    strip_workbook_images,
)

PLAIN_FONT = Font(name="Calibri", size=11)
PLAIN_HEADER_FONT = Font(name="Calibri", size=11, bold=True)
NUMBER_HEADERS = frozenset(
    {"cost_cny", "cost_usd", "sell_cny", "sell_usd", "amount_usd"}
)
# Blush / rose fills used by the retired cute books — must not appear here.
RETIRED_FILL_RGB = frozenset({"F7C6D5", "C43B6E", "FFF7FA", "F3E6EE", "FFF3B0"})
INVENTORY_WIDTHS = {
    "ma": 10,
    "kind": 8,
    "kind_vi": 12,
    "size": 12,
    "color": 16,
    "color_note": 16,
    "color_pieces": 22,
    "blurb": 22,
    "cost_cny": 11,
    "cost_usd": 11,
    "cost_currency": 12,
    "sell_cny": 11,
    "sell_usd": 11,
    "sell_currency": 12,
    "source_link": 36,
    "status": 12,
    "square": 12,
    "created_at": 22,
    "updated_at": 22,
    "photo_link": 40,
}
CANDIDATES_WIDTHS = {
    "title_note": 22,
    "source_link": 36,
    "size_note": 14,
    "color_note": 14,
    "cost_note": 12,
    "status": 12,
    "photo_note": 22,
}
ORDERS_WIDTHS = {
    "date": 12,
    "ma": 10,
    "customer_note": 20,
    "pay_method": 12,
    "amount_usd": 12,
    "ship_or_meetup": 16,
    "status": 12,
    "notes": 24,
}


class ExportError(RuntimeError):
    """Live export could not be read or did not match the locked headers."""


def fetch_export_csv(url: str) -> str:
    request = urllib.request.Request(
        url,
        headers={
            "Accept": "text/csv, text/plain;q=0.9, */*;q=0.1",
            "User-Agent": "sassy-closet-excel-kit/plain-data",
        },
        method="GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=45) as response:
            raw = response.read()
    except urllib.error.URLError as exc:
        raise ExportError(f"could not fetch {url}: {exc}") from exc
    return raw.decode("utf-8-sig")


def read_export_text(text: str) -> tuple[list[str], list[dict[str, object]]]:
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ExportError("export has no headers")
    headers = [str(name).strip() for name in reader.fieldnames]
    expected = list(PLAIN_INVENTORY)
    missing = [name for name in expected if name not in headers]
    extra = [name for name in headers if name not in expected]
    if missing or extra or headers != expected:
        raise ExportError(
            "export headers must match Inventory columns exactly. "
            f"missing={missing} extra={extra} got={headers}"
        )
    rows: list[dict[str, object]] = []
    skipped_empty_ma = 0
    for raw in reader:
        ma = "" if raw.get("ma") is None else str(raw.get("ma")).strip()
        if not ma:
            if any(str(raw.get(h) or "").strip() for h in expected):
                skipped_empty_ma += 1
            continue
        item: dict[str, object] = {}
        for name in expected:
            if name == "ma":
                item[name] = ma
                continue
            if name == "photo_link":
                item[name] = photo_folder_link(ma)
                continue
            item[name] = coerce_cell(name, raw.get(name))
        rows.append(item)
    if skipped_empty_ma:
        print(
            f"note: skipped {skipped_empty_ma} export row(s) with empty mã "
            "(never invent)",
            file=sys.stderr,
        )
    return expected, rows


def coerce_cell(header: str, value: object) -> object:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    if header in NUMBER_HEADERS:
        try:
            number = float(text)
        except ValueError:
            return text
        if number.is_integer():
            return int(number)
        return number
    return text


def style_plain_header(ws: Worksheet, headers: tuple[str, ...]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.font = PLAIN_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.number_format = FORMAT_TEXT
    ws.row_dimensions[1].height = 18
    last = get_column_letter(len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last}{max(ws.max_row or 1, 1)}"


def apply_named_widths(ws: Worksheet, headers: tuple[str, ...], widths: dict[str, float]) -> None:
    for index, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(name, 12)


def write_inventory(ws: Worksheet, rows: list[dict[str, object]]) -> None:
    ws.title = "Inventory"
    for col, name in enumerate(PLAIN_INVENTORY, start=1):
        ws.cell(1, col, name)
    for row_index, item in enumerate(rows, start=2):
        for col, name in enumerate(PLAIN_INVENTORY, start=1):
            cell = ws.cell(row_index, col, item.get(name))
            cell.font = PLAIN_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=name in {"blurb", "color_pieces"})
            if name in NUMBER_HEADERS and isinstance(item.get(name), (int, float)):
                cell.number_format = "0.##"
            elif name in {"created_at", "updated_at", "ma", "photo_link", "source_link"}:
                cell.number_format = FORMAT_TEXT
    style_plain_header(ws, PLAIN_INVENTORY)
    last_row = max(len(rows) + 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PLAIN_INVENTORY))}{last_row}"
    apply_named_widths(ws, PLAIN_INVENTORY, INVENTORY_WIDTHS)


def write_blank_template(
    ws: Worksheet,
    title: str,
    headers: tuple[str, ...],
    widths: dict[str, float],
    dropdowns: list[tuple[str, tuple[str, ...]]],
) -> None:
    ws.title = title
    for col, name in enumerate(headers, start=1):
        ws.cell(1, col, name)
    for row in range(2, PLAIN_TEMPLATE_ROWS + 2):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col, None)
            cell.font = PLAIN_FONT
            cell.alignment = Alignment(vertical="center")
    style_plain_header(ws, headers)
    last = PLAIN_TEMPLATE_ROWS + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    apply_named_widths(ws, headers, widths)
    for header, choices in dropdowns:
        add_list_dropdown(
            ws,
            header_index(headers, header) + 1,
            choices,
            start_row=2,
            end_row=last,
            prompt=header,
        )


def new_plain_workbook() -> WorkbookType:
    workbook = Workbook()
    workbook.properties.creator = "Sassy Closet excel-kit"
    workbook.properties.title = "Sassy Closet Data"
    return workbook


def build_workbook(rows: list[dict[str, object]]) -> WorkbookType:
    workbook = new_plain_workbook()
    inventory = workbook.active
    write_inventory(inventory, rows)
    candidates = workbook.create_sheet("Candidates")
    write_blank_template(
        candidates,
        "Candidates",
        PLAIN_CANDIDATES,
        CANDIDATES_WIDTHS,
        [("status", PLAIN_CANDIDATE_STATUS)],
    )
    orders = workbook.create_sheet("Orders")
    write_blank_template(
        orders,
        "Orders",
        PLAIN_ORDERS,
        ORDERS_WIDTHS,
        [
            ("pay_method", PLAIN_ORDER_PAY),
            ("status", PLAIN_ORDER_STATUS),
        ],
    )
    order = list(PLAIN_SHEETS)
    for index, name in enumerate(order):
        workbook.move_sheet(name, offset=index - workbook.sheetnames.index(name))
    strip_workbook_images(workbook)
    return workbook


def unzip_test(path: Path) -> None:
    proc = subprocess.run(
        ["unzip", "-t", str(path)],
        check=False,
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        raise AssertionError(f"unzip -t failed:\n{proc.stdout}\n{proc.stderr}")
    if "No errors detected" not in proc.stdout:
        raise AssertionError(f"unzip -t did not confirm a clean zip:\n{proc.stdout}")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    media = [name for name in names if "/media/" in name.lower() or name.lower().endswith((".png", ".jpg", ".jpeg", ".gif"))]
    if media:
        raise AssertionError(f"embedded media in xlsx zip: {media}")


def _fill_rgb(fill: object) -> str | None:
    if fill is None:
        return None
    pattern = getattr(fill, "patternType", None)
    if pattern in (None, "none"):
        return None
    fg = getattr(fill, "fgColor", None)
    if fg is None:
        return None
    rgb = getattr(fg, "rgb", None)
    if not rgb:
        return None
    text = str(rgb).upper()
    if len(text) == 8:
        text = text[2:]
    return text


def inventory_values(ws: Worksheet, header: str) -> list[str]:
    idx = header_index(list(PLAIN_INVENTORY), header) + 1
    values: list[str] = []
    for row in range(2, (ws.max_row or 1) + 1):
        value = ws.cell(row, idx).value
        if value is None or str(value).strip() == "":
            continue
        values.append(str(value).strip())
    return values


def nonempty_row_count(ws: Worksheet, headers: tuple[str, ...], header_row: int = 1) -> int:
    count = 0
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        values = [ws.cell(row, col).value for col in range(1, len(headers) + 1)]
        if any(value is not None and str(value).strip() != "" for value in values):
            count += 1
    return count


def verify_plain_book(path: Path, *, expected_mas: Iterable[str] | None = None) -> list[str]:
    unzip_test(path)
    workbook = load_workbook(path)
    notes: list[str] = []
    try:
        if list(workbook.sheetnames) != list(PLAIN_SHEETS):
            raise AssertionError(f"sheets must be {list(PLAIN_SHEETS)}, got {workbook.sheetnames}")
        forbidden = {"Dashboard", "Bot_Activity", "Wishlist", "START HERE", "How_to_use"}
        extra = forbidden.intersection(workbook.sheetnames)
        if extra:
            raise AssertionError(f"retired/forbidden sheets present: {sorted(extra)}")

        inventory = workbook["Inventory"]
        headers = [inventory.cell(1, col).value for col in range(1, len(PLAIN_INVENTORY) + 1)]
        if headers != list(PLAIN_INVENTORY):
            raise AssertionError(f"Inventory headers drifted: {headers}")
        if inventory.freeze_panes != "A2":
            raise AssertionError(f"Inventory freeze_panes={inventory.freeze_panes}")
        if not inventory.auto_filter.ref or not str(inventory.auto_filter.ref).startswith("A1:"):
            raise AssertionError(f"Inventory autofilter={inventory.auto_filter.ref}")
        if inventory.merged_cells.ranges:
            raise AssertionError(f"Inventory merged cells forbidden: {inventory.merged_cells.ranges}")

        mas = inventory_values(inventory, "ma")
        if expected_mas is not None:
            got = list(mas)
            want = [str(ma).strip() for ma in expected_mas]
            if got != want:
                raise AssertionError(f"Inventory mã mismatch: {got} != {want}")
        leftover = [ma for ma in mas if ma.upper() in {code.upper() for code in PLAIN_RETIRED_MA}]
        if leftover:
            raise AssertionError(f"leftover retired mã {leftover} (A02 was renamed to P02)")

        links = inventory_values(inventory, "photo_link")
        for ma, link in zip(mas, inventory_values(inventory, "photo_link"), strict=False):
            want = photo_folder_link(ma)
            if link != want:
                raise AssertionError(f"photo_link for {ma} must be {want!r}, got {link!r}")
            if not link.endswith(f"Photos/{ma}/"):
                raise AssertionError(f"photo_link must end Photos/{ma}/, got {link!r}")
        if len(links) != len(mas):
            raise AssertionError("every Inventory row needs a photo_link")

        candidates = workbook["Candidates"]
        cand_headers = [candidates.cell(1, col).value for col in range(1, len(PLAIN_CANDIDATES) + 1)]
        if cand_headers != list(PLAIN_CANDIDATES):
            raise AssertionError(f"Candidates headers drifted: {cand_headers}")
        if candidates.freeze_panes != "A2":
            raise AssertionError(f"Candidates freeze_panes={candidates.freeze_panes}")
        if nonempty_row_count(candidates, PLAIN_CANDIDATES) != 0:
            raise AssertionError("Candidates must be a blank template (no invented finds)")
        if not candidates.data_validations.dataValidation:
            raise AssertionError("Candidates missing status dropdown")

        orders = workbook["Orders"]
        order_headers = [orders.cell(1, col).value for col in range(1, len(PLAIN_ORDERS) + 1)]
        if order_headers != list(PLAIN_ORDERS):
            raise AssertionError(f"Orders headers drifted: {order_headers}")
        if orders.freeze_panes != "A2":
            raise AssertionError(f"Orders freeze_panes={orders.freeze_panes}")
        if nonempty_row_count(orders, PLAIN_ORDERS) != 0:
            raise AssertionError("Orders must be a blank log (no invented customers)")
        if len(orders.data_validations.dataValidation) < 2:
            raise AssertionError("Orders needs pay_method + status dropdowns")

        for sheet in workbook.worksheets:
            if getattr(sheet, "_images", None):
                raise AssertionError(f"{sheet.title}: embedded images forbidden")
            if sheet.merged_cells.ranges:
                raise AssertionError(f"{sheet.title}: merged title art is retired")
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 1, 40), max_col=sheet.max_column or 1):
                for cell in row:
                    rgb = _fill_rgb(cell.fill)
                    if rgb and rgb in RETIRED_FILL_RGB:
                        raise AssertionError(f"{sheet.title}!{cell.coordinate} uses retired cute fill {rgb}")
                    font = cell.font
                    if font and font.size and font.size > 14:
                        raise AssertionError(f"{sheet.title}!{cell.coordinate} title-art font size {font.size}")

        notes.append(f"inventory_rows {len(mas)}")
        notes.append(f"mas {','.join(mas)}")
        notes.append(f"photo_root {ONEDRIVE_PHOTOS}/{{ma}}/")
    finally:
        workbook.close()
    return notes


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_rows(from_csv: Path | None, export_url: str) -> list[dict[str, object]]:
    if from_csv is not None:
        text = Path(from_csv).read_text(encoding="utf-8-sig")
        _headers, rows = read_export_text(text)
        return rows
    text = fetch_export_csv(export_url)
    _headers, rows = read_export_text(text)
    return rows


def build(out_dir: Path, *, from_csv: Path | None = None, export_url: str = PLAIN_EXPORT_URL) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / PLAIN_XLSX_NAME
    rows = load_rows(from_csv, export_url)
    retired = [
        str(row["ma"]).strip()
        for row in rows
        if str(row["ma"]).strip().upper() in {code.upper() for code in PLAIN_RETIRED_MA}
    ]
    if retired:
        raise ExportError(
            f"leftover retired mã {retired} in export. "
            "A02 was renamed to P02 (not P05). Never invent mã."
        )
    workbook = build_workbook(rows)
    workbook.save(path)
    workbook.close()
    notes = verify_plain_book(path, expected_mas=[str(row["ma"]) for row in rows])
    digest = sha256_file(path)
    print(f"wrote {path}")
    print(f"sha256 {digest}")
    for note in notes:
        print(note)
    if from_csv is None:
        print(f"source {export_url}")
    else:
        print(f"source csv {from_csv}")
    print(f"sync target: OneDrive {ONEDRIVE_PLAIN_DATA}")
    print("cute / pink / emoji / phone Excel is retired")
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=Path("out"),
        help="folder for Sassy_Closet_Data.xlsx (default: ./out)",
    )
    parser.add_argument(
        "--from-csv",
        type=Path,
        default=None,
        help="use a local export CSV instead of fetching the live site (tests)",
    )
    parser.add_argument(
        "--export-url",
        default=PLAIN_EXPORT_URL,
        help=f"live GET CSV (default: {PLAIN_EXPORT_URL})",
    )
    args = parser.parse_args(argv)
    try:
        build(args.out_dir, from_csv=args.from_csv, export_url=args.export_url)
    except (ExportError, AssertionError, OSError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
