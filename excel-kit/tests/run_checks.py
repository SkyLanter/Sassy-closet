#!/usr/bin/env python3
"""Kit checks: py_compile, CLI --help, no fake inventory in git.

Run from repo root:
  python3 excel-kit/tests/run_checks.py
"""

from __future__ import annotations

import csv
import hashlib
import py_compile
import subprocess
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
KIT = REPO / "excel-kit"
if str(KIT) not in sys.path:
    sys.path.insert(0, str(KIT))

from schema import (  # noqa: E402
    APPEND_OFFICIAL_ROW_EXPORTS,
    ASK_STOCK_MA,
    CANDIDATES,
    EXPORT_HEADERS,
    HUB_ALL,
    HUB_KIND_SHEETS,
    HUB_LINK_COLS,
    HUB_ORDERS,
    HUB_REQUIRED_BACKUP_COLS,
    HUB_SHEETS,
    HUB_XLSX_NAME,
    MA_LIST,
    OFFICIAL_TO_MA_LIST_STATUS,
    ONEDRIVE_FROM_GF,
    ONEDRIVE_SHOP_DIR,
    ONEDRIVE_HUB,
    ONEDRIVE_HUB_DIR,
    ONEDRIVE_HUB_PHOTOS,
    ORDER_STATUS,
    SOT_DASHBOARD_BRIEF_CELL,
    SOT_OFFICIAL_STATUS,
    SOT_WISHLIST,
    SQUARE_IMPORT_HEADERS,
    SQUARE_SOT_SHORT,
    STAY_OFF_SQUARE,
    MissingMaError,
    assert_hub_backup_columns,
    hub_photo_folder,
    is_hub_ma,
    looks_like_demo_row,
    official_status_or_raise,
    parse_ma,
    photo_filename_for_ma,
    require_ma,
    resolve_sheet_name,
)

SCRIPTS = [
    KIT / "schema.py",
    KIT / "clean_sot_demo.py",
    KIT / "build_boutique_desktop.py",
    KIT / "build_sassycloset_hub.py",
    KIT / "sot" / "workbook.py",
    KIT / "sot" / "append_official_row.py",
    KIT / "sot" / "append_wishlist_row.py",
    KIT / "sot" / "append_order_row.py",
    KIT / "sot" / "append_bot_activity.py",
    KIT / "sot" / "dashboard_brief.py",
    KIT / "sot" / "build_sot_desktop.py",
    KIT / "sot" / "gf_intake_apply.py",
    KIT / "sot" / "onedrive_from_gf_link.py",
    KIT / "square" / "validate_import.py",
    KIT / "tests" / "run_checks.py",
]

CLIS = [
    KIT / "clean_sot_demo.py",
    KIT / "build_boutique_desktop.py",
    KIT / "build_sassycloset_hub.py",
    KIT / "sot" / "append_official_row.py",
    KIT / "sot" / "append_wishlist_row.py",
    KIT / "sot" / "append_order_row.py",
    KIT / "sot" / "append_bot_activity.py",
    KIT / "sot" / "dashboard_brief.py",
    KIT / "sot" / "build_sot_desktop.py",
    KIT / "sot" / "gf_intake_apply.py",
    KIT / "sot" / "onedrive_from_gf_link.py",
    KIT / "square" / "validate_import.py",
]

FAKE_TOKENS = (
    "fake stock",
    "demo buyer",
    "test buyer",
    "ao999",
    "qu999",
    "va999",
    "~25 sample",
)


def _run(args: list[str], *, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        args,
        cwd=cwd or REPO,
        text=True,
        capture_output=True,
        check=False,
    )


def check_py_compile() -> None:
    for path in SCRIPTS:
        py_compile.compile(str(path), doraise=True)
    print(f"py_compile ok ({len(SCRIPTS)} files)")


def check_help() -> None:
    for path in CLIS:
        proc = _run([sys.executable, str(path), "--help"])
        if proc.returncode != 0:
            raise AssertionError(f"{path.name} --help failed:\n{proc.stderr}")
        if "usage:" not in proc.stdout.lower() and "usage:" not in proc.stderr.lower():
            raise AssertionError(f"{path.name} --help produced no usage")
        if path.name == "append_official_row.py":
            blob = proc.stdout + proc.stderr
            for status in SOT_OFFICIAL_STATUS:
                if status not in blob:
                    raise AssertionError(
                        f"append_official_row --help missing status {status!r}"
                    )
    print(f"--help ok ({len(CLIS)} CLIs)")


def check_schema_contract() -> None:
    assert MA_LIST[-1] == "photo_link"
    assert CANDIDATES[-1] == "photo_link"
    assert SOT_WISHLIST[-1] == "photo_link"
    assert "source" in SOT_WISHLIST
    assert ORDER_STATUS == (
        "Inquiry",
        "Reserved",
        "Paid",
        "Shipped",
        "Picked up",
        "Cancelled",
    )
    assert SOT_DASHBOARD_BRIEF_CELL == "B43"
    assert ONEDRIVE_FROM_GF.endswith("From GF")
    assert "Square" in STAY_OFF_SQUARE
    assert "Save" in STAY_OFF_SQUARE
    try:
        require_ma(None)
        raise AssertionError("require_ma(None) must fail")
    except MissingMaError as exc:
        assert "Ask Stock" in str(exc) or "ask Stock" in str(exc) or "ASK" in ASK_STOCK_MA
    assert require_ma("ao015") == "AO015"
    assert looks_like_demo_row(["AO999", "demo buyer"]) is True
    assert looks_like_demo_row(["", None, ""]) is False
    # append_official_row import contract — keep through A01 schema cutover.
    import schema as schema_mod

    missing = [name for name in APPEND_OFFICIAL_ROW_EXPORTS if not hasattr(schema_mod, name)]
    if missing:
        raise AssertionError(f"schema.py dropped append_official_row exports: {missing}")
    assert ASK_STOCK_MA
    assert "Square Free" in SQUARE_SOT_SHORT
    assert SOT_OFFICIAL_STATUS == (
        "Available",
        "Reserved",
        "Sold",
        "Hold",
        "Damaged",
        "Donated",
    )
    assert OFFICIAL_TO_MA_LIST_STATUS["Available"] == "in_stock"
    assert OFFICIAL_TO_MA_LIST_STATUS["Reserved"] == "held"
    assert OFFICIAL_TO_MA_LIST_STATUS["Sold"] == "sold"
    assert OFFICIAL_TO_MA_LIST_STATUS["Hold"] == "held"
    assert OFFICIAL_TO_MA_LIST_STATUS["Damaged"] == "archived"
    assert OFFICIAL_TO_MA_LIST_STATUS["Donated"] == "archived"
    assert official_status_or_raise("available") == "Available"
    try:
        official_status_or_raise("in_stock")
        raise AssertionError("Ma_List status must not pass official_status_or_raise")
    except ValueError:
        pass
    # Legacy AO001 parse_ma — do not treat A01 as valid here (hub mã is separate).
    assert parse_ma("AO001") == ("AO", 1)
    assert parse_ma("ao015") == ("AO", 15)
    assert parse_ma("A01") is None
    assert parse_ma("P02") is None
    assert photo_filename_for_ma("AO001") == "AO001.jpg"
    assert EXPORT_HEADERS[-1] == "photo_link"
    assert HUB_ALL[0] == "ma"
    assert HUB_ALL[1] == "source_link"
    assert HUB_ALL == (
        "ma",
        "source_link",
        "kind",
        "colors",
        "sell_usd",
        "cost",
        "currency",
        "square",
        "status",
        "flag",
        "next_desk",
        "photo_folder",
    )
    for name in HUB_REQUIRED_BACKUP_COLS:
        assert name in HUB_ALL
    for name in HUB_LINK_COLS:
        assert name in HUB_ALL
    assert_hub_backup_columns(HUB_ALL, "HUB_ALL")
    assert HUB_ORDERS == (
        "date",
        "ma",
        "customer",
        "pay",
        "amount_usd",
        "ship_or_meetup",
        "status",
        "notes",
    )
    assert HUB_SHEETS[0] == "All"
    assert HUB_SHEETS[-2:] == ("Orders", "Readme")
    assert len(HUB_SHEETS) == 14
    assert [name for name, _letter in HUB_KIND_SHEETS] == [
        "A_Ao",
        "Q_Quan",
        "V_Vay",
        "K_Khoac",
        "G_Giay",
        "B_Tui",
        "P_PhuKien",
        "S_Set",
        "O_Khac",
        "H_Toc",
        "J_TrangSuc",
    ]
    assert ONEDRIVE_HUB == f"{ONEDRIVE_HUB_DIR}/{HUB_XLSX_NAME}"
    assert ONEDRIVE_HUB_DIR == ONEDRIVE_SHOP_DIR
    assert ONEDRIVE_HUB_DIR == "Documents/Sassy Closet"
    assert ONEDRIVE_HUB == "Documents/Sassy Closet/sassycloset.xlsx"
    assert ONEDRIVE_HUB_PHOTOS == "Documents/Sassy Closet/Photos"
    assert ONEDRIVE_HUB_DIR != "Documents/sassycloset"
    assert is_hub_ma("A01")
    assert is_hub_ma("P02")
    assert is_hub_ma("P05")
    assert is_hub_ma("A100")
    assert is_hub_ma("A02") is True  # leftover shape; builder still refuses the row
    assert is_hub_ma("AO001") is False
    assert hub_photo_folder("P02") == f"{ONEDRIVE_HUB_PHOTOS}/P02/"
    assert hub_photo_folder("P05") == f"{ONEDRIVE_HUB_PHOTOS}/P05/"
    try:
        hub_photo_folder("")
        raise AssertionError("hub_photo_folder must refuse an empty mã")
    except ValueError:
        pass
    try:
        hub_photo_folder("SHIRT1")
        raise AssertionError("hub_photo_folder must refuse an invented mã")
    except ValueError:
        pass
    try:
        photo_filename_for_ma("not-a-ma")
        raise AssertionError("photo_filename_for_ma must refuse invented mã")
    except ValueError:
        pass
    assert resolve_sheet_name(["Official", "Wishlist"], "official") == "Official"
    assert resolve_sheet_name(["Ma_List"], "official") == "Ma_List"
    print("schema contract ok")


def check_no_fake_inventory_in_git() -> None:
    skip_suffix = {".md", ".py"}
    # Markdown/prompts may mention AO999 as a forbidden example — that is not a data row.
    # Scan committed non-doc payloads (csv / json / xlsx would be the leak).
    leaked: list[str] = []
    for path in REPO.rglob("*"):
        if not path.is_file():
            continue
        if ".git" in path.parts or "out" in path.parts or "__pycache__" in path.parts:
            continue
        if path.suffix.lower() == ".xlsx":
            leaked.append(f"committed xlsx (forbidden): {path.relative_to(REPO)}")
            continue
        if path.suffix.lower() not in {".csv", ".tsv", ".json"}:
            continue
        text = path.read_text(encoding="utf-8", errors="replace").lower()
        for token in FAKE_TOKENS:
            if token in text:
                leaked.append(f"{path.relative_to(REPO)} contains {token!r}")
    if leaked:
        raise AssertionError("fake inventory in git:\n  " + "\n  ".join(leaked))
    print("no fake inventory rows committed")


def check_square_template() -> None:
    path = KIT / "square" / "square_import_template.csv"
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        rows = [r for r in reader if any(c.strip() for c in r)]
    if len(rows) != 1:
        raise AssertionError(f"square template must be headers-only, got {len(rows)} rows")
    headers = rows[0]
    for required in ("Token", "Item Name", "Variation Name", "SKU", "Stockable"):
        if required not in headers:
            raise AssertionError(f"square template missing {required}")
    proc = _run([sys.executable, str(KIT / "square" / "validate_import.py"), str(path)])
    if proc.returncode != 0:
        raise AssertionError(proc.stdout + proc.stderr)
    if "headers-only" not in proc.stdout:
        raise AssertionError("validate_import should report headers-only")
    print("square template headers-only ok")


def _write_harness_export(path: Path) -> None:
    """Runtime-only CSV. Not committed. Uses live-shaped mã; never invents new ones."""
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(EXPORT_HEADERS))
        writer.writeheader()
        writer.writerow(
            {
                "ma": "A01",
                "kind": "A",
                "kind_vi": "Áo",
                "size": "S M L",
                "color": "Kem",
                "color_note": "",
                "color_pieces": "P1: Kem",
                "blurb": "harness row",
                "cost_cny": "10",
                "cost_usd": "1.5",
                "cost_currency": "USD",
                "sell_cny": "",
                "sell_usd": "20",
                "sell_currency": "USD",
                "source_link": "https://example.com/harness-a01",
                "status": "staged",
                "square": "not_square",
                "created_at": "2026-09-07T00:00:00.000Z",
                "updated_at": "2026-09-07T00:00:00.000Z",
                "photo_link": "replace-me",
            }
        )
        writer.writerow(
            {
                "ma": "P02",
                "kind": "P",
                "kind_vi": "Phụ kiện",
                "size": "",
                "color": "Đỏ",
                "color_note": "",
                "color_pieces": "P1: Đỏ",
                "blurb": "",
                "cost_cny": "63",
                "cost_usd": "9.39",
                "cost_currency": "CNY",
                "sell_cny": "",
                "sell_usd": "",
                "sell_currency": "",
                "source_link": "https://example.com/harness-p02",
                "status": "staged",
                "square": "not_square",
                "created_at": "2026-09-08T00:00:00.000Z",
                "updated_at": "2026-09-08T00:00:00.000Z",
                "photo_link": "also-replace-me",
            }
        )
        writer.writerow(
            {
                "ma": "P05",
                "kind": "P",
                "kind_vi": "Phụ kiện",
                "size": "",
                "color": "Hồng",
                "color_note": "",
                "color_pieces": "P1: Hồng",
                "blurb": "",
                "cost_cny": "75",
                "cost_usd": "11.18",
                "cost_currency": "CNY",
                "sell_cny": "",
                "sell_usd": "23",
                "sell_currency": "USD",
                "source_link": "https://example.com/harness-p05",
                "status": "staged",
                "square": "not_square",
                "created_at": "2026-09-08T00:00:00.000Z",
                "updated_at": "2026-09-08T00:00:00.000Z",
                "photo_link": "p05-replace",
            }
        )
        writer.writerow({"ma": "", "blurb": "empty mã must be skipped"})


def check_sassycloset_hub() -> None:
    prompt = KIT / "prompts" / "SASSYCLOSET_HUB_2026-09-07.md"
    kit_md = KIT / "KIT.md"
    photos_md = KIT / "PHOTOS.md"
    kit_sh = KIT / "kit.sh"
    root_kit = REPO / "kit.sh"
    builder = KIT / "build_sassycloset_hub.py"
    for path in (prompt, kit_md, photos_md, kit_sh, root_kit, builder):
        if not path.is_file():
            raise AssertionError(f"missing {path.relative_to(REPO)}")

    prompt_text = prompt.read_text(encoding="utf-8")
    for needle in (
        "sassycloset.xlsx",
        "api/export",
        "photo_folder",
        "A_Ao",
        "P_PhuKien",
        "J_TrangSuc",
        "Documents/Sassy Closet",
        "kit.sh save",
        "kit.sh run",
        "P02 and P05",
    ):
        if needle not in prompt_text:
            raise AssertionError(f"SASSYCLOSET_HUB_2026-09-07.md should keep spec line {needle!r}")
    if "offline backup" not in prompt_text.lower() or "source_link" not in prompt_text:
        raise AssertionError("hub spec must say Excel is an offline backup and keep source_link")

    kit_text = kit_md.read_text(encoding="utf-8")
    photos_text = photos_md.read_text(encoding="utf-8")
    if "sassycloset" not in kit_text.lower() or "teammates" not in kit_text.lower():
        raise AssertionError("KIT.md must name the sassycloset hub for teammates")
    if "kit.sh save" not in kit_text or "Documents/Sassy Closet" not in kit_text:
        raise AssertionError("KIT.md must document kit.sh save and the OneDrive land path")
    if "Documents/sassycloset" in kit_text:
        raise AssertionError("KIT.md must not land on Documents/sassycloset/")
    if "offline backup" not in kit_text.lower() or "source_link" not in kit_text:
        raise AssertionError("KIT.md must say Excel is the offline backup and keep source_link")
    if "Documents/Sassy Closet/Photos/{MA}/" not in photos_text:
        raise AssertionError("PHOTOS.md must use Documents/Sassy Closet/Photos/{MA}/")

    syntax = _run(["bash", "-n", str(kit_sh)])
    if syntax.returncode != 0:
        raise AssertionError(f"kit.sh bash -n failed:\n{syntax.stderr}")
    root_syntax = _run(["bash", "-n", str(root_kit)])
    if root_syntax.returncode != 0:
        raise AssertionError(f"root kit.sh bash -n failed:\n{root_syntax.stderr}")
    help_proc = _run(["bash", str(kit_sh), "--help"])
    if help_proc.returncode != 0:
        raise AssertionError(help_proc.stderr)
    help_blob = help_proc.stdout + help_proc.stderr
    if "save" not in help_blob or "run" not in help_blob:
        raise AssertionError("kit.sh --help should mention save and run")
    if "Documents/Sassy Closet" not in help_blob:
        raise AssertionError("kit.sh --help should document the OneDrive land path")
    if "Documents/sassycloset" in help_blob:
        raise AssertionError("kit.sh --help must not land on Documents/sassycloset/")
    cute = _run(["bash", str(kit_sh), "wishlist"])
    if cute.returncode == 0:
        raise AssertionError("kit.sh wishlist must be refused (cute pull retired)")
    if "retired" not in (cute.stdout + cute.stderr).lower():
        raise AssertionError("kit.sh wishlist refusal should say retired")

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        csv_path = out / "export.csv"
        _write_harness_export(csv_path)
        built = _run(
            [
                sys.executable,
                str(builder),
                "--from-csv",
                str(csv_path),
                "--skip-photos",
                "--out-dir",
                str(out),
            ]
        )
        if built.returncode != 0:
            raise AssertionError(built.stdout + built.stderr)
        book = out / HUB_XLSX_NAME
        if not book.is_file():
            raise AssertionError("hub builder did not write sassycloset.xlsx")
        blob = built.stdout + built.stderr
        if "sha256" not in blob:
            raise AssertionError("builder should print a checksum")
        if "Documents/Sassy Closet" not in blob:
            raise AssertionError("builder should print the OneDrive land path")
        if "Documents/sassycloset" in blob:
            raise AssertionError("builder must not land on Documents/sassycloset/")
        readme_txt = out / "README.txt"
        if not readme_txt.is_file():
            raise AssertionError("hub builder must write README.txt")
        readme_body = readme_txt.read_text(encoding="utf-8")
        if "Documents/Sassy Closet" not in readme_body:
            raise AssertionError("README.txt must document the OneDrive land path")
        if "Documents/sassycloset" in readme_body:
            raise AssertionError("README.txt must not land on Documents/sassycloset/")
        for ma in ("A01", "P02", "P05"):
            folder = out / "Photos" / ma
            if not folder.is_dir():
                raise AssertionError(f"skip-photos should still create Photos/{ma}/")
        if (out / "Photos" / "A02").exists():
            raise AssertionError("must not create an A02 photo folder")

        from openpyxl import load_workbook

        wb = load_workbook(book)
        try:
            if list(wb.sheetnames) != list(HUB_SHEETS):
                raise AssertionError(f"unexpected sheets: {wb.sheetnames}")
            if "Candidates" in wb.sheetnames or "Inventory" in wb.sheetnames or "Track" in wb.sheetnames:
                raise AssertionError("Candidates / Inventory / Track must not be hub sheets")
            all_sheet = wb["All"]
            headers = [all_sheet.cell(1, c).value for c in range(1, len(HUB_ALL) + 1)]
            if headers != list(HUB_ALL):
                raise AssertionError(f"All headers drifted: {headers}")
            if headers[0] != "ma" or headers[1] != "source_link":
                raise AssertionError(f"All must lead with ma + source_link, got {headers[:2]}")
            if "source_link" not in headers or "photo_folder" not in headers:
                raise AssertionError("never drop link columns")
            mas = []
            for row in range(2, (all_sheet.max_row or 1) + 1):
                value = all_sheet.cell(row, 1).value
                if value is None or str(value).strip() == "":
                    continue
                mas.append(str(value).strip())
            if mas != ["A01", "P02", "P05"]:
                raise AssertionError(f"harness All mã should be A01,P02,P05 (no A02), got {mas}")
            photo_col = list(HUB_ALL).index("photo_folder") + 1
            link_col = list(HUB_ALL).index("source_link") + 1
            cost_col = list(HUB_ALL).index("cost") + 1
            curr_col = list(HUB_ALL).index("currency") + 1
            flag_col = list(HUB_ALL).index("flag") + 1
            desk_col = list(HUB_ALL).index("next_desk") + 1
            colors_col = list(HUB_ALL).index("colors") + 1
            if all_sheet.cell(2, link_col).value != "https://example.com/harness-a01":
                raise AssertionError("source_link must copy the export Taobao/source URL")
            if all_sheet.cell(3, link_col).value != "https://example.com/harness-p02":
                raise AssertionError("P02 source_link must stay; never drop link columns")
            for row_index, ma in enumerate(mas, start=2):
                folder = str(all_sheet.cell(row_index, photo_col).value)
                if folder != f"{ONEDRIVE_HUB_PHOTOS}/{ma}/":
                    raise AssertionError(f"photo_folder for {ma} should be hub path, got {folder}")
            if all_sheet.cell(2, colors_col).value != "Kem":
                raise AssertionError("colors should come from export color, not invented")
            if all_sheet.cell(2, cost_col).value != 1.5 or str(all_sheet.cell(2, curr_col).value) != "USD":
                raise AssertionError("A01 cost should follow export cost_currency=USD")
            if all_sheet.cell(3, cost_col).value != 63 or str(all_sheet.cell(3, curr_col).value) != "CNY":
                raise AssertionError("P02 cost should follow export cost_currency=CNY")
            if all_sheet.cell(2, flag_col).value not in (None, ""):
                raise AssertionError("flag must stay empty unless export already has one")
            if all_sheet.cell(2, desk_col).value not in (None, ""):
                raise AssertionError("next_desk must stay empty by default")
            if any(cell.value == "A02" for row in all_sheet.iter_rows() for cell in row):
                raise AssertionError("no mã cell A02")
            if getattr(all_sheet, "_images", None):
                raise AssertionError("All embeds forbidden")
            if all_sheet.freeze_panes != "A2" or wb["Orders"].freeze_panes != "A2":
                raise AssertionError("All/Orders must freeze the header row")
            if wb["Q_Quan"].freeze_panes != "A2":
                raise AssertionError("empty category sheets still freeze the header")
            ao = [
                str(wb["A_Ao"].cell(row, 1).value).strip()
                for row in range(2, (wb["A_Ao"].max_row or 1) + 1)
                if wb["A_Ao"].cell(row, 1).value
            ]
            pk = [
                str(wb["P_PhuKien"].cell(row, 1).value).strip()
                for row in range(2, (wb["P_PhuKien"].max_row or 1) + 1)
                if wb["P_PhuKien"].cell(row, 1).value
            ]
            quan = [
                str(wb["Q_Quan"].cell(row, 1).value).strip()
                for row in range(2, (wb["Q_Quan"].max_row or 1) + 1)
                if wb["Q_Quan"].cell(row, 1).value
            ]
            if ao != ["A01"] or pk != ["P02", "P05"] or quan != []:
                raise AssertionError(f"category filters wrong: A_Ao={ao} P_PhuKien={pk} Q_Quan={quan}")
            readme_lines = [
                str(wb["Readme"].cell(r, 1).value).strip()
                for r in range(1, 8)
                if wb["Readme"].cell(r, 1).value
            ]
            if not (4 <= len(readme_lines) <= 6):
                raise AssertionError(f"Readme should be hub-rule lines, got {readme_lines}")
        finally:
            wb.close()

        leftover = out / "leftover.csv"
        with leftover.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(EXPORT_HEADERS))
            writer.writeheader()
            writer.writerow({"ma": "A02", "kind": "A", "kind_vi": "Áo", "status": "staged"})
        refuse = _run(
            [
                sys.executable,
                str(builder),
                "--from-csv",
                str(leftover),
                "--skip-photos",
                "--out-dir",
                str(out / "bad"),
            ]
        )
        if refuse.returncode == 0:
            raise AssertionError("builder must refuse leftover A02 (renamed to P02, not P05)")

    print("sassycloset hub ok")


def check_builders_and_append() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        sot = _run([sys.executable, str(KIT / "sot" / "build_sot_desktop.py"), "--out-dir", str(out)])
        if sot.returncode != 0:
            raise AssertionError(sot.stdout + sot.stderr)
        book = out / "Sassy_Closet_SoT.xlsx"
        if not book.is_file():
            raise AssertionError("SoT builder did not write Sassy_Closet_SoT.xlsx")

        desk = _run([sys.executable, str(KIT / "build_boutique_desktop.py"), "--out-dir", str(out)])
        if desk.returncode != 0:
            raise AssertionError(desk.stdout + desk.stderr)

        missing = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_official_row.py"),
                "-w",
                str(book),
                "--name-vi",
                "should fail",
            ]
        )
        if missing.returncode == 0:
            raise AssertionError("Official append without --ma must fail")
        err = missing.stderr + missing.stdout
        if "ma" not in err.lower() and "mã" not in err.lower():
            raise AssertionError(f"missing-mã error unclear:\n{err}")

        bad_ma = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_official_row.py"),
                "-w",
                str(book),
                "--ma",
                "SHIRT1",
            ]
        )
        if bad_ma.returncode == 0:
            raise AssertionError("invented / invalid mã must fail")

        ok = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_official_row.py"),
                "-w",
                str(book),
                "--ma",
                "AO001",
                "--name-vi",
                "Áo test-harness",
                "--size",
                "M",
                "--color",
                "đen",
                "--qty",
                "1",
                "--status",
                "Available",
            ]
        )
        if ok.returncode != 0:
            raise AssertionError(ok.stdout + ok.stderr)

        dup = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_official_row.py"),
                "-w",
                str(book),
                "--ma",
                "AO001",
                "--name-vi",
                "dup",
            ]
        )
        if dup.returncode == 0:
            raise AssertionError("duplicate mã must fail")

        wish_fail = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_wishlist_row.py"),
                "-w",
                str(book),
                "--what-vi",
                "no link",
            ]
        )
        if wish_fail.returncode == 0:
            raise AssertionError("wishlist without source/--no-source must fail")

        wish_ok = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_wishlist_row.py"),
                "-w",
                str(book),
                "--source",
                "https://item.taobao.com/item.htm?id=1",
                "--what-vi",
                "Váy harness",
                "--status",
                "candidate",
            ]
        )
        if wish_ok.returncode != 0:
            raise AssertionError(wish_ok.stdout + wish_ok.stderr)

        order_bad = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_order_row.py"),
                "-w",
                str(book),
                "--status",
                "Maybe",
            ]
        )
        if order_bad.returncode == 0:
            raise AssertionError("invalid order status must fail")

        order_ok = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_order_row.py"),
                "-w",
                str(book),
                "--status",
                "Inquiry",
                "--ma",
                "AO001",
                "--channel",
                "Facebook",
            ]
        )
        if order_ok.returncode != 0:
            raise AssertionError(order_ok.stdout + order_ok.stderr)

        bot_ok = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_bot_activity.py"),
                "-w",
                str(book),
                "--bot",
                "Stock",
                "--action",
                "draft",
                "--summary",
                "harness log",
                "--ma",
                "AO001",
            ]
        )
        if bot_ok.returncode != 0:
            raise AssertionError(bot_ok.stdout + bot_ok.stderr)

        # Lean Official desktop always has Bot_Activity; strip it to prove skip-note.
        from openpyxl import load_workbook

        skip_book = out / "no_bot.xlsx"
        wb = load_workbook(book)
        if "Bot_Activity" in wb.sheetnames:
            del wb["Bot_Activity"]
        wb.save(skip_book)
        wb.close()
        skip = _run(
            [
                sys.executable,
                str(KIT / "sot" / "append_bot_activity.py"),
                "-w",
                str(skip_book),
                "--bot",
                "Stock",
                "--action",
                "draft",
                "--summary",
                "should skip",
            ]
        )
        if skip.returncode != 0:
            raise AssertionError(skip.stdout + skip.stderr)
        if "skipped" not in (skip.stdout + skip.stderr).lower():
            raise AssertionError(f"expected skip note, got:\n{skip.stdout}\n{skip.stderr}")

        brief = _run([sys.executable, str(KIT / "sot" / "dashboard_brief.py"), "-w", str(book)])
        if brief.returncode != 0:
            raise AssertionError(brief.stdout + brief.stderr)
        blob = brief.stdout + brief.stderr
        if "B43" not in blob:
            raise AssertionError("dashboard_brief must mention B43")

        clean = _run(
            [sys.executable, str(KIT / "clean_sot_demo.py"), str(book), "--check"]
        )
        if clean.returncode != 0:
            raise AssertionError(clean.stdout + clean.stderr)

    print("builders + append CLIs ok")


GF_INTAKE = KIT / "sot" / "gf_intake_apply.py"
GF_CONTRACT = KIT / "prompts" / "GF_CLOTHES_INTAKE.md"
GF_HOW_TO = KIT / "templates" / "from_gf" / "HOW_TO_UPLOAD.md"
GF_TEMPLATE = KIT / "templates" / "from_gf" / "INTAKE_TEMPLATE.txt"
GF_CANDIDATE = KIT / "templates" / "from_gf" / "examples" / "candidate_looking_vay"
GF_BOUGHT = KIT / "templates" / "from_gf" / "examples" / "bought_waiting_stock"
GF_REPLIES = KIT / "inbox" / "gf_intake_reply_templates.md"
GF_WATCH = KIT / "sot" / "FROM_GF_WATCH.md"
GF_LINKING = KIT / "templates" / "from_gf" / "LINKING.md"
GF_LINK_CLI = KIT / "sot" / "onedrive_from_gf_link.py"


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _nonempty(book: Path, sheet: str, logical: str) -> tuple[list[str], list[str]]:
    from openpyxl import load_workbook

    from schema import header_row_for_sheet, read_headers, resolve_header_key

    wb = load_workbook(book)
    ws = wb[sheet]
    header_row = header_row_for_sheet(ws.title)
    headers = read_headers(ws, header_row)
    idx = resolve_header_key(headers, logical)
    values: list[str] = []
    if idx is not None:
        for row in range(header_row + 1, (ws.max_row or header_row) + 1):
            value = ws.cell(row, idx + 1).value
            if value is not None and str(value).strip() != "":
                values.append(str(value).strip())
    wb.close()
    return headers, values


def check_gf_intake() -> None:
    from sot.gf_intake_apply import (  # noqa: E402
        IntakeError,
        parse_intake_text,
        sanitize_ma_field,
    )

    for path in (
        GF_CONTRACT,
        GF_HOW_TO,
        GF_TEMPLATE,
        GF_REPLIES,
        GF_WATCH,
        GF_LINKING,
        GF_CANDIDATE / "INTAKE_TEMPLATE.txt",
        GF_BOUGHT / "note.txt",
    ):
        if not path.is_file():
            raise AssertionError(f"missing {path.relative_to(REPO)}")

    how = GF_HOW_TO.read_text(encoding="utf-8")
    for needle in (
        "From GF",
        "chưa lên Square",
        "share",
        "Messenger",
        "Dial Bot",
        "Microsoft Form",
    ):
        if needle not in how:
            raise AssertionError(f"HOW_TO_UPLOAD.md should mention {needle!r}")
    if "LINKING.md" not in how:
        raise AssertionError("HOW_TO_UPLOAD.md should point Boss/Kit at LINKING.md")

    replies = GF_REPLIES.read_text(encoding="utf-8").lower()
    if "vẫn chưa lên square" not in replies or "still off square" not in replies:
        raise AssertionError("reply templates must not claim Square stock (need off-Square wording in VI+EN)")

    parsed = parse_intake_text((GF_CANDIDATE / "INTAKE_TEMPLATE.txt").read_text(encoding="utf-8"))
    if parsed.get("kind") != "looking":
        raise AssertionError(f"candidate packet kind should be looking, got {parsed.get('kind')!r}")
    if "PLACEHOLDER" not in (parsed.get("source") or ""):
        raise AssertionError("candidate packet should keep a placeholder source URL, not a live shop")
    if parsed.get("ma"):
        raise AssertionError("candidate packet must not carry a live mã")
    if parsed.get("requested_by") != "GF":
        raise AssertionError("candidate packet requested_by should be GF")

    bought_parsed = parse_intake_text((GF_BOUGHT / "note.txt").read_text(encoding="utf-8"))
    if bought_parsed.get("kind") != "bought":
        raise AssertionError("bought packet kind should be bought")
    if bought_parsed.get("ma"):
        raise AssertionError("bought example must not invent a live mã")

    try:
        sanitize_ma_field("next")
        raise AssertionError("ma=next must be refused (never invent)")
    except IntakeError as exc:
        blob = str(exc).lower()
        if "invent" not in blob and "mã" not in str(exc) and "ma" not in blob:
            raise AssertionError(f"invent-mã error unclear: {exc}") from exc

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        built = _run([sys.executable, str(KIT / "sot" / "build_sot_desktop.py"), "--out-dir", str(out)])
        if built.returncode != 0:
            raise AssertionError(built.stdout + built.stderr)
        book = out / "Sassy_Closet_SoT.xlsx"
        photos = out / "Photos"
        before = _digest(book)

        dry = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--from-gf",
                str(GF_CANDIDATE),
                "-w",
                str(book),
                "--photos-out",
                str(photos),
                "--dry-run",
            ]
        )
        if dry.returncode != 0:
            raise AssertionError(dry.stdout + dry.stderr)
        if _digest(book) != before:
            raise AssertionError("dry-run must not write xlsx")
        if photos.exists() and any(photos.iterdir()):
            raise AssertionError("dry-run must not copy photos")
        dry_blob = dry.stdout + dry.stderr
        if "stay off square" not in dry_blob.lower():
            raise AssertionError("success path must print stay-off-Square reminder")

        refuse_ma = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--from-gf",
                str(GF_CANDIDATE),
                "-w",
                str(book),
                "--ma",
                "AO001",
                "--dry-run",
            ]
        )
        if refuse_ma.returncode == 0:
            raise AssertionError("looking packet with --ma must fail (no live mã on Wishlist)")
        if refuse_ma.returncode != 2:
            raise AssertionError(f"looking+ma should exit 2, got {refuse_ma.returncode}")

        invent = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--kind",
                "bought",
                "--no-source",
                "--what-vi",
                "Áo",
                "--ma",
                "next",
                "-w",
                str(book),
                "--dry-run",
            ]
        )
        if invent.returncode == 0:
            raise AssertionError("ma=next must fail — never invent mã")

        bad_link = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--kind",
                "looking",
                "--no-source",
                "--what-en",
                "bag",
                "--photo-link",
                "/tmp/From GF/secret.jpg",
                "-w",
                str(book),
                "--dry-run",
            ]
        )
        if bad_link.returncode == 0:
            raise AssertionError("local / From GF photo_link must be refused")

        root_refuse = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--from-gf",
                str(KIT / "templates" / "from_gf"),
                "-w",
                str(book),
                "--dry-run",
            ]
        )
        if root_refuse.returncode == 0:
            raise AssertionError("From GF inbox root must be refused (pass a packet folder)")

        live = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--from-gf",
                str(GF_CANDIDATE),
                "-w",
                str(book),
                "--photos-out",
                str(photos),
            ]
        )
        if live.returncode != 0:
            raise AssertionError(live.stdout + live.stderr)
        wish_headers, wish_what = _nonempty(book, "Wishlist", "what_vi")
        if wish_headers[-1] != "photo_link":
            raise AssertionError(f"Wishlist photo_link must stay last, got {wish_headers[-1]!r}")
        if not any("Váy hoa hồng" in v or "ví dụ" in v for v in wish_what):
            raise AssertionError(f"candidate row missing from Wishlist: {wish_what}")
        _, wish_status = _nonempty(book, "Wishlist", "status")
        if "candidate" not in wish_status:
            raise AssertionError(f"looking packet should be status=candidate, got {wish_status}")
        _, wish_links = _nonempty(book, "Wishlist", "photo_link")
        if any("from gf" in v.lower() or (v.startswith("/") and "http" not in v.lower()) for v in wish_links):
            raise AssertionError(f"photo_link must not be a From GF / local path: {wish_links}")
        copied = list(photos.glob("#*.jpg")) if photos.is_dir() else []
        if not copied:
            raise AssertionError("live candidate intake should copy/rename a wishlist photo to #NNN.jpg")
        _, official_ma = _nonempty(book, "Official", "ma")
        if official_ma:
            raise AssertionError(f"looking packet must not write Official mã: {official_ma}")

        bought = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--from-gf",
                str(GF_BOUGHT),
                "-w",
                str(book),
                "--photos-out",
                str(photos),
            ]
        )
        if bought.returncode != 0:
            raise AssertionError(bought.stdout + bought.stderr)
        blob = bought.stdout + bought.stderr
        if "ASK STOCK" not in blob and "Ask Stock" not in blob:
            raise AssertionError("bought without mã must ASK STOCK")
        _, statuses = _nonempty(book, "Wishlist", "status")
        if "bought" not in statuses:
            raise AssertionError(f"bought without mã should stay on Wishlist status=bought, got {statuses}")
        _, official_ma = _nonempty(book, "Official", "ma")
        if official_ma:
            raise AssertionError(
                f"bought without mã must not write Official (got {official_ma}). Never invent mã."
            )
        if "stay off square" not in blob.lower():
            raise AssertionError("bought success must print stay-off-Square reminder")

        linked = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--kind",
                "looking",
                "--source",
                "https://item.taobao.com/item.htm?id=PLACEHOLDER2",
                "--what-en",
                "flag-only scarf",
                "--photo-link",
                "https://1drv.ms/placeholder-photo",
                "-w",
                str(book),
            ]
        )
        if linked.returncode != 0:
            raise AssertionError(linked.stdout + linked.stderr)
        _, links = _nonempty(book, "Wishlist", "photo_link")
        if "https://1drv.ms/placeholder-photo" not in links:
            raise AssertionError(f"supplied photo_link should be stored, got {links}")

        official = _run(
            [
                sys.executable,
                str(GF_INTAKE),
                "--kind",
                "bought",
                "--source",
                "https://item.taobao.com/item.htm?id=PLACEHOLDER3",
                "--what-vi",
                "Áo harness Official",
                "--ma",
                "AO001",
                "--qty",
                "1",
                "-w",
                str(book),
                "--photos-out",
                str(photos),
            ]
        )
        if official.returncode != 0:
            raise AssertionError(official.stdout + official.stderr)
        _, official_ma = _nonempty(book, "Official", "ma")
        if "AO001" not in official_ma:
            raise AssertionError(f"bought with Stock mã should append Official, got {official_ma}")
        if "stay off square" not in (official.stdout + official.stderr).lower():
            raise AssertionError("Official staging success still prints stay-off-Square (no Save)")

    print("GF clothes intake ok")


def check_onedrive_from_gf_link() -> None:
    from sot.onedrive_from_gf_link import (  # noqa: E402
        EXAMPLE_BROKEN_WEB_URL,
        ONEDRIVE_FROM_GF as LINK_FROM_GF,
        build_plan,
        looks_like_broken_live_browse,
    )

    if LINK_FROM_GF != ONEDRIVE_FROM_GF:
        raise AssertionError("onedrive_from_gf_link.ONEDRIVE_FROM_GF must match schema")

    if not GF_LINKING.is_file():
        raise AssertionError("missing templates/from_gf/LINKING.md")
    linking = GF_LINKING.read_text(encoding="utf-8")
    for needle in (
        "onedrive.live.com?cid=",
        "%21",
        "createLink",
        "Can edit",
        "/?",
        "1drv.ms",
        "404",
    ):
        if needle not in linking:
            raise AssertionError(f"LINKING.md should warn/teach {needle!r}")
    watch = GF_WATCH.read_text(encoding="utf-8")
    if "LINKING.md" not in watch or "onedrive_from_gf_link.py" not in watch:
        raise AssertionError("FROM_GF_WATCH.md should point at LINKING.md and the helper")
    prompts = (KIT / "PROMPTS.md").read_text(encoding="utf-8")
    if "onedrive_from_gf_link.py" not in prompts or "LINKING.md" not in prompts:
        raise AssertionError("PROMPTS.md should point at LINKING.md / helper")

    if not looks_like_broken_live_browse(EXAMPLE_BROKEN_WEB_URL):
        raise AssertionError("Boss Graph webUrl must be detected as the failed browse shape")

    plan = build_plan(web_url=EXAMPLE_BROKEN_WEB_URL)
    want = (
        "https://onedrive.live.com/?cid=7a74d53e91d4f05f"
        "&id=7A74D53E91D4F05F%21sae89effd79104c01a63f092754d1c9f9"
    )
    if plan.canonical_browse != want:
        raise AssertionError(f"canonical browse mismatch: {plan.canonical_browse!r}")
    if "%21" not in (plan.canonical_browse or ""):
        raise AssertionError("canonical browse must encode ! as %21")
    if "onedrive.live.com?" in (plan.canonical_browse or "") and "onedrive.live.com/?" not in (
        plan.canonical_browse or ""
    ):
        raise AssertionError("canonical browse must include /? after .com")
    guid_want = "https://onedrive.live.com/?id=ae89effd-7910-4c01-a63f-092754d1c9f9&cid=7a74d53e91d4f05f"
    if plan.guid_browse != guid_want:
        raise AssertionError(f"GUID browse mismatch: {plan.guid_browse!r}")
    if "d.docs.live.net/7a74d53e91d4f05f/Documents/Sassy%20Closet/From%20GF" not in (
        plan.path_browse or ""
    ):
        raise AssertionError(f"path-style From GF URL missing: {plan.path_browse!r}")

    flags = build_plan(
        cid="7a74d53e91d4f05f",
        item_id="7A74D53E91D4F05F!sae89effd79104c01a63f092754d1c9f9",
    )
    if flags.canonical_browse != want:
        raise AssertionError("--cid/--item-id should match --web-url canonical form")

    already = build_plan(web_url=want)
    if already.canonical_browse != want:
        raise AssertionError("already-canonical URL should stay canonical")

    share = build_plan(web_url="https://1drv.ms/f/s!placeholder-do-not-invent")
    if share.guest_share != "https://1drv.ms/f/s!placeholder-do-not-invent":
        raise AssertionError("1drv.ms must pass through — never rewrite into a fake token")
    if share.canonical_browse and "1drv.ms" not in (share.guest_share or ""):
        raise AssertionError("guest share URL should remain the 1drv.ms input")

    help_proc = _run([sys.executable, str(GF_LINK_CLI), "--help"])
    if help_proc.returncode != 0:
        raise AssertionError(help_proc.stderr)
    help_blob = (help_proc.stdout + help_proc.stderr).lower()
    for needle in ("--cid", "--item-id", "--web-url", "usage:"):
        if needle not in help_blob:
            raise AssertionError(f"onedrive_from_gf_link.py --help should mention {needle!r}")

    cli = _run(
        [
            sys.executable,
            str(GF_LINK_CLI),
            "--web-url",
            EXAMPLE_BROKEN_WEB_URL,
        ]
    )
    if cli.returncode != 0:
        raise AssertionError(cli.stdout + cli.stderr)
    blob = cli.stdout + cli.stderr
    if want not in blob:
        raise AssertionError("CLI must print the canonical browse URL")
    if "createLink" not in blob:
        raise AssertionError("no Graph auth in repo — CLI must print createLink instructions")
    if "Can edit" not in blob:
        raise AssertionError("CLI must print Boss Share → Can edit steps")
    if "password" in blob.lower() and "do not" not in blob.lower():
        raise AssertionError("CLI must not ask for a password")

    missing = _run([sys.executable, str(GF_LINK_CLI)])
    if missing.returncode == 0:
        raise AssertionError("helper with no args should fail (need --cid/--item-id or --web-url)")

    print("From GF OneDrive link helper ok")


def main() -> int:
    try:
        check_py_compile()
        check_help()
        check_schema_contract()
        check_no_fake_inventory_in_git()
        check_square_template()
        check_sassycloset_hub()
        check_builders_and_append()
        check_gf_intake()
        check_onedrive_from_gf_link()
    except Exception as exc:  # noqa: BLE001 — kit runner prints and exits
        print(f"FAIL: {exc}", file=sys.stderr)
        return 1
    print("all excel-kit checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
