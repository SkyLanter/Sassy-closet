"""Builder smoke: unzip, openpyxl load, headers, freeze, empty rebuild."""

from __future__ import annotations

import hashlib
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from schema import (
    FINANCE_EXPENSES,
    FINANCE_FEES,
    FINANCE_PAYOUTS,
    FINANCE_SALES,
    FINANCE_SHEETS,
    FINANCE_TAX_SUMMARY_HEADERS,
    SQUARE_ON_HAND,
    SQUARE_SHEETS,
    SQUARE_SOLD_LOG,
    SQUARE_TEMPLATE_ROWS,
)
from build_square_finance import build, verify_finance_book, verify_square_book

KIT = Path(__file__).resolve().parents[1]
REPO = KIT.parent
BUILDER = KIT / "build_square_finance.py"
KIT_SH = KIT / "kit.sh"


def _headers(ws) -> list[object]:
    values = []
    col = 1
    while True:
        value = ws.cell(1, col).value
        if value in (None, ""):
            break
        values.append(value)
        col += 1
    return values


def test_builder_writes_empty_books(tmp_path: Path) -> None:
    square, finance = build(tmp_path)
    assert square.is_file()
    assert finance.is_file()
    verify_square_book(square)
    verify_finance_book(finance)


def test_cli_out_dir(tmp_path: Path) -> None:
    proc = subprocess.run(
        [sys.executable, str(BUILDER), "--out-dir", str(tmp_path)],
        check=False,
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "wrote" in proc.stdout
    assert (tmp_path / "Square.xlsx").is_file()
    assert (tmp_path / "Finance.xlsx").is_file()


def test_kit_sh_square(tmp_path: Path) -> None:
    proc = subprocess.run(
        ["bash", str(KIT_SH), "square", str(tmp_path)],
        check=False,
        capture_output=True,
        text=True,
        cwd=str(REPO),
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert (tmp_path / "Square.xlsx").is_file()
    assert (tmp_path / "Finance.xlsx").is_file()


def test_kit_sh_help_documents_books() -> None:
    proc = subprocess.run(
        ["bash", str(KIT_SH), "--help"],
        check=False,
        capture_output=True,
        text=True,
    )
    blob = proc.stdout + proc.stderr
    assert proc.returncode == 0
    for needle in ("square", "finance", "books", "Square.xlsx", "Finance.xlsx", "Documents/Sassy Closet"):
        assert needle in blob
    cute = subprocess.run(
        ["bash", str(KIT_SH), "cute"],
        check=False,
        capture_output=True,
        text=True,
    )
    assert cute.returncode == 2


def test_unzip_and_sheet_contract(tmp_path: Path) -> None:
    square, finance = build(tmp_path)
    for path in (square, finance):
        unzip = subprocess.run(
            ["unzip", "-t", str(path)],
            check=False,
            capture_output=True,
            text=True,
        )
        assert unzip.returncode == 0
        assert "No errors detected" in unzip.stdout

    sw = load_workbook(square)
    try:
        assert list(sw.sheetnames) == list(SQUARE_SHEETS)
        assert _headers(sw["On_Hand"]) == list(SQUARE_ON_HAND)
        assert _headers(sw["Sold_Log"]) == list(SQUARE_SOLD_LOG)
        assert sw["On_Hand"].freeze_panes == "A2"
        assert str(sw["On_Hand"].auto_filter.ref).startswith("A1:")
        assert sw["Sold_Log"].freeze_panes == "A2"
    finally:
        sw.close()

    fw = load_workbook(finance)
    try:
        assert list(fw.sheetnames) == list(FINANCE_SHEETS)
        assert _headers(fw["Sales"]) == list(FINANCE_SALES)
        assert _headers(fw["Fees"]) == list(FINANCE_FEES)
        assert _headers(fw["Payouts_Transfers"]) == list(FINANCE_PAYOUTS)
        assert _headers(fw["Expenses"]) == list(FINANCE_EXPENSES)
        assert _headers(fw["Tax_Summary"]) == list(FINANCE_TAX_SUMMARY_HEADERS)
        assert fw["Sales"].freeze_panes == "A2"
        assert str(fw["Sales"].auto_filter.ref).startswith("A1:")
        assert fw["Tax_Summary"].freeze_panes == "A2"
    finally:
        fw.close()


def test_idempotent_rebuild_does_not_invent_sales(tmp_path: Path) -> None:
    first = tmp_path / "a"
    second = tmp_path / "b"
    square_a, finance_a = build(first)
    square_b, finance_b = build(second)

    for path in (square_a, finance_a, square_b, finance_b):
        wb = load_workbook(path)
        try:
            for title in wb.sheetnames:
                if title in {"Readme", "Tax_Summary"}:
                    continue
                ws = wb[title]
                for row in ws.iter_rows(min_row=2, max_row=SQUARE_TEMPLATE_ROWS + 1):
                    for cell in row:
                        value = cell.value
                        if value is None or str(value).strip() == "":
                            continue
                        if isinstance(value, str) and value.startswith("="):
                            continue
                        raise AssertionError(f"{path.name} / {title} {cell.coordinate} invented {value!r}")
        finally:
            wb.close()

    # Structure is stable even if xlsx zip timestamps change the digest.
    wa = load_workbook(finance_a)
    wb = load_workbook(finance_b)
    try:
        assert wa.sheetnames == wb.sheetnames
        assert _headers(wa["Sales"]) == _headers(wb["Sales"])
    finally:
        wa.close()
        wb.close()


def test_verify_only(tmp_path: Path) -> None:
    square, finance = build(tmp_path)
    proc = subprocess.run(
        [sys.executable, str(BUILDER), "--verify-only", str(square), str(finance)],
        check=False,
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr


def test_rebuild_hashes_are_files(tmp_path: Path) -> None:
    square, finance = build(tmp_path)
    assert len(hashlib.sha256(square.read_bytes()).hexdigest()) == 64
    assert len(hashlib.sha256(finance.read_bytes()).hexdigest()) == 64


@pytest.mark.parametrize("alias", ["books", "finance"])
def test_kit_aliases(tmp_path: Path, alias: str) -> None:
    dest = tmp_path / alias
    dest.mkdir()
    proc = subprocess.run(
        ["bash", str(KIT_SH), alias, str(dest)],
        check=False,
        capture_output=True,
        text=True,
        cwd=str(REPO),
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert (dest / "Finance.xlsx").is_file()
