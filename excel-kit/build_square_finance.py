#!/usr/bin/env python3
"""Build empty Square.xlsx + Finance.xlsx (bought tracker + tax-ready).

On_Hand starts empty — staged site mãs are not bought.
Finance starts empty — no invented sales or $.
San Jose, CA combined sales tax is 10.000% (CDTFA, effective Apr 1, 2026).
Plain freeze + AutoFilter. No cute / pink / embeds.
Square Free remains on-hand inventory truth. No Square Save.

Usage:
  python3 excel-kit/build_square_finance.py --out-dir ./out
"""

from __future__ import annotations

import argparse
import hashlib
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (  # noqa: E402
    FINANCE_CHANNEL,
    FINANCE_EXPENSE_CATEGORY,
    FINANCE_EXPENSE_PAY,
    FINANCE_EXPENSES,
    FINANCE_FEE_SOURCE,
    FINANCE_FEE_TYPE,
    FINANCE_FEES,
    FINANCE_FROM_METHOD,
    FINANCE_PAY_METHOD,
    FINANCE_PAYOUT_KIND,
    FINANCE_PAYOUTS,
    FINANCE_README_LINES,
    FINANCE_SALES,
    FINANCE_SHEETS,
    FINANCE_TAX_CATEGORY,
    FINANCE_TAX_SUMMARY_HEADERS,
    FINANCE_TAX_SUMMARY_METRICS,
    FINANCE_TAX_SUMMARY_NOTE_ROWS,
    FINANCE_TAX_SUMMARY_ROWS,
    FINANCE_TAX_YEAR,
    FINANCE_XLSX_NAME,
    ONEDRIVE_FINANCE,
    ONEDRIVE_PHOTOS,
    ONEDRIVE_SQUARE,
    SAN_JOSE_JURISDICTION,
    SAN_JOSE_SALES_TAX_RATE,
    SAN_JOSE_SALES_TAX_RATE_DISPLAY,
    SQUARE_COST_CURRENCY,
    SQUARE_FORMULA_LAST_ROW,
    SQUARE_ON_HAND,
    SQUARE_ON_HAND_STATUS,
    SQUARE_README_LINES,
    SQUARE_SHEETS,
    SQUARE_SOLD_LOG,
    SQUARE_TEMPLATE_ROWS,
    SQUARE_TRACK_ON,
    SQUARE_XLSX_NAME,
    add_list_dropdown,
    finance_metric_month_formula,
    finance_net_usd_formula,
    finance_sales_tax_rate_formula,
    finance_sales_tax_usd_formula,
    finance_taxable_base_formula,
    finance_ytd_formula,
    header_index,
    square_photo_folder_formula,
    strip_workbook_images,
)

PLAIN_FONT = Font(name="Calibri", size=11)
PLAIN_HEADER_FONT = Font(name="Calibri", size=11, bold=True)
CUTE_FILL_RGB = frozenset({"F7C6D5", "C43B6E", "FFF7FA", "F3E6EE", "FFF3B0"})
CUTE_EMOJI = "🎀♡💗✨"
FORBIDDEN_SHEETS = frozenset(
    {
        "Dashboard",
        "Bot_Activity",
        "Wishlist",
        "START HERE",
        "How_to_use",
        "Candidates",
        "Inventory",
        "Track",
        "All",
    }
)
# Word-boundary so "Shipping" is not treated as a PIN.
SECRET_RE = re.compile(r"\b(password|passwd|pin|otp|routing number|ssn)\b", re.I)

ON_HAND_WIDTHS = {
    "ma": 10,
    "kind": 8,
    "colors": 18,
    "size": 12,
    "qty_on_hand": 12,
    "cost_cny": 11,
    "cost_usd": 11,
    "cost_currency": 13,
    "buy_date": 12,
    "source_link": 36,
    "photo_folder": 42,
    "square_item_name": 20,
    "track_on": 10,
    "status": 12,
    "sold_date": 12,
    "notes": 24,
}
SOLD_LOG_WIDTHS = {
    "ma": 10,
    "sold_date": 12,
    "qty": 8,
    "finance_ref": 16,
    "notes": 24,
}
SALES_WIDTHS = {
    "date": 12,
    "ma": 10,
    "description": 22,
    "qty": 8,
    "gross_usd": 11,
    "ship_usd": 11,
    "discount_usd": 13,
    "net_usd": 11,
    "pay_method": 14,
    "pay_ref": 16,
    "customer_note": 18,
    "channel": 12,
    "square_xlsx_ma": 16,
    "tax_category": 14,
    "cogs_usd": 11,
    "taxable_base_usd": 14,
    "sales_tax_rate": 14,
    "sales_tax_usd": 13,
    "notes": 24,
}
FEES_WIDTHS = {
    "date": 12,
    "source": 14,
    "amount_usd": 12,
    "fee_type": 16,
    "related_sale_ref": 16,
    "notes": 24,
}
PAYOUTS_WIDTHS = {
    "date": 12,
    "from_method": 16,
    "to_account_note": 20,
    "amount_usd": 12,
    "confirmation": 16,
    "notes": 24,
}
EXPENSES_WIDTHS = {
    "date": 12,
    "vendor": 18,
    "category": 16,
    "amount_usd": 12,
    "payment_method": 16,
    "receipt_note": 18,
    "notes": 24,
}
MONEY_HEADERS = frozenset(
    {
        "qty_on_hand",
        "cost_cny",
        "cost_usd",
        "qty",
        "gross_usd",
        "ship_usd",
        "discount_usd",
        "net_usd",
        "amount_usd",
        "cogs_usd",
        "taxable_base_usd",
        "sales_tax_usd",
    }
)
RATE_HEADERS = frozenset({"sales_tax_rate"})
DATE_HEADERS = frozenset({"buy_date", "sold_date", "date"})
FORMULA_HEADERS = frozenset(
    {
        "net_usd",
        "taxable_base_usd",
        "sales_tax_rate",
        "sales_tax_usd",
        "photo_folder",
    }
)


def style_plain_header(ws: Worksheet, headers: tuple[str, ...], last_row: int) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.font = PLAIN_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.number_format = FORMAT_TEXT
    ws.row_dimensions[1].height = 18
    last = get_column_letter(len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last}{max(last_row, 1)}"


def apply_named_widths(ws: Worksheet, headers: tuple[str, ...], widths: dict[str, float]) -> None:
    for index, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(name, 12)


def write_empty_grid(ws: Worksheet, headers: tuple[str, ...], rows: int) -> None:
    for row in range(2, rows + 2):
        for col, name in enumerate(headers, start=1):
            cell = ws.cell(row, col, None)
            cell.font = PLAIN_FONT
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=name in {"notes", "description", "source_link", "photo_folder", "receipt_note"},
            )
            if name in DATE_HEADERS:
                cell.number_format = "YYYY-MM-DD"
            elif name in RATE_HEADERS:
                cell.number_format = "0.00%"
            elif name in MONEY_HEADERS:
                cell.number_format = "0.00"
            else:
                cell.number_format = FORMAT_TEXT


def write_readme(ws: Worksheet, title: str, lines: tuple[str, ...]) -> None:
    ws.title = title
    ws.column_dimensions["A"].width = 108
    for row, line in enumerate(lines, start=1):
        if any(mark in line for mark in CUTE_EMOJI):
            raise AssertionError(f"Readme must stay plain, got cute mark in {line!r}")
        cell = ws.cell(row, 1, line)
        cell.font = PLAIN_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 32


def new_plain_workbook(title: str) -> WorkbookType:
    workbook = Workbook()
    workbook.properties.creator = "Sassy Closet excel-kit"
    workbook.properties.title = title
    return workbook


def _dropdown(
    ws: Worksheet,
    headers: tuple[str, ...],
    header: str,
    choices: tuple[str, ...],
    last_row: int,
) -> None:
    add_list_dropdown(
        ws,
        header_index(headers, header) + 1,
        choices,
        start_row=2,
        end_row=last_row,
        prompt=header,
    )


def _write_formula_column(
    ws: Worksheet,
    headers: tuple[str, ...],
    header: str,
    last_row: int,
    formula_for_row,
    number_format: str,
) -> None:
    col = header_index(headers, header) + 1
    for row in range(2, last_row + 1):
        cell = ws.cell(row, col, formula_for_row(row))
        cell.font = PLAIN_FONT
        cell.number_format = number_format
        cell.alignment = Alignment(vertical="center")


def write_on_hand(ws: Worksheet) -> None:
    ws.title = "On_Hand"
    last_row = SQUARE_TEMPLATE_ROWS + 1
    write_empty_grid(ws, SQUARE_ON_HAND, SQUARE_TEMPLATE_ROWS)
    photo_col = header_index(SQUARE_ON_HAND, "photo_folder") + 1
    for row in range(2, last_row + 1):
        cell = ws.cell(row, photo_col, square_photo_folder_formula(row))
        cell.font = PLAIN_FONT
        cell.number_format = FORMAT_TEXT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    style_plain_header(ws, SQUARE_ON_HAND, last_row)
    apply_named_widths(ws, SQUARE_ON_HAND, ON_HAND_WIDTHS)
    _dropdown(ws, SQUARE_ON_HAND, "cost_currency", SQUARE_COST_CURRENCY, last_row)
    _dropdown(ws, SQUARE_ON_HAND, "track_on", SQUARE_TRACK_ON, last_row)
    _dropdown(ws, SQUARE_ON_HAND, "status", SQUARE_ON_HAND_STATUS, last_row)


def write_sold_log(ws: Worksheet) -> None:
    ws.title = "Sold_Log"
    last_row = SQUARE_TEMPLATE_ROWS + 1
    write_empty_grid(ws, SQUARE_SOLD_LOG, SQUARE_TEMPLATE_ROWS)
    style_plain_header(ws, SQUARE_SOLD_LOG, last_row)
    apply_named_widths(ws, SQUARE_SOLD_LOG, SOLD_LOG_WIDTHS)


def build_square_book() -> WorkbookType:
    workbook = new_plain_workbook("Square")
    write_on_hand(workbook.active)
    write_sold_log(workbook.create_sheet("Sold_Log"))
    write_readme(workbook.create_sheet("Readme"), "Readme", SQUARE_README_LINES)
    for index, name in enumerate(SQUARE_SHEETS):
        workbook.move_sheet(name, offset=index - workbook.sheetnames.index(name))
    strip_workbook_images(workbook)
    return workbook


def write_sales(ws: Worksheet) -> None:
    ws.title = "Sales"
    last_row = SQUARE_TEMPLATE_ROWS + 1
    write_empty_grid(ws, FINANCE_SALES, SQUARE_TEMPLATE_ROWS)
    _write_formula_column(ws, FINANCE_SALES, "net_usd", last_row, finance_net_usd_formula, "0.00")
    _write_formula_column(
        ws, FINANCE_SALES, "taxable_base_usd", last_row, finance_taxable_base_formula, "0.00"
    )
    _write_formula_column(
        ws, FINANCE_SALES, "sales_tax_rate", last_row, finance_sales_tax_rate_formula, "0.00%"
    )
    _write_formula_column(
        ws, FINANCE_SALES, "sales_tax_usd", last_row, finance_sales_tax_usd_formula, "0.00"
    )
    style_plain_header(ws, FINANCE_SALES, last_row)
    apply_named_widths(ws, FINANCE_SALES, SALES_WIDTHS)
    _dropdown(ws, FINANCE_SALES, "pay_method", FINANCE_PAY_METHOD, last_row)
    _dropdown(ws, FINANCE_SALES, "channel", FINANCE_CHANNEL, last_row)
    _dropdown(ws, FINANCE_SALES, "tax_category", FINANCE_TAX_CATEGORY, last_row)


def write_fees(ws: Worksheet) -> None:
    ws.title = "Fees"
    last_row = SQUARE_TEMPLATE_ROWS + 1
    write_empty_grid(ws, FINANCE_FEES, SQUARE_TEMPLATE_ROWS)
    style_plain_header(ws, FINANCE_FEES, last_row)
    apply_named_widths(ws, FINANCE_FEES, FEES_WIDTHS)
    _dropdown(ws, FINANCE_FEES, "source", FINANCE_FEE_SOURCE, last_row)
    _dropdown(ws, FINANCE_FEES, "fee_type", FINANCE_FEE_TYPE, last_row)


def write_payouts(ws: Worksheet) -> None:
    ws.title = "Payouts_Transfers"
    last_row = SQUARE_TEMPLATE_ROWS + 1
    write_empty_grid(ws, FINANCE_PAYOUTS, SQUARE_TEMPLATE_ROWS)
    style_plain_header(ws, FINANCE_PAYOUTS, last_row)
    apply_named_widths(ws, FINANCE_PAYOUTS, PAYOUTS_WIDTHS)
    _dropdown(ws, FINANCE_PAYOUTS, "from_method", FINANCE_FROM_METHOD, last_row)


def write_expenses(ws: Worksheet) -> None:
    ws.title = "Expenses"
    last_row = SQUARE_TEMPLATE_ROWS + 1
    write_empty_grid(ws, FINANCE_EXPENSES, SQUARE_TEMPLATE_ROWS)
    style_plain_header(ws, FINANCE_EXPENSES, last_row)
    apply_named_widths(ws, FINANCE_EXPENSES, EXPENSES_WIDTHS)
    _dropdown(ws, FINANCE_EXPENSES, "category", FINANCE_EXPENSE_CATEGORY, last_row)
    _dropdown(ws, FINANCE_EXPENSES, "payment_method", FINANCE_EXPENSE_PAY, last_row)


def write_tax_summary(ws: Worksheet) -> None:
    ws.title = "Tax_Summary"
    headers = FINANCE_TAX_SUMMARY_HEADERS
    last_row = 1 + len(FINANCE_TAX_SUMMARY_METRICS) + len(FINANCE_TAX_SUMMARY_NOTE_ROWS)
    style_plain_header(ws, headers, last_row)
    ws.column_dimensions["A"].width = 64
    for month in range(1, 13):
        ws.column_dimensions[get_column_letter(month + 1)].width = 12
    ws.column_dimensions["N"].width = 12
    ws.column_dimensions["O"].width = 56

    notes = {
        "gross": "US Schedule C-style rollup for an accountant. Not legal advice.",
        "shipping": "Customer ship $ only when paid. Never invent a flat ship rate.",
        "sales_tax": (
            f"{SAN_JOSE_JURISDICTION} {SAN_JOSE_SALES_TAX_RATE_DISPLAY} on taxable sell price. "
            "Liability, not income. Do not use 9.375% as current."
        ),
        "cogs": "Copy Square.xlsx On_Hand cost_usd onto Sales cogs_usd when the piece sells. Do not also re-enter the same dollar as Expenses inventory_cogs.",
        "profit": "Income tax is on profit, not COGS. Sales tax is excluded from this line.",
    }
    for index, (label, kind) in enumerate(FINANCE_TAX_SUMMARY_METRICS, start=2):
        left = ws.cell(index, 1, label)
        left.font = PLAIN_FONT
        left.alignment = Alignment(vertical="center", wrap_text=True)
        for month in range(1, 13):
            formula = finance_metric_month_formula(kind, FINANCE_TAX_YEAR, month)
            cell = ws.cell(index, month + 1, formula)
            cell.font = PLAIN_FONT
            cell.number_format = "0.00"
            cell.alignment = Alignment(vertical="center")
        ytd = ws.cell(index, 14, finance_ytd_formula(index))
        ytd.font = PLAIN_FONT
        ytd.number_format = "0.00"
        note = ws.cell(index, 15, notes.get(kind))
        note.font = PLAIN_FONT
        note.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[index].height = 22

    note_start = 2 + len(FINANCE_TAX_SUMMARY_METRICS)
    for offset, (label, note_text) in enumerate(FINANCE_TAX_SUMMARY_NOTE_ROWS):
        row = note_start + offset
        left = ws.cell(row, 1, label)
        left.font = PLAIN_FONT
        left.alignment = Alignment(vertical="center", wrap_text=True)
        note = ws.cell(row, 15, note_text)
        note.font = PLAIN_FONT
        note.alignment = Alignment(wrap_text=True, vertical="center")
        if label == "Combined sales tax rate":
            rate_cell = ws.cell(row, 14, float(SAN_JOSE_SALES_TAX_RATE))
            rate_cell.font = PLAIN_FONT
            rate_cell.number_format = "0.00%"
        ws.row_dimensions[row].height = 22


def build_finance_book() -> WorkbookType:
    workbook = new_plain_workbook("Finance")
    write_sales(workbook.active)
    write_fees(workbook.create_sheet("Fees"))
    write_payouts(workbook.create_sheet("Payouts_Transfers"))
    write_expenses(workbook.create_sheet("Expenses"))
    write_tax_summary(workbook.create_sheet("Tax_Summary"))
    write_readme(workbook.create_sheet("Readme"), "Readme", FINANCE_README_LINES)
    for index, name in enumerate(FINANCE_SHEETS):
        workbook.move_sheet(name, offset=index - workbook.sheetnames.index(name))
    strip_workbook_images(workbook)
    return workbook


def unzip_test(path: Path) -> None:
    proc = subprocess.run(
        ["unzip", "-t", str(path)],
        check=False,
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        raise AssertionError(f"unzip -t failed:\n{proc.stdout}\n{proc.stderr}")
    if "No errors detected" not in proc.stdout:
        raise AssertionError(f"unzip -t did not confirm a clean zip:\n{proc.stdout}")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    media = [
        name
        for name in names
        if "/media/" in name.lower() or name.lower().endswith((".png", ".jpg", ".jpeg", ".gif"))
    ]
    if media:
        raise AssertionError(f"embedded media in xlsx zip: {media}")


def _fill_rgb(fill: object) -> str | None:
    if fill is None:
        return None
    if getattr(fill, "patternType", None) in (None, "none"):
        return None
    fg = getattr(fill, "fgColor", None)
    if fg is None:
        return None
    rgb = getattr(fg, "rgb", None)
    if not rgb:
        return None
    text = str(rgb).upper()
    if len(text) == 8:
        text = text[2:]
    return text


def _assert_no_secrets(ws: Worksheet) -> None:
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, max_col=ws.max_column or 1):
        for cell in row:
            text = "" if cell.value is None else str(cell.value)
            if SECRET_RE.search(text):
                if "never store bank passwords" in text.lower() or "never store" in text.lower():
                    continue
                raise AssertionError(
                    f"{ws.title} {cell.coordinate} looks like a secret field: {cell.value!r}"
                )


def _assert_plain_data_sheet(ws: Worksheet, headers: tuple[str, ...], title: str) -> None:
    if ws.title != title:
        raise AssertionError(f"sheet title {ws.title!r} != {title!r}")
    got = [ws.cell(1, col).value for col in range(1, len(headers) + 1)]
    if got != list(headers):
        raise AssertionError(f"{title} headers drifted: {got}")
    extra = ws.cell(1, len(headers) + 1).value
    if extra not in (None, ""):
        raise AssertionError(f"{title} extra header {extra!r}")
    if ws.freeze_panes != "A2":
        raise AssertionError(f"{title} freeze_panes={ws.freeze_panes}")
    if not ws.auto_filter.ref or not str(ws.auto_filter.ref).startswith("A1:"):
        raise AssertionError(f"{title} autofilter={ws.auto_filter.ref}")
    if ws.merged_cells.ranges:
        raise AssertionError(f"{title} merged cells forbidden: {ws.merged_cells.ranges}")
    scan_rows = min(ws.max_row or 1, SQUARE_TEMPLATE_ROWS + 1)
    for row in ws.iter_rows(min_row=1, max_row=scan_rows, max_col=len(headers)):
        for cell in row:
            rgb = _fill_rgb(cell.fill)
            if rgb in CUTE_FILL_RGB:
                raise AssertionError(f"{title} {cell.coordinate}: cute fill {rgb}")
            value = "" if cell.value is None else str(cell.value)
            if any(mark in value for mark in CUTE_EMOJI):
                raise AssertionError(f"{title} {cell.coordinate}: cute mark {value!r}")
    _assert_no_secrets(ws)


def _typed_values(ws: Worksheet, headers: tuple[str, ...], header: str) -> list[str]:
    idx = header_index(list(headers), header) + 1
    values: list[str] = []
    for row in range(2, (ws.max_row or 1) + 1):
        value = ws.cell(row, idx).value
        if value is None or str(value).strip() == "":
            continue
        text = str(value).strip()
        if text.startswith("="):
            continue
        values.append(text)
    return values


def _assert_no_invented_ma(ws: Worksheet, headers: tuple[str, ...], title: str) -> None:
    for header in ("ma", "square_xlsx_ma", "related_sale_ref", "related_ma", "finance_ref"):
        if header not in headers:
            continue
        found = _typed_values(ws, headers, header)
        if found:
            raise AssertionError(f"{title} must ship empty of mã, found {found!r}")


def _assert_no_invented_money(ws: Worksheet, headers: tuple[str, ...], title: str) -> None:
    money = [name for name in headers if name in MONEY_HEADERS and name not in FORMULA_HEADERS]
    for header in money:
        found = _typed_values(ws, headers, header)
        if found:
            raise AssertionError(f"{title} must ship empty of $, found {header}={found!r}")


def _assert_no_workbook_images(workbook: WorkbookType, label: str) -> None:
    if getattr(workbook, "_images", None):
        raise AssertionError(f"{label}: workbook-level images forbidden")
    for ws in workbook.worksheets:
        if getattr(ws, "_images", None):
            raise AssertionError(f"{label} / {ws.title}: embedded images forbidden")


def verify_square_book(path: Path) -> None:
    unzip_test(path)
    workbook = load_workbook(path)
    try:
        if list(workbook.sheetnames) != list(SQUARE_SHEETS):
            raise AssertionError(f"Square sheets must be {list(SQUARE_SHEETS)}, got {workbook.sheetnames}")
        extra = FORBIDDEN_SHEETS.intersection(workbook.sheetnames)
        if extra:
            raise AssertionError(f"Square forbidden sheets present: {sorted(extra)}")
        on_hand = workbook["On_Hand"]
        _assert_plain_data_sheet(on_hand, SQUARE_ON_HAND, "On_Hand")
        _assert_no_invented_ma(on_hand, SQUARE_ON_HAND, "On_Hand")
        _assert_no_invented_money(on_hand, SQUARE_ON_HAND, "On_Hand")
        photo_col = header_index(SQUARE_ON_HAND, "photo_folder") + 1
        formula = on_hand.cell(2, photo_col).value
        if not isinstance(formula, str) or not formula.startswith("="):
            raise AssertionError(f"On_Hand photo_folder must be a formula, got {formula!r}")
        if ONEDRIVE_PHOTOS not in formula or "A2" not in formula:
            raise AssertionError(f"On_Hand photo_folder formula drifted: {formula!r}")
        sold = workbook["Sold_Log"]
        _assert_plain_data_sheet(sold, SQUARE_SOLD_LOG, "Sold_Log")
        _assert_no_invented_ma(sold, SQUARE_SOLD_LOG, "Sold_Log")
        if "finance_ref" not in SQUARE_SOLD_LOG:
            raise AssertionError("Sold_Log must keep finance_ref for the sold → Finance path")
        readme = "\n".join(str(workbook["Readme"].cell(row, 1).value or "") for row in range(1, 8)).lower()
        for needle in ("staged site", "square free", "never square save", "not bought", "finance_ref"):
            if needle not in readme:
                raise AssertionError(f"Square Readme missing {needle!r}")
        _assert_no_workbook_images(workbook, "Square.xlsx")
    finally:
        workbook.close()


def verify_finance_book(path: Path) -> None:
    unzip_test(path)
    workbook = load_workbook(path)
    try:
        if list(workbook.sheetnames) != list(FINANCE_SHEETS):
            raise AssertionError(f"Finance sheets must be {list(FINANCE_SHEETS)}, got {workbook.sheetnames}")
        extra = FORBIDDEN_SHEETS.intersection(workbook.sheetnames)
        if extra:
            raise AssertionError(f"Finance forbidden sheets present: {sorted(extra)}")
        sales = workbook["Sales"]
        _assert_plain_data_sheet(sales, FINANCE_SALES, "Sales")
        _assert_no_invented_ma(sales, FINANCE_SALES, "Sales")
        _assert_no_invented_money(sales, FINANCE_SALES, "Sales")
        for header, expected in (
            ("net_usd", finance_net_usd_formula(2)),
            ("taxable_base_usd", finance_taxable_base_formula(2)),
            ("sales_tax_rate", finance_sales_tax_rate_formula(2)),
            ("sales_tax_usd", finance_sales_tax_usd_formula(2)),
        ):
            value = sales.cell(2, header_index(FINANCE_SALES, header) + 1).value
            if value != expected:
                raise AssertionError(f"Sales {header} drifted: {value!r}")
        _assert_plain_data_sheet(workbook["Fees"], FINANCE_FEES, "Fees")
        _assert_no_invented_ma(workbook["Fees"], FINANCE_FEES, "Fees")
        _assert_plain_data_sheet(workbook["Payouts_Transfers"], FINANCE_PAYOUTS, "Payouts_Transfers")
        _assert_plain_data_sheet(workbook["Expenses"], FINANCE_EXPENSES, "Expenses")
        summary = workbook["Tax_Summary"]
        _assert_plain_data_sheet(summary, FINANCE_TAX_SUMMARY_HEADERS, "Tax_Summary")
        got_labels = [
            summary.cell(row, 1).value
            for row in range(2, 2 + len(FINANCE_TAX_SUMMARY_METRICS))
        ]
        want_labels = [label for label, _kind in FINANCE_TAX_SUMMARY_METRICS]
        if got_labels != want_labels:
            raise AssertionError(f"Tax_Summary labels drifted: {got_labels}")
        jan_gross = finance_metric_month_formula("gross", FINANCE_TAX_YEAR, 1)
        if summary.cell(2, 2).value != jan_gross:
            raise AssertionError(f"Tax_Summary B2 drifted: {summary.cell(2, 2).value!r}")
        if summary.cell(2, 14).value != finance_ytd_formula(2):
            raise AssertionError(f"Tax_Summary N2 YTD drifted: {summary.cell(2, 14).value!r}")
        ytd_labels = [label for label, _formula in FINANCE_TAX_SUMMARY_ROWS[1:]]
        if ytd_labels != want_labels:
            raise AssertionError("FINANCE_TAX_SUMMARY_ROWS labels must match monthly metrics")
        note_start = 2 + len(FINANCE_TAX_SUMMARY_METRICS)
        got_notes = [
            summary.cell(note_start + offset, 1).value
            for offset in range(len(FINANCE_TAX_SUMMARY_NOTE_ROWS))
        ]
        want_notes = [label for label, _text in FINANCE_TAX_SUMMARY_NOTE_ROWS]
        if got_notes != want_notes:
            raise AssertionError(f"Tax_Summary helper labels drifted: {got_notes}")
        rate_cell = summary.cell(note_start + 1, 14).value
        if float(rate_cell) != float(SAN_JOSE_SALES_TAX_RATE):
            raise AssertionError(f"Tax_Summary San Jose rate drifted: {rate_cell!r}")
        readme = "\n".join(str(workbook["Readme"].cell(row, 1).value or "") for row in range(1, 10)).lower()
        for needle in (
            "tax-ready",
            "no invented",
            "never square save",
            "square_xlsx_ma",
            "san jose",
            "10.000%",
            "tax on",
            "profit",
            "thang tien huynh",
            "finance_ref",
        ):
            if needle not in readme:
                raise AssertionError(f"Finance Readme missing {needle!r}")
        if "9.375" in readme and "not" not in readme:
            raise AssertionError("Finance Readme must not present 9.375% as current")
        # Keep unused import live so schema deletions fail at import time.
        if FINANCE_PAYOUT_KIND[0] != "square_payout":
            raise AssertionError("FINANCE_PAYOUT_KIND drifted")
        _assert_no_workbook_images(workbook, "Finance.xlsx")
    finally:
        workbook.close()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build(out_dir: Path) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    square_path = out_dir / SQUARE_XLSX_NAME
    finance_path = out_dir / FINANCE_XLSX_NAME

    square = build_square_book()
    square.save(square_path)
    square.close()
    verify_square_book(square_path)

    finance = build_finance_book()
    finance.save(finance_path)
    finance.close()
    verify_finance_book(finance_path)

    print(f"wrote {square_path}")
    print(f"wrote {finance_path}")
    print(f"sha256 Square.xlsx  {_sha256(square_path)}")
    print(f"sha256 Finance.xlsx {_sha256(finance_path)}")
    print("On_Hand empty on purpose — staged site mãs are not bought")
    print("Finance empty on purpose — no invented sales or $")
    print(f"San Jose combined sales tax {SAN_JOSE_SALES_TAX_RATE_DISPLAY} — not 9.375%")
    print(f"sync target: OneDrive {ONEDRIVE_SQUARE}")
    print(f"sync target: OneDrive {ONEDRIVE_FINANCE}")
    return square_path, finance_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=Path("out"),
        help="folder for Square.xlsx and Finance.xlsx (default: ./out)",
    )
    parser.add_argument(
        "--verify-only",
        nargs=2,
        metavar=("SQUARE", "FINANCE"),
        help="verify two existing books instead of rebuilding",
    )
    args = parser.parse_args(argv)
    if args.verify_only:
        square, finance = (Path(p) for p in args.verify_only)
        verify_square_book(square)
        verify_finance_book(finance)
        print(f"verified {square}")
        print(f"verified {finance}")
        print("Square Free = on-hand source of truth. No Square Save.")
        return 0
    if SQUARE_FORMULA_LAST_ROW < SQUARE_TEMPLATE_ROWS + 1:
        raise AssertionError("formula range shorter than the empty runway")
    build(args.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
