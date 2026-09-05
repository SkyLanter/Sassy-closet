#!/usr/bin/env python3
"""Build desktop Official + Wishlist workbooks (no image embeds).

Writes empty-header books with freeze, AutoFilter, and status dropdowns.
Does not invent mã or live inventory rows. photo_link is last on Ma_List
and Candidates — paste a OneDrive share URL; never drop a picture in.

Usage:
  python3 build_boutique_desktop.py
  python3 build_boutique_desktop.py --out-dir ./out
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl.styles import Alignment, PatternFill
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

_KIT_DIR = Path(__file__).resolve().parent
if str(_KIT_DIR) not in sys.path:
    sys.path.insert(0, str(_KIT_DIR))

from schema import (
    BLUSH_ROW,
    BODY_FONT,
    BOT_ACTIVITY,
    BOT_ACTIVITY_WIDTHS,
    BOTS_DRAFT_ONLY,
    CANDIDATE_STATUS,
    CANDIDATE_TYPES,
    CANDIDATES,
    CANDIDATES_WIDTHS,
    MA_LIST,
    MA_LIST_STATUS,
    MA_LIST_WIDTHS,
    ONEDRIVE_OFFICIAL_DESKTOP,
    ONEDRIVE_PHOTOS,
    ONEDRIVE_SOT,
    ONEDRIVE_WISHLIST_DESKTOP,
    ORDER_STATUS,
    ORDERS,
    ORDERS_WIDTHS,
    PHOTO_RULE,
    SQUARE_SOT_LINE,
    add_list_dropdown,
    apply_widths,
    header_index,
    new_boutique_workbook,
    paint_tab,
    style_header_row,
    strip_workbook_images,
    write_how_to_use,
)

DIVIDER = "·  ·  ·  ·  ·  ·  ·  ·  ·  ·  ·  ·"


def _prep_input_grid(ws: Worksheet, headers: tuple[str, ...], rows: int = 40) -> None:
    """Leave blank typed rows (status/link cols stay empty — no fake stock)."""
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", fgColor=BLUSH_ROW)
    for row in range(2, rows + 2):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col, None)
            cell.font = BODY_FONT
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")
    # AutoFilter covers header + reserved blank rows so desktop arrows work.
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{rows + 1}"


def _build_official() -> Workbook:
    wb = new_boutique_workbook()
    cover = wb.active
    cover.title = "How_to_use"
    write_how_to_use(
        cover,
        [
            ("title", "🎀 Sassy Closet Official ✨"),
            ("sub", "Your soft shop book — real pieces, real orders, nothing invented."),
            ("divider", DIVIDER),
            ("section", "How to use"),
            (
                "rule",
                "1. Open Ma_List when a piece is already yours. One row, one mã. "
                "Boss assigns AO / QU / VA / AK / GI / PK / SET + 3 digits — never invent one.",
            ),
            (
                "rule",
                "2. Log every Facebook hello on Orders. Use the status dropdown: "
                "Inquiry → Reserved → Paid → Shipped or Picked up. Cancelled if it falls through.",
            ),
            (
                "rule",
                "3. Bots only draft. They write Bot_Activity in Pacific Time. They never Save in Square.",
            ),
            (
                "rule",
                "4. Desktop: click the arrows on the pink header to filter or sort. "
                "The header stays frozen while you scroll. Paste a URL in photo_link — Excel makes it clickable.",
            ),
            (
                "rule",
                f"5. Keep pictures out of Excel. Drop the file in {ONEDRIVE_PHOTOS}/, "
                "then paste the OneDrive share URL in photo_link.",
            ),
            (
                "rule",
                "6. Square Free is on-hand source of truth. Official Excel / Ma_List is a working copy "
                "and photo index, not live stock. After a good save, Kit syncs this file back to OneDrive.",
            ),
            ("divider", DIVIDER),
            ("section", "What’s on each sheet"),
            ("body", "Ma_List — owned pieces, ma through photo_link"),
            ("body", "Orders — Facebook hellos through pickup or cancel"),
            ("body", "Bot_Activity — drafts only, Pacific Time"),
            ("body", "How_to_use — this cover, text and pastel only"),
            ("divider", DIVIDER),
            ("section", "STATUS legend"),
            ("body", "Ma_List: in_stock · held · sold · archived"),
            ("body", "Wishlist: candidate · watching · skip · bought"),
            ("body", "Orders: Inquiry → Reserved → Paid → Shipped / Picked up / Cancelled"),
            ("divider", DIVIDER),
            ("section", "Photos"),
            ("body", f"Folder: OneDrive {ONEDRIVE_PHOTOS}/ — paste share links only."),
            ("body", "Owned files: AO001.jpg, AO001_2.jpg (match ma). Wishlist finds: #001.jpg, #001_2.jpg."),
            ("body", "notes = captions. photo_link = OneDrive share URL. Never embed a picture in this workbook."),
            ("body", PHOTO_RULE),
            ("divider", DIVIDER),
            ("section", "Shop law"),
            ("body", SQUARE_SOT_LINE),
            ("body", BOTS_DRAFT_ONLY),
            ("body", f"Cursor Cloud Agents build the kit in GitHub. Kit syncs files to OneDrive ({ONEDRIVE_SOT})."),
            ("divider", DIVIDER),
            ("foot", "♡ Empty data sheets are intentional until a real find or buy."),
        ],
    )

    ma = wb.create_sheet("Ma_List")
    style_header_row(ma, MA_LIST)
    apply_widths(ma, MA_LIST_WIDTHS)
    paint_tab(ma)
    add_list_dropdown(ma, header_index(MA_LIST, "status") + 1, MA_LIST_STATUS, prompt="Ma_List status")
    _prep_input_grid(ma, MA_LIST)

    orders = wb.create_sheet("Orders")
    style_header_row(orders, ORDERS, underline_photo_link=False)
    apply_widths(orders, ORDERS_WIDTHS)
    paint_tab(orders)
    add_list_dropdown(orders, header_index(ORDERS, "status") + 1, ORDER_STATUS, prompt="order status")
    _prep_input_grid(orders, ORDERS)

    bots = wb.create_sheet("Bot_Activity")
    style_header_row(bots, BOT_ACTIVITY, underline_photo_link=False)
    apply_widths(bots, BOT_ACTIVITY_WIDTHS)
    paint_tab(bots)
    _prep_input_grid(bots, BOT_ACTIVITY)

    strip_workbook_images(wb)
    return wb


def _build_wishlist() -> Workbook:
    wb = new_boutique_workbook()
    cover = wb.active
    cover.title = "How_to_use"
    write_how_to_use(
        cover,
        [
            ("title", "🎀 Sassy Closet Wishlist ♡"),
            ("sub", "Pretty possibilities only — finds you might buy, not inventory yet."),
            ("divider", DIVIDER),
            ("section", "How to use"),
            (
                "rule",
                "1. One row per find on Candidates. Paste the full Taobao / source URL in "
                "Link (reopen) first — never lose it.",
            ),
            (
                "rule",
                "2. Status starts as candidate. Use the dropdown to move to watching, skip, or bought.",
            ),
            (
                "rule",
                "3. Asia size + cm only. Cost in ¥, sell in $. Photos / notes are captions — "
                "not pictures in this file.",
            ),
            (
                "rule",
                "4. Desktop: filter or sort the pink headers. Paste the OneDrive URL in photo_link — "
                f"Excel makes it clickable. Pictures live in {ONEDRIVE_PHOTOS}/ as #001.jpg.",
            ),
            (
                "rule",
                "5. Stay off Square until Boss confirms bought and assigns a real mã. "
                "Candidates are not inventory.",
            ),
            (
                "rule",
                "6. After a good save, Kit syncs this file back to OneDrive. "
                "Cursor Cloud Agents build the kit in GitHub.",
            ),
            ("divider", DIVIDER),
            ("section", "What’s on each sheet"),
            ("body", "Candidates — # through photo_link, one find per row"),
            ("body", "How_to_use — this cover, text and pastel only"),
            ("divider", DIVIDER),
            ("section", "STATUS legend"),
            ("body", "Wishlist: candidate · watching · skip · bought"),
            ("body", "Orders: Inquiry → Reserved → Paid → Shipped / Picked up / Cancelled"),
            ("divider", DIVIDER),
            ("section", "Photos"),
            ("body", f"Folder: OneDrive {ONEDRIVE_PHOTOS}/ — paste share links only."),
            ("body", "Wishlist files: #001.jpg, #001_2.jpg (match Candidates #). Owned later: AO001.jpg."),
            ("body", "Photos / notes = captions. photo_link = OneDrive share URL. Never embed a picture."),
            ("body", PHOTO_RULE),
            ("divider", DIVIDER),
            ("section", "Shop law"),
            ("body", SQUARE_SOT_LINE),
            ("body", BOTS_DRAFT_ONLY),
            ("body", "Never invent stock. A bought find gets a Boss-assigned mã on Official, then Square Save."),
            ("divider", DIVIDER),
            ("foot", "♡ Empty data sheets are intentional until a real find or buy."),
        ],
    )

    cand = wb.create_sheet("Candidates")
    style_header_row(cand, CANDIDATES)
    apply_widths(cand, CANDIDATES_WIDTHS)
    paint_tab(cand)
    add_list_dropdown(cand, header_index(CANDIDATES, "Status") + 1, CANDIDATE_STATUS, prompt="wishlist status")
    add_list_dropdown(cand, header_index(CANDIDATES, "Type") + 1, CANDIDATE_TYPES, prompt="type")
    _prep_input_grid(cand, CANDIDATES)

    strip_workbook_images(wb)
    return wb


def _assert_no_inventory_values(ws: Worksheet, headers: list[object]) -> None:
    """Builders must not ship live or demo stock in data rows."""
    from schema import looks_like_demo_row, parse_ma

    max_col = len(headers)
    for row in range(2, (ws.max_row or 1) + 1):
        values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
        if looks_like_demo_row(values):
            raise AssertionError(f"{ws.title} row {row}: demo/fake inventory is forbidden")
        for header, value in zip(headers, values, strict=False):
            if value is None or str(value).strip() == "":
                continue
            name = str(header or "").strip().lower()
            if name in {"ma", "mã"} and parse_ma(value) is not None:
                raise AssertionError(
                    f"{ws.title} row {row}: builder must not commit mã {value!r}"
                )


def verify_desktop_book(path: Path, *, kind: str) -> None:
    """Fail loud if freeze / filter / dropdowns / photo_link contract broke."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        if getattr(wb, "_images", None):
            raise AssertionError(f"{path.name}: workbook-level images are forbidden")
        for ws in wb.worksheets:
            if getattr(ws, "_images", None):
                raise AssertionError(f"{path.name} / {ws.title}: embedded images are forbidden")
        if kind == "official":
            ws = wb["Ma_List"]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            if headers != list(MA_LIST):
                raise AssertionError(f"Ma_List headers drifted: {headers}")
            if not headers or headers[-1] != "photo_link":
                raise AssertionError(f"Ma_List photo_link must be last, got {headers[-1]!r}")
            _assert_no_inventory_values(ws, headers)
            if ws.freeze_panes != "A2":
                raise AssertionError(f"Ma_List freeze_panes={ws.freeze_panes}")
            if not ws.auto_filter.ref or not ws.auto_filter.ref.startswith("A1:"):
                raise AssertionError(f"Ma_List autofilter={ws.auto_filter.ref}")
            if not ws.data_validations.dataValidation:
                raise AssertionError("Ma_List missing status dropdown")
            if wb["Orders"].freeze_panes != "A2":
                raise AssertionError("Orders freeze missing")
            if not wb["Bot_Activity"].auto_filter.ref:
                raise AssertionError("Bot_Activity autofilter missing")
        else:
            ws = wb["Candidates"]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            if headers != list(CANDIDATES):
                raise AssertionError(f"Candidates headers drifted: {headers}")
            if not headers or headers[-1] != "photo_link":
                raise AssertionError("Candidates photo_link must be last")
            _assert_no_inventory_values(ws, headers)
            if ws.freeze_panes != "A2":
                raise AssertionError(f"Candidates freeze_panes={ws.freeze_panes}")
            if not ws.auto_filter.ref:
                raise AssertionError("Candidates autofilter missing")
            if len(ws.data_validations.dataValidation) < 2:
                raise AssertionError("Candidates needs Status + Type dropdowns")
    finally:
        wb.close()


def build(out_dir: Path) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    official_path = out_dir / Path(ONEDRIVE_OFFICIAL_DESKTOP).name
    wishlist_path = out_dir / Path(ONEDRIVE_WISHLIST_DESKTOP).name

    official = _build_official()
    official.save(official_path)
    official.close()
    verify_desktop_book(official_path, kind="official")

    wishlist = _build_wishlist()
    wishlist.save(wishlist_path)
    wishlist.close()
    verify_desktop_book(wishlist_path, kind="wishlist")

    print(f"wrote {official_path}")
    print(f"wrote {wishlist_path}")
    print("headers empty on purpose — no live inventory, no customer PII")
    print(f"sync target: OneDrive {ONEDRIVE_OFFICIAL_DESKTOP}")
    print(f"sync target: OneDrive {ONEDRIVE_WISHLIST_DESKTOP}")
    return official_path, wishlist_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=Path("out"),
        help="folder for the two desktop .xlsx files (default: ./out)",
    )
    parser.add_argument(
        "--verify-only",
        nargs=2,
        metavar=("OFFICIAL", "WISHLIST"),
        help="verify two existing desktop books instead of rebuilding",
    )
    args = parser.parse_args(argv)
    if args.verify_only:
        official, wishlist = (Path(p) for p in args.verify_only)
        verify_desktop_book(official, kind="official")
        verify_desktop_book(wishlist, kind="wishlist")
        print(f"verified {official}")
        print(f"verified {wishlist}")
        print(SQUARE_SOT_LINE)
        return 0
    build(args.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
