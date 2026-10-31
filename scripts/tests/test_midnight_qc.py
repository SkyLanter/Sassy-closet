"""Midnight QC on fixtures. No live Taobao, Blob, or sell host."""

from __future__ import annotations

import json
import subprocess
import sys
import urllib.request
from pathlib import Path

from openpyxl import Workbook

import freshness_check
from freshness_check import ALIAS_ORIGIN, PRIMARY_ORIGIN, SELL_ORIGINS, SourceState
from held_policy import load_held_policy
from midnight_qc import (
    PR105_URL,
    SUMMARY_KEYS,
    HttpResult,
    fetch_night,
    main,
    parse_home_prices,
    summary_from_payload,
)

ROOT = Path(__file__).resolve().parents[2]
HELD_BYTES = (ROOT / "data" / "held.json").read_bytes()
EXAMPLE_KEYS = (
    "date",
    "checkedAt",
    "timezone",
    "overall",
    "dataOk",
    "sellSite",
    "metaLight",
    "blobCount",
    "sellCount",
    "hubCount",
    "intakeCount",
    "heldCount",
    "s06Usd",
    "a02TitleEn",
    "k06OnBlob",
    "k06Pdp",
    "v05OnBlob",
    "v05Pdp",
    "v05IntakeCostCny",
    "v05CreatedAt",
    "v05SellUsd",
    "blobUpdatedAt",
    "blobSha256",
    "blobEtag",
    "blobLastModifiedHeader",
    "blobAgeSecondsAtCheckedAt",
    "hubSha256",
    "hubLastModified",
    "homeSha256",
    "aliasFirstFetchSha256",
    "homeRefetchSha256",
    "mainSha",
    "mainHeldList",
    "mainIncludesK06",
    "mainIncludesV05",
    "mainNeverList",
    "saveRouteCallsHeldCheck",
    "shopMintRefusesHeld",
    "shopSaveRefusesOffLockS06",
    "blobHubUsdMismatches",
    "intakeVsBlobMismatches",
    "heldAbsent",
    "q02Absent",
    "b01Referenced",
    "imagesChecked",
    "imagesOk",
    "newIntakeSince20261007",
    "newIntakeSince20261009",
    "newIntakeSince20261011",
    "pr105Url",
    "pr105Draft",
    "pr105Head",
    "pr105AddsV05",
    "pr101Head",
    "pr101Draft",
    "pr101State",
    "newDraftPr",
    "newDraftPrHead",
    "bossPing",
    "changedSince20261007",
    "changedSince20261009",
    "changedSince20261010",
    "changedSince20261011",
    "unchangedOpenPrs",
    "notes",
    "freshnessGuards",
    "freshnessHub",
)


def _hosts(status: int, price: str | None = None) -> dict:
    cell: dict[str, object] = {"status": status}
    if price is not None:
        cell["price_usd"] = price
    return {origin: dict(cell) for origin in SELL_ORIGINS}


def _passing_payload() -> dict:
    policy = load_held_policy()
    pdp = {code: _hosts(404) for code in [*policy.held, "S16", "Q02"]}
    pdp["S06"] = _hosts(200, "25")
    pdp["A02"] = _hosts(200, "23")
    prices = {"S06": "25", "A02": "23"}
    return {
        "day": "2026-10-14",
        "checkedAt": "2026-10-14T08:17:39Z",
        "timezone": "America/Los_Angeles",
        "origins": list(SELL_ORIGINS),
        "blob": {
            "state": "ok",
            "updatedAt": "2026-09-30T15:56:17.273Z",
            "sha256": "abc",
            "etag": "etag",
            "lastModified": "Wed, 07 Oct 2026 16:54:37 GMT",
            "ageSeconds": 573782,
            "products": [
                {"ma": "S06", "priceUsd": 25, "titleEn": "Washed distressed overalls", "status": "available"},
                {"ma": "A02", "priceUsd": 23, "titleEn": "Polka-dot top", "status": "available"},
            ],
        },
        "hub": {
            "state": "ok",
            "sha256": "hub",
            "lastModified": "2026-09-30T15:58:28Z",
            "rows": [
                {"ma": "S06", "price_usd": "25"},
                {"ma": "A02", "price_usd": "23"},
            ],
        },
        "intake": {
            "state": "ok",
            "submissions": [
                {"ma": "S06", "cost_cny": "0", "sell_usd": "", "created_at": "2026-09-01T00:00:00Z"},
                {"ma": "A02", "cost_cny": "0", "sell_usd": "22", "created_at": "2026-09-01T00:00:00Z"},
                {"ma": "K06", "cost_cny": "0", "sell_usd": "", "created_at": "2026-10-06T00:00:00Z"},
                {
                    "ma": "V05",
                    "cost_cny": "0",
                    "sell_usd": None,
                    "created_at": "2026-10-09T00:00:00Z",
                },
                {"ma": "S16", "cost_cny": "0", "sell_usd": "", "created_at": "2026-10-12T00:00:00Z"},
            ],
        },
        "pdp": pdp,
        "homes": {
            PRIMARY_ORIGIN: {"sha256": "home", "refetchSha256": "home-again", "prices": dict(prices)},
            ALIAS_ORIGIN: {"sha256": "alias", "prices": dict(prices)},
        },
        "images": [{"ok": True, "status": 206}, {"ok": True, "status": 206}],
    }


def _summary(payload: dict | None = None) -> tuple[dict, str]:
    return summary_from_payload(payload or _passing_payload())


def test_summary_keys_match_the_midnight_qc_shape() -> None:
    assert SUMMARY_KEYS == EXAMPLE_KEYS
    summary, _markdown = _summary()
    assert tuple(summary) == EXAMPLE_KEYS


def test_passing_fixture_checks_counts_s06_and_held_404s() -> None:
    summary, markdown = _summary()
    assert summary["overall"] == "PASS"
    assert summary["dataOk"] == "PASS"
    assert summary["sellSite"] == "PASS"
    assert summary["metaLight"] == "NOTE"
    assert summary["bossPing"] is False
    assert summary["s06Usd"] == 25
    assert summary["a02TitleEn"] == "Polka-dot top"
    assert summary["blobCount"] == 2
    assert summary["hubCount"] == 2
    assert summary["sellCount"] == 2
    assert summary["intakeCount"] == 5
    assert summary["heldCount"] == 3
    assert summary["k06OnBlob"] is False
    assert summary["k06Pdp"] == 404
    assert summary["v05OnBlob"] is False
    assert summary["v05Pdp"] == 404
    assert summary["v05IntakeCostCny"] == "0"
    assert summary["v05CreatedAt"] == "2026-10-09T00:00:00Z"
    assert summary["v05SellUsd"] is None
    assert summary["blobHubUsdMismatches"] == []
    assert summary["intakeVsBlobMismatches"] == ["A02 22→23"]
    assert summary["heldAbsent"] is True
    assert summary["q02Absent"] is True
    assert summary["b01Referenced"] is False
    assert summary["imagesChecked"] == 2
    assert summary["imagesOk"] == 2
    assert summary["freshnessGuards"] == "PASS"
    assert summary["freshnessHub"] == "ok"
    assert summary["newIntakeSince20261007"] == ["V05", "S16"]
    assert summary["newIntakeSince20261009"] == ["V05", "S16"]
    assert summary["newIntakeSince20261011"] == ["S16"]
    assert summary["mainIncludesK06"] is True
    assert summary["mainIncludesV05"] is True
    assert summary["mainNeverList"] == []
    held = json.loads(HELD_BYTES)
    assert summary["mainHeldList"] == held["heldIncomplete"]
    assert summary["saveRouteCallsHeldCheck"] is False
    assert summary["shopMintRefusesHeld"] is True
    assert summary["shopSaveRefusesOffLockS06"] is True
    assert summary["pr105Url"] == PR105_URL
    assert summary["pr105Head"] is None
    assert summary["homeSha256"] == "home"
    assert summary["aliasFirstFetchSha256"] == "alias"
    assert summary["homeRefetchSha256"] == "home-again"
    assert "No Boss ping" in str(summary["notes"])
    assert "No live price was changed" in str(summary["notes"])
    assert "Meta is not CLEAR" in str(summary["notes"])
    assert "S06 is $25" in str(summary["notes"])
    for heading in (
        "# DATA OK — Midnight QC 2026-10-14",
        "## Counts",
        "## Blob↔hub mismatch table",
        "## Intake mismatch table",
        "## WARN / NOTE",
        "## FAIL",
        "## Every live mã",
        "## Held, and Q02",
        "## Sources",
    ):
        assert heading in markdown
    assert "\nNone.\n" in markdown
    assert "W-held-gap" not in markdown
    assert "S16" in markdown
    assert (ROOT / "data" / "held.json").read_bytes() == HELD_BYTES


def test_held_gap_follows_the_live_held_list() -> None:
    payload = _passing_payload()
    payload["intake"]["submissions"].append(
        {"ma": "Z90", "cost_cny": "0", "sell_usd": "", "created_at": "2026-01-01T00:00:00Z"}
    )
    payload["pdp"]["Z90"] = _hosts(404)
    summary, markdown = _summary(payload)
    assert "S16" in summary["mainHeldList"]
    assert "Z90" not in summary["mainHeldList"]
    gap_lines = [line for line in markdown.splitlines() if "W-held-gap" in line]
    assert len(gap_lines) == 1
    assert "Z90" in gap_lines[0]
    assert "S16" not in gap_lines[0]
    assert "K06" not in gap_lines[0]
    assert "A27" not in gap_lines[0]
    assert (ROOT / "data" / "held.json").read_bytes() == HELD_BYTES


def test_context_cannot_claim_meta_clear_or_a_boss_ping() -> None:
    payload = _passing_payload()
    payload["context"] = {"bossPing": True, "metaLight": "CLEAR", "pr105Head": "abc", "pr105Draft": True}
    summary, _markdown = _summary(payload)
    assert summary["bossPing"] is False
    assert summary["metaLight"] == "NOTE"
    assert summary["pr105Head"] == "abc"
    assert summary["pr105Draft"] is True


def test_s06_not_25_fails() -> None:
    payload = _passing_payload()
    for product in payload["blob"]["products"]:
        if product["ma"] == "S06":
            product["priceUsd"] = 24
    for row in payload["hub"]["rows"]:
        if row["ma"] == "S06":
            row["price_usd"] = "24"
    payload["pdp"]["S06"] = _hosts(200, "24")
    payload["homes"][PRIMARY_ORIGIN]["prices"]["S06"] = "24"
    payload["homes"][ALIAS_ORIGIN]["prices"]["S06"] = "24"
    summary, markdown = _summary(payload)
    assert summary["overall"] == "FAIL"
    assert summary["dataOk"] == "FAIL"
    assert summary["s06Usd"] == 24
    assert summary["freshnessGuards"] == "FAIL"
    assert "S06 is $24" in markdown


def test_held_pdp_200_fails() -> None:
    payload = _passing_payload()
    payload["pdp"]["K06"][PRIMARY_ORIGIN] = {"status": 200, "price_usd": "48"}
    summary, _markdown = _summary(payload)
    assert summary["overall"] == "FAIL"
    assert summary["sellSite"] == "FAIL"
    assert summary["k06Pdp"] == 200
    assert summary["heldAbsent"] is False
    assert summary["freshnessGuards"] == "FAIL"


def test_blob_hub_usd_mismatch_fails() -> None:
    payload = _passing_payload()
    for row in payload["hub"]["rows"]:
        if row["ma"] == "A02":
            row["price_usd"] = "99"
    summary, markdown = _summary(payload)
    assert summary["overall"] == "FAIL"
    assert summary["dataOk"] == "FAIL"
    assert summary["sellSite"] == "PASS"
    assert summary["blobHubUsdMismatches"] == ["A02 23→99"]
    assert "MISMATCH" in markdown


def test_intake_lag_is_a_warn_and_not_a_fail() -> None:
    summary, markdown = _summary()
    assert summary["overall"] == "PASS"
    assert summary["intakeVsBlobMismatches"] == ["A02 22→23"]
    assert "W-intake" in markdown
    assert "LAG 22→23" in markdown


def test_hub_skip_does_not_fail_or_rewrite_the_parser() -> None:
    payload = _passing_payload()
    payload["hub"] = {"state": "skip", "detail": "Official sheet is missing ma or list_price"}
    summary, markdown = _summary(payload)
    assert summary["overall"] == "PASS"
    assert summary["freshnessHub"] == "skip"
    assert summary["hubCount"] is None
    assert summary["blobHubUsdMismatches"] == []
    assert "W-hub" in markdown
    assert "list_price" in markdown
    assert "hub parser was not rewritten" in markdown


def test_row1_sell_usd_workbook_is_the_live_hub_header(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Official"
    sheet.cell(1, 1, "ma")
    sheet.cell(1, 2, "sell_usd")
    sheet.cell(2, 1, "S06")
    sheet.cell(2, 2, 25)
    sheet.cell(3, 1, "A02")
    sheet.cell(3, 2, 23)
    path = tmp_path / "row1.xlsx"
    book.save(path)
    payload = _passing_payload()
    del payload["hub"]
    payload["hubFile"] = str(path)
    summary, markdown = _summary(payload)
    assert summary["overall"] == "PASS"
    assert summary["freshnessHub"] == "ok"
    assert summary["hubCount"] == 2
    assert "missing ma or list_price" not in markdown


def test_row4_list_price_workbook_is_the_parser_result(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Official"
    sheet.cell(4, 1, "ma")
    sheet.cell(4, 2, "list_price")
    sheet.cell(5, 1, "S06")
    sheet.cell(5, 2, 25)
    sheet.cell(6, 1, "A02")
    sheet.cell(6, 2, 23)
    path = tmp_path / "row4.xlsx"
    book.save(path)
    payload = _passing_payload()
    del payload["hub"]
    payload["hubFile"] = str(path)
    summary, _markdown = _summary(payload)
    assert summary["overall"] == "PASS"
    assert summary["freshnessHub"] == "ok"
    assert summary["hubCount"] == 2
    assert summary["blobHubUsdMismatches"] == []
    assert summary["s06Usd"] == 25


def test_ops_held_pdp_200_fails_without_editing_held_json() -> None:
    payload = _passing_payload()
    payload["pdp"]["S16"][ALIAS_ORIGIN] = {"status": 200, "price_usd": "30"}
    summary, markdown = _summary(payload)
    assert summary["overall"] == "FAIL"
    assert summary["heldAbsent"] is False
    assert "S16" in markdown
    assert (ROOT / "data" / "held.json").read_bytes() == HELD_BYTES


def test_q02_404_is_not_a_fail_and_q02_on_blob_is() -> None:
    absent, _markdown = _summary()
    assert absent["q02Absent"] is True
    assert absent["overall"] == "PASS"
    payload = _passing_payload()
    payload["blob"]["products"].append({"ma": "Q02", "priceUsd": 10, "titleEn": "Rejected", "status": "available"})
    payload["hub"]["rows"].append({"ma": "Q02", "price_usd": "10"})
    payload["pdp"]["Q02"] = _hosts(200, "10")
    for origin in SELL_ORIGINS:
        payload["homes"][origin]["prices"]["Q02"] = "10"
    live, _markdown = _summary(payload)
    assert live["q02Absent"] is False
    assert live["overall"] == "FAIL"
    assert live["dataOk"] == "FAIL"


def test_alias_and_home_price_mismatches_fail_sell() -> None:
    alias = _passing_payload()
    alias["pdp"]["A02"][ALIAS_ORIGIN] = {"status": 200, "price_usd": "99"}
    alias_summary, _markdown = _summary(alias)
    assert alias_summary["overall"] == "FAIL"
    assert alias_summary["sellSite"] == "FAIL"
    home = _passing_payload()
    home["homes"][PRIMARY_ORIGIN]["prices"]["A02"] = "99"
    home_summary, _markdown = _summary(home)
    assert home_summary["overall"] == "FAIL"
    assert home_summary["sellSite"] == "FAIL"
    assert home_summary["dataOk"] == "PASS"


def test_gail_in_catalog_text_sets_b01_referenced() -> None:
    payload = _passing_payload()
    payload["blob"]["products"][1]["titleEn"] = "Polka-dot top gail"
    summary, _markdown = _summary(payload)
    assert summary["b01Referenced"] is True
    assert summary["overall"] == "PASS"


def test_parse_home_prices_reads_embedded_price_usd() -> None:
    html = '{"ma":"S06","titleEn":"Washed distressed overalls","priceUsd":25}'
    assert parse_home_prices(html)["S06"] == 25


def test_cli_writes_reports_and_does_not_touch_held_or_network(tmp_path: Path, monkeypatch) -> None:
    def explode(*_args, **_kwargs):
        raise AssertionError("network")

    monkeypatch.setattr(urllib.request, "urlopen", explode)
    fixture = tmp_path / "night.json"
    fixture.write_text(json.dumps(_passing_payload()), encoding="utf-8")
    out = tmp_path / "qc"
    passed = main(["--fixture", str(fixture), "--out", str(out), "--date", "2026-10-14"])
    assert passed == 0
    summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    assert summary["overall"] == "PASS"
    assert (out / "data-ok.md").is_file()
    assert (ROOT / "data" / "held.json").read_bytes() == HELD_BYTES

    fail_payload = _passing_payload()
    fail_payload["pdp"]["V05"][PRIMARY_ORIGIN] = {"status": 200, "price_usd": "27"}
    fail_file = tmp_path / "fail.json"
    fail_file.write_text(json.dumps(fail_payload), encoding="utf-8")
    failed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "midnight_qc.py"),
            "--fixture",
            str(fail_file),
            "--out",
            str(out),
            "--date",
            "2026-10-14",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert failed.returncode == 1, failed.stderr
    assert json.loads((out / "summary.json").read_text(encoding="utf-8"))["overall"] == "FAIL"
    assert (ROOT / "data" / "held.json").read_bytes() == HELD_BYTES


def test_live_fetch_uses_freshness_parsers_and_skips_taobao(monkeypatch) -> None:
    calls: list[str] = []

    def fake(url: str, timeout: float = 20) -> HttpResult:
        del timeout
        calls.append(url)
        if "taobao" in url or "facebook" in url or "instagram" in url:
            raise AssertionError(url)
        if url.endswith("/api/submissions"):
            body = json.dumps(
                {
                    "submissions": [
                        {
                            "ma": "S06",
                            "sell_usd": "",
                            "cost_cny": "0",
                            "link": "https://example.invalid/item.htm?id=1",
                        }
                    ]
                }
            ).encode()
            return HttpResult(200, body, {}, None)
        if "/m/" in url:
            if url.endswith("/S06"):
                html = (
                    '<script type="application/ld+json">'
                    '{"@type":"Product","offers":{"price":"25.00","priceCurrency":"USD"}}'
                    "</script>"
                )
                return HttpResult(200, html.encode(), {}, None)
            return HttpResult(404, b"no", {}, None)
        if "catalog.v1.json" in url:
            body = json.dumps(
                {
                    "updatedAt": "2026-09-30T15:56:17.273Z",
                    "products": [
                        {"ma": "S06", "priceUsd": 25, "titleEn": "Washed distressed overalls", "status": "available"}
                    ],
                }
            ).encode()
            return HttpResult(
                200,
                body,
                {"etag": '"abc"', "last-modified": "Wed, 07 Oct 2026 16:54:37 GMT", "age": "5"},
                None,
            )
        if url.rstrip("/") in SELL_ORIGINS:
            html = '{"ma":"S06","titleEn":"Washed distressed overalls","priceUsd":25}'
            return HttpResult(200, html.encode(), {}, None)
        raise AssertionError(url)

    monkeypatch.setattr("midnight_qc.http_get", fake)
    monkeypatch.setattr(
        freshness_check,
        "_fetch_hub",
        lambda: (None, SourceState("skip", "Official sheet is missing ma or list_price")),
    )
    payload = fetch_night("2026-10-14", check_images=False)
    summary, _markdown = summary_from_payload(payload)
    assert all("taobao" not in url for url in calls)
    assert not any(url.endswith("item.htm?id=1") for url in calls)
    assert summary["overall"] == "PASS"
    assert summary["s06Usd"] == 25
    assert summary["blobCount"] == 1
    assert summary["sellCount"] == 1
    assert summary["freshnessHub"] == "skip"
    assert summary["blobEtag"] == "abc"
    assert summary["blobAgeSecondsAtCheckedAt"] == 5
    assert summary["k06Pdp"] == 404
    assert summary["v05Pdp"] == 404
    assert (ROOT / "data" / "held.json").read_bytes() == HELD_BYTES


def test_script_does_not_duplicate_freshness_or_meta() -> None:
    source = (ROOT / "scripts" / "midnight_qc.py").read_text(encoding="utf-8")
    assert "def read_hub_official" not in source
    assert "header_row_for_sheet" not in source
    assert "fetch_live(" not in source
    assert "list_price" not in source
    assert "taobao" + ".com" not in source
    assert "facebook.com" not in source
    assert "instagram.com" not in source
    assert "tesla" not in source.lower()
    assert "read_hub_official" in source
    assert "evaluate(" in source
    assert 'method="GET"' in source
    assert "does not write data/held.json" in source
